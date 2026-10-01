import hashlib
import hmac
import json
import io
import mimetypes
import os
import threading
import base64
import httpx
from datetime import datetime, timezone, timedelta
from uuid import UUID, uuid4
from typing import List

from fastapi import FastAPI, File, UploadFile, Request, Form, HTTPException, BackgroundTasks
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse
from supabase import create_client, Client

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload


# ============================================================
# DATAVAULT DLP | GM INGENIEROS Y CONSULTORES
# ============================================================

app = FastAPI(
    title="DataVault DLP API | GM Ingenieros y Consultores"
)

# Exponer la carpeta física "images" para archivos estáticos
app.mount("/images", StaticFiles(directory="images"), name="images")

_cors_public_base = os.getenv("PUBLIC_BASE_URL", "").strip().rstrip("/")
CORS_ALLOWED_ORIGINS = [
    origin
    for origin in (
        _cors_public_base,
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    )
    if origin
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
    allow_headers=["Authorization", "Content-Type", "X-Telegram-Bot-Api-Secret-Token"],
)


# ============================================================
# CONFIGURACIÓN DESDE RAILWAY Y GOOGLE DRIVE
# ============================================================

TELEGRAM_TOKEN = os.getenv(
    "TELEGRAM_TOKEN",
    ""
).strip()


TELEGRAM_WEBHOOK_SECRET = os.getenv(
    "TELEGRAM_WEBHOOK_SECRET",
    ""
).strip()


AUTHORIZED_CHAT_IDS_RAW = os.getenv(
    "AUTHORIZED_CHAT_IDS",
    "8893414961,6718944855"
).strip()


SUPABASE_URL = os.getenv(
    "SUPABASE_URL",
    "https://crujlbbhtkcithullgfs.supabase.co"
).strip()


SUPABASE_KEY = os.getenv(
    "SUPABASE_KEY",
    ""
).strip()


SUPABASE_SERVICE_ROLE_KEY = os.getenv(
    "SUPABASE_SERVICE_ROLE_KEY",
    ""
).strip()


PUBLIC_BASE_URL = os.getenv(
    "PUBLIC_BASE_URL",
    ""
).strip().rstrip("/")


RAILWAY_PUBLIC_DOMAIN = os.getenv(
    "RAILWAY_PUBLIC_DOMAIN",
    ""
).strip()


GOOGLE_CLIENT_ID = os.getenv(
    "GOOGLE_CLIENT_ID",
    ""
).strip()


GOOGLE_CLIENT_SECRET = os.getenv(
    "GOOGLE_CLIENT_SECRET",
    ""
).strip()


GOOGLE_REFRESH_TOKEN = os.getenv(
    "GOOGLE_REFRESH_TOKEN",
    ""
).strip()


GOOGLE_FOLDER_ID = os.getenv(
    "GOOGLE_FOLDER_ID",
    ""
).strip()


# Raíz autorizada para la nueva navegación. "root" permite trabajar con toda
# Mi unidad usando únicamente las credenciales del backend.
GOOGLE_DRIVE_ROOT_ID = os.getenv(
    "GOOGLE_DRIVE_ROOT_ID",
    "root"
).strip() or "root"


DRIVE_CACHE_TTL_SECONDS = max(
    10,
    int(os.getenv("DRIVE_CACHE_TTL_SECONDS", "60"))
)


SUPABASE_TEMP_BUCKET = os.getenv(
    "SUPABASE_TEMP_BUCKET",
    "custodia-pendiente"
).strip() or "custodia-pendiente"


if not PUBLIC_BASE_URL and RAILWAY_PUBLIC_DOMAIN:

    PUBLIC_BASE_URL = (
        f"https://{RAILWAY_PUBLIC_DOMAIN}"
    ).rstrip("/")


DECISION_LOCK = threading.Lock()

# Límites de prueba para carga por lotes. Se pueden cambiar en Railway.
MAX_BATCH_FILES = int(os.getenv("MAX_BATCH_FILES", "100"))
MAX_FILE_SIZE_MB = int(os.getenv("MAX_FILE_SIZE_MB", "25"))
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

# Cliente HTTP reutilizable para Supabase Auth / Telegram.
# Evita abrir una conexión nueva en cada llamada externa.
HTTP_TIMEOUT_SECONDS = float(os.getenv("HTTP_TIMEOUT_SECONDS", "20"))
HTTP_CLIENT = httpx.Client(
    timeout=httpx.Timeout(HTTP_TIMEOUT_SECONDS),
    limits=httpx.Limits(
        max_connections=40,
        max_keepalive_connections=20,
    ),
    follow_redirects=True,
)


# ============================================================
# ESTADOS / ERRORES NORMALIZADOS (PUNTOS 27-28)
# ============================================================

ESTADO_PENDIENTE = "PENDIENTE"
ESTADO_APROBADO = "APROBADO"
ESTADO_RECHAZADO = "RECHAZADO"
ESTADO_ERROR = "ERROR"

ESTADO_ARCHIVO_ACTIVO = "ACTIVO"
ESTADO_ARCHIVO_ELIMINADO = "ELIMINADO"
ESTADO_ARCHIVO_ELIMINADO_EXTERNAMENTE = "ELIMINADO_EXTERNAMENTE"


def normalizar_estado(valor, predeterminado=ESTADO_PENDIENTE) -> str:
    texto = str(valor or "").strip().upper().replace(" ", "_")
    alias = {
        "PENDING": ESTADO_PENDIENTE,
        "EN_COLA": ESTADO_PENDIENTE,
        "PENDIENTE": ESTADO_PENDIENTE,
        "APPROVED": ESTADO_APROBADO,
        "APROBADA": ESTADO_APROBADO,
        "APROBADO": ESTADO_APROBADO,
        "REJECTED": ESTADO_RECHAZADO,
        "RECHAZADA": ESTADO_RECHAZADO,
        "RECHAZADO": ESTADO_RECHAZADO,
        "FAILED": ESTADO_ERROR,
        "FALLO": ESTADO_ERROR,
        "ERROR": ESTADO_ERROR,
    }
    return alias.get(texto, texto or predeterminado)


def normalizar_estado_archivo(valor, predeterminado=ESTADO_ARCHIVO_ACTIVO) -> str:
    texto = str(valor or "").strip().upper().replace(" ", "_")
    alias = {
        "ACTIVE": ESTADO_ARCHIVO_ACTIVO,
        "ACTIVO": ESTADO_ARCHIVO_ACTIVO,
        "DELETED": ESTADO_ARCHIVO_ELIMINADO,
        "ELIMINADO": ESTADO_ARCHIVO_ELIMINADO,
        "TRASHED": ESTADO_ARCHIVO_ELIMINADO,
        "ELIMINADO_EXTERNAMENTE": ESTADO_ARCHIVO_ELIMINADO_EXTERNAMENTE,
    }
    return alias.get(texto, texto or predeterminado)


def detalle_error(codigo: str, mensaje: str, contexto=None) -> dict:
    detalle = {
        "code": str(codigo or "ERROR").strip().upper(),
        "message": str(mensaje or "Error no especificado."),
    }
    if contexto not in (None, {}, []):
        detalle["context"] = contexto
    return detalle


def error_api(status_code: int, codigo: str, mensaje: str, contexto=None):
    return HTTPException(
        status_code=status_code,
        detail=detalle_error(codigo, mensaje, contexto),
    )


# ============================================================
# GOOGLE DRIVE UPLOAD
# ============================================================

def google_drive_configurado() -> bool:
    return all([
        GOOGLE_CLIENT_ID,
        GOOGLE_CLIENT_SECRET,
        GOOGLE_REFRESH_TOKEN,
    ])


def obtener_servicio_google_drive():
    if not google_drive_configurado():
        raise RuntimeError("Faltan variables de Google Drive")

    creds = Credentials(
        token=None,
        refresh_token=GOOGLE_REFRESH_TOKEN,
        token_uri="https://oauth2.googleapis.com/token",
        client_id=GOOGLE_CLIENT_ID,
        client_secret=GOOGLE_CLIENT_SECRET,
    )

    return build(
        "drive",
        "v3",
        credentials=creds,
        cache_discovery=False
    )


def crear_carpeta_google_drive(
    service,
    nombre: str,
    parent_id: str,
    app_properties: dict | None = None,
):
    metadata = {
        "name": nombre,
        "mimeType": "application/vnd.google-apps.folder",
        "parents": [parent_id],
    }
    if app_properties:
        metadata["appProperties"] = app_properties

    respuesta = (
        service
        .files()
        .create(
            body=metadata,
            fields="id,name",
            supportsAllDrives=True,
        )
        .execute()
    )

    carpeta_id = respuesta.get("id")
    if not carpeta_id:
        raise RuntimeError(f"Google Drive no devolvió ID para la carpeta {nombre}")

    try:
        invalidar_cache_drive()
    except Exception:
        pass

    return carpeta_id


def subir_archivo_google_drive_en_carpeta(
    service,
    nombre_archivo: str,
    contenido_bytes: bytes,
    parent_id: str,
    app_properties: dict | None = None,
):
    if not contenido_bytes:
        raise RuntimeError(f"El archivo {nombre_archivo} está vacío")

    mime_type = (
        mimetypes.guess_type(nombre_archivo)[0]
        or "application/octet-stream"
    )

    media = MediaIoBaseUpload(
        io.BytesIO(contenido_bytes),
        mimetype=mime_type,
        resumable=False,
    )

    respuesta = (
        service
        .files()
        .create(
            body={
                "name": nombre_archivo,
                "parents": [parent_id],
                **({"appProperties": app_properties} if app_properties else {}),
            },
            media_body=media,
            fields="id,name",
            supportsAllDrives=True,
        )
        .execute()
    )

    archivo_id = respuesta.get("id")
    if not archivo_id:
        raise RuntimeError(f"Google Drive no devolvió ID para {nombre_archivo}")

    try:
        invalidar_cache_drive()
    except Exception:
        pass

    return archivo_id


def subir_a_google_drive(
    nombre_archivo: str,
    contenido_bytes: bytes,
    auditoria_id: str | None = None,
    parent_id: str | None = None,
):
    """Sube un archivo individual directamente desde RAM a Google Drive."""

    if not google_drive_configurado():
        print("[DRIVE CONFIG ERROR] Faltan variables de Google Drive")
        return None

    if not contenido_bytes:
        print("[DRIVE ERROR] El contenido recibido está vacío")
        return None

    destino_parent_id = (
        str(parent_id or GOOGLE_FOLDER_ID).strip()
        or GOOGLE_FOLDER_ID
    )

    try:
        service = obtener_servicio_google_drive()
        drive_id = subir_archivo_google_drive_en_carpeta(
            service,
            nombre_archivo,
            contenido_bytes,
            destino_parent_id,
            app_properties=(
                {"datavault_auditoria_id": str(auditoria_id)}
                if auditoria_id else None
            ),
        )

        print(
            f"[DRIVE SUCCESS] {nombre_archivo} "
            f"subido con ID: {drive_id}"
        )
        return drive_id

    except Exception as e:
        print(f"[DRIVE EXCEPTION] Error al transferir a Drive: {e}")
        return None


def obtener_carpeta_desde_ruta(ruta_relativa: str):
    ruta = str(ruta_relativa or "").replace("\\", "/").strip("/")
    partes = [p for p in ruta.split("/") if p not in ("", ".", "..")]

    # webkitRelativePath de una carpeta incluye: Carpeta/archivo.ext
    if len(partes) < 2:
        return None

    return partes[0]


def subir_lote_carpeta_a_drive(
    documentos: list,
    parent_id: str | None = None,
):
    """Sube una carpeta pendiente desde Supabase Storage a Google Drive.

    Los archivos ya no dependen de la RAM de Railway. Cada objeto se descarga
    temporalmente de Supabase Storage justo antes de enviarlo a Drive y se
    libera antes de continuar con el siguiente archivo.
    """

    if not documentos:
        return None

    if not google_drive_configurado():
        print("[DRIVE LOTE CONFIG ERROR] Faltan variables de Google Drive")
        return None

    preparados = []

    for documento in documentos:
        auditoria_id = str(documento.get("id") or "").strip()
        temp_path = str(documento.get("temp_storage_path") or "").strip()

        if not auditoria_id or not temp_path:
            print(
                "[DRIVE LOTE ERROR] Archivo temporal no disponible:",
                auditoria_id,
            )
            return None

        ruta = str(
            documento.get("ruta_relativa")
            or documento.get("nombre_archivo")
            or "archivo"
        ).replace("\\", "/").strip("/")

        preparados.append({
            "documento": documento,
            "ruta": ruta,
            "temp_path": temp_path,
            "temp_bucket": (
                documento.get("temp_storage_bucket")
                or SUPABASE_TEMP_BUCKET
            ),
        })

    nombre_carpeta = obtener_carpeta_desde_ruta(preparados[0]["ruta"])
    if not nombre_carpeta:
        lote_id = str(documentos[0].get("lote_id") or "lote")
        nombre_carpeta = f"LOTE_{lote_id[:8]}"

    destino_parent_id = (
        str(parent_id or GOOGLE_FOLDER_ID).strip()
        or GOOGLE_FOLDER_ID
    )

    service = None
    carpeta_raiz_id = None

    try:
        service = obtener_servicio_google_drive()
        lote_id_actual = str(documentos[0].get("lote_id") or "")
        carpeta_raiz_id = crear_carpeta_google_drive(
            service,
            nombre_carpeta,
            destino_parent_id,
            app_properties={
                "datavault_lote_id": lote_id_actual,
                "datavault_tipo": "carpeta_lote",
            },
        )

        carpetas_cache = {"": carpeta_raiz_id}
        archivos_drive = []

        for item in preparados:
            ruta = item["ruta"]
            partes = [p for p in ruta.split("/") if p]

            if partes and partes[0] == nombre_carpeta:
                partes = partes[1:]

            if not partes:
                partes = [item["documento"].get("nombre_archivo") or "archivo"]

            nombre_drive = partes[-1]
            subdirectorios = partes[:-1]
            parent_actual = carpeta_raiz_id
            ruta_cache = ""

            for directorio in subdirectorios:
                ruta_cache = f"{ruta_cache}/{directorio}" if ruta_cache else directorio

                if ruta_cache not in carpetas_cache:
                    carpetas_cache[ruta_cache] = crear_carpeta_google_drive(
                        service,
                        directorio,
                        parent_actual,
                    )

                parent_actual = carpetas_cache[ruta_cache]

            contenido = leer_archivo_temporal(
                item["temp_path"],
                bucket=item["temp_bucket"],
            )

            archivo_drive_id = subir_archivo_google_drive_en_carpeta(
                service,
                nombre_drive,
                contenido,
                parent_actual,
                app_properties={
                    "datavault_auditoria_id": str(item["documento"].get("id") or ""),
                    "datavault_lote_id": str(item["documento"].get("lote_id") or ""),
                },
            )
            archivos_drive.append(archivo_drive_id)
            del contenido

        print(
            f"[DRIVE LOTE SUCCESS] carpeta={nombre_carpeta} "
            f"archivos={len(archivos_drive)} id={carpeta_raiz_id}"
        )

        return {
            "folder_id": carpeta_raiz_id,
            "folder_name": nombre_carpeta,
            "file_ids": archivos_drive,
        }

    except Exception as error:
        print(f"[DRIVE LOTE EXCEPTION] {error}")

        if service and carpeta_raiz_id:
            try:
                (
                    service
                    .files()
                    .delete(
                        fileId=carpeta_raiz_id,
                        supportsAllDrives=True,
                    )
                    .execute()
                )
                print(f"[DRIVE LOTE ROLLBACK] carpeta eliminada {carpeta_raiz_id}")
            except Exception as rollback_error:
                print(f"[DRIVE LOTE ROLLBACK ERROR] {rollback_error}")

        return None


# ============================================================
# CARGAR IDS DE TELEGRAM
# ============================================================

def cargar_ids_autorizados():

    ids = []

    for valor in AUTHORIZED_CHAT_IDS_RAW.split(","):

        valor = valor.strip()

        if not valor:
            continue

        try:

            ids.append(
                int(valor)
            )

        except ValueError:

            print(
                f"[CONFIG] Telegram ID inválido "
                f"ignorado: {valor}"
            )

    return ids


AUTHORIZED_CHAT_IDS = cargar_ids_autorizados()


# ============================================================
# SUPABASE
# ============================================================

if not SUPABASE_URL:

    raise RuntimeError(
        "SUPABASE_URL no está configurado."
    )


if not SUPABASE_KEY:

    raise RuntimeError(
        "SUPABASE_KEY no está configurado en Railway."
    )


supabase: Client = create_client(
    SUPABASE_URL,
    SUPABASE_KEY
)


# Cliente exclusivo del backend para operaciones administrativas.
# Si la service role está disponible, permite gestionar solicitudes_operacion
# aunque la tabla tenga RLS activado. Nunca se expone esta clave al navegador.
if not SUPABASE_SERVICE_ROLE_KEY:
    raise RuntimeError(
        "SUPABASE_SERVICE_ROLE_KEY no está configurado en Railway. "
        "Es obligatorio para las operaciones internas del backend."
    )

supabase_admin: Client = create_client(
    SUPABASE_URL,
    SUPABASE_SERVICE_ROLE_KEY
)


# ============================================================
# MOTOR DLP DEL BACKEND (PUNTOS 13-17)
# ============================================================

# Estas reglas son la base segura del backend. Las configuraciones guardadas
# en Supabase pueden modificar valores configurables, pero las reglas marcadas
# como obligatorias nunca se desactivan aunque la fila de la BD se altere.
DLP_POLITICAS_DEFAULT = {
    "APROBACION_SUBIDA": {
        "nombre": "Aprobación obligatoria de subidas",
        "descripcion": "Toda subida debe pasar por custodia antes de llegar a Google Drive.",
        "tipo": "BOOLEAN",
        "valor": True,
        "activa": True,
        "obligatoria": True,
    },
    "APROBACION_MOVIMIENTO": {
        "nombre": "Aprobación obligatoria de movimientos",
        "descripcion": "Mover archivos o carpetas requiere autorización del custodio.",
        "tipo": "BOOLEAN",
        "valor": True,
        "activa": True,
        "obligatoria": True,
    },
    "APROBACION_ELIMINACION": {
        "nombre": "Aprobación obligatoria de eliminaciones",
        "descripcion": "Eliminar archivos o carpetas requiere autorización del custodio.",
        "tipo": "BOOLEAN",
        "valor": True,
        "activa": True,
        "obligatoria": True,
    },
    "VALIDAR_SHA256": {
        "nombre": "Integridad SHA-256",
        "descripcion": "Calcula y registra SHA-256 para cada archivo que ingresa al sistema.",
        "tipo": "BOOLEAN",
        "valor": True,
        "activa": True,
        "obligatoria": True,
    },
    "MAX_FILE_SIZE_MB": {
        "nombre": "Tamaño máximo por archivo",
        "descripcion": "Límite máximo permitido para cada archivo recibido por el backend.",
        "tipo": "NUMBER",
        "valor": MAX_FILE_SIZE_MB,
        "activa": True,
        "obligatoria": False,
    },
    "EXTENSIONES_BLOQUEADAS": {
        "nombre": "Extensiones bloqueadas",
        "descripcion": "Bloquea extensiones consideradas de riesgo antes de guardarlas.",
        "tipo": "LIST",
        "valor": [
            "exe", "bat", "cmd", "com", "scr", "msi",
            "ps1", "vbs", "vbe", "wsf", "wsh"
        ],
        "activa": True,
        "obligatoria": False,
    },
    "ALERTA_TELEGRAM": {
        "nombre": "Alertas por Telegram",
        "descripcion": "Envía avisos de solicitudes críticas a los custodios autorizados.",
        "tipo": "BOOLEAN",
        "valor": True,
        "activa": True,
        "obligatoria": False,
    },
    "ALERTA_LOGIN_SUBORDINADO": {
        "nombre": "Alerta de inicio de sesión",
        "descripcion": "Permite alertar por Telegram cuando inicia sesión un subordinado.",
        "tipo": "BOOLEAN",
        "valor": True,
        "activa": True,
        "obligatoria": False,
    },
}

DLP_CODIGOS_OBLIGATORIOS = {
    codigo
    for codigo, config in DLP_POLITICAS_DEFAULT.items()
    if config.get("obligatoria")
}


def _dlp_codigo(valor) -> str:
    return str(valor or "").strip().upper()


def _dlp_extension(nombre_archivo: str) -> str:
    extension = os.path.splitext(str(nombre_archivo or ""))[1].lower().strip()
    return extension.lstrip(".")


def _dlp_lista_extensiones(valor) -> list:
    if valor is None:
        return []

    if isinstance(valor, str):
        # Acepta JSON, CSV o texto separado por saltos de línea.
        texto = valor.strip()
        if not texto:
            return []
        try:
            parsed = json.loads(texto)
            if isinstance(parsed, list):
                valor = parsed
            else:
                valor = texto.replace("\n", ",").split(",")
        except Exception:
            valor = texto.replace("\n", ",").split(",")

    if not isinstance(valor, (list, tuple, set)):
        valor = [valor]

    salida = []
    vistos = set()
    for item in valor:
        ext = str(item or "").strip().lower().lstrip(".")
        if not ext or ext in vistos:
            continue
        vistos.add(ext)
        salida.append(ext)

    return salida


def _dlp_valor_desde_fila(fila: dict, fallback):
    """Lee el valor nuevo (valor_json) y tolera columnas antiguas si existen."""
    if not isinstance(fila, dict):
        return fallback

    if fila.get("valor_json") is not None:
        return fila.get("valor_json")

    for columna in (
        "valor",
        "valor_booleano",
        "valor_numero",
        "valor_texto",
    ):
        if fila.get(columna) is not None:
            valor = fila.get(columna)

            if columna == "valor_texto":
                texto = str(valor).strip()
                try:
                    return json.loads(texto)
                except Exception:
                    return texto

            return valor

    return fallback


def obtener_politicas_dlp() -> dict:
    """Devuelve políticas efectivas; nunca deja caer las reglas obligatorias."""
    politicas = {
        codigo: {
            "codigo": codigo,
            **config,
        }
        for codigo, config in DLP_POLITICAS_DEFAULT.items()
    }

    try:
        respuesta = (
            supabase_admin
            .table("dlp_politicas")
            .select("*")
            .execute()
        )
        filas = respuesta.data or []
    except Exception as error:
        # Si Supabase tiene una incidencia, la seguridad esencial sigue activa
        # mediante los valores por defecto del backend.
        print("[DLP POLICIES LOAD ERROR]", error)
        return politicas

    for fila in filas:
        codigo = _dlp_codigo(fila.get("codigo"))
        if not codigo:
            continue

        base = politicas.get(codigo, {
            "codigo": codigo,
            "nombre": fila.get("nombre") or codigo,
            "descripcion": fila.get("descripcion") or "",
            "tipo": fila.get("tipo") or "TEXT",
            "valor": None,
            "activa": True,
            "obligatoria": False,
        })

        obligatoria = bool(
            base.get("obligatoria")
            or fila.get("obligatoria")
            or codigo in DLP_CODIGOS_OBLIGATORIOS
        )

        activa_db = fila.get("activa")
        activa = True if obligatoria else (
            bool(activa_db)
            if activa_db is not None
            else bool(base.get("activa", True))
        )

        base.update({
            "codigo": codigo,
            "nombre": fila.get("nombre") or base.get("nombre") or codigo,
            "descripcion": (
                fila.get("descripcion")
                if fila.get("descripcion") is not None
                else base.get("descripcion", "")
            ),
            "tipo": fila.get("tipo") or base.get("tipo") or "TEXT",
            "valor": _dlp_valor_desde_fila(
                fila,
                base.get("valor"),
            ),
            "activa": activa,
            "obligatoria": obligatoria,
            "updated_at": fila.get("updated_at"),
            "updated_by": fila.get("updated_by"),
        })

        politicas[codigo] = base

    # Defensa adicional: incluso si una fila obligatoria fue manipulada,
    # el backend la considera activa.
    for codigo in DLP_CODIGOS_OBLIGATORIOS:
        if codigo in politicas:
            politicas[codigo]["activa"] = True
            politicas[codigo]["obligatoria"] = True

    return politicas


def obtener_politica_dlp(codigo: str) -> dict:
    codigo = _dlp_codigo(codigo)
    politicas = obtener_politicas_dlp()

    if codigo not in politicas:
        raise HTTPException(
            status_code=404,
            detail=f"Política DLP no encontrada: {codigo}",
        )

    return politicas[codigo]


def politica_dlp_activa(codigo: str) -> bool:
    codigo = _dlp_codigo(codigo)

    # Las esenciales no dependen de la BD.
    if codigo in DLP_CODIGOS_OBLIGATORIOS:
        return True

    try:
        politica = obtener_politica_dlp(codigo)
        return bool(politica.get("activa"))
    except HTTPException:
        return False


def obtener_max_file_size_dlp_mb() -> int:
    try:
        politica = obtener_politica_dlp("MAX_FILE_SIZE_MB")
        if not politica.get("activa"):
            return MAX_FILE_SIZE_MB

        valor = politica.get("valor")
        numero = int(float(valor))
        if numero < 1:
            raise ValueError("El límite debe ser mayor que cero.")
        # Tope defensivo para evitar configuraciones accidentales extremas.
        return min(numero, 2048)
    except Exception as error:
        print("[DLP MAX FILE SIZE FALLBACK]", error)
        return MAX_FILE_SIZE_MB


def obtener_extensiones_bloqueadas_dlp() -> set:
    try:
        politica = obtener_politica_dlp("EXTENSIONES_BLOQUEADAS")
        if not politica.get("activa"):
            return set()

        return set(
            _dlp_lista_extensiones(
                politica.get("valor")
            )
        )
    except Exception as error:
        print("[DLP EXTENSIONS FALLBACK]", error)
        return set(
            _dlp_lista_extensiones(
                DLP_POLITICAS_DEFAULT["EXTENSIONES_BLOQUEADAS"]["valor"]
            )
        )


def validar_archivo_dlp(
    nombre_archivo: str,
    tamano_bytes: int,
):
    """Aplica DLP real antes de persistir el archivo."""
    nombre_archivo = str(nombre_archivo or "archivo_sin_nombre").strip()
    tamano_bytes = int(tamano_bytes or 0)

    if tamano_bytes <= 0:
        raise HTTPException(
            status_code=400,
            detail="El archivo está vacío.",
        )

    limite_mb = obtener_max_file_size_dlp_mb()
    limite_bytes = limite_mb * 1024 * 1024

    if tamano_bytes > limite_bytes:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Política DLP: el archivo supera el límite de "
                f"{limite_mb} MB."
            ),
        )

    extension = _dlp_extension(nombre_archivo)
    bloqueadas = obtener_extensiones_bloqueadas_dlp()

    if extension and extension in bloqueadas:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Política DLP: la extensión .{extension} "
                "no está permitida."
            ),
        )

    return {
        "permitido": True,
        "limite_mb": limite_mb,
        "extension": extension or None,
    }


def normalizar_valor_politica_dlp(codigo: str, valor):
    codigo = _dlp_codigo(codigo)
    config = DLP_POLITICAS_DEFAULT.get(codigo)

    if not config:
        return valor

    tipo = str(config.get("tipo") or "TEXT").upper()

    if codigo == "MAX_FILE_SIZE_MB":
        try:
            numero = int(float(valor))
        except (TypeError, ValueError):
            raise HTTPException(
                status_code=400,
                detail="MAX_FILE_SIZE_MB debe ser un número.",
            )

        if numero < 1 or numero > 2048:
            raise HTTPException(
                status_code=400,
                detail="MAX_FILE_SIZE_MB debe estar entre 1 y 2048.",
            )

        return numero

    if codigo == "EXTENSIONES_BLOQUEADAS":
        return _dlp_lista_extensiones(valor)

    if tipo == "BOOLEAN":
        if isinstance(valor, bool):
            return valor

        texto = str(valor or "").strip().lower()
        if texto in ("true", "1", "si", "sí", "on"):
            return True
        if texto in ("false", "0", "no", "off"):
            return False

        raise HTTPException(
            status_code=400,
            detail=f"{codigo} requiere un valor booleano.",
        )

    return valor


def registrar_auditoria_cambio_dlp(
    admin: dict,
    codigo: str,
    accion: str,
    anterior,
    nuevo,
):
    payload = {
        "politica_codigo": _dlp_codigo(codigo),
        "accion": str(accion or "ACTUALIZAR").upper(),
        "valor_anterior": anterior,
        "valor_nuevo": nuevo,
        "actor_id": str(admin.get("id") or "") or None,
        "actor_nombre": admin.get("nombre"),
        "actor_correo": admin.get("correo"),
        "fecha": ahora_iso(),
    }

    try:
        (
            supabase_admin
            .table("dlp_auditoria")
            .insert(payload)
            .execute()
        )
    except Exception as error:
        print("[DLP AUDIT ERROR]", error)
        # El fallo queda visible y no se oculta al administrador.
        raise HTTPException(
            status_code=500,
            detail=(
                "La política fue procesada, pero no se pudo registrar "
                "la auditoría DLP. Revisa la tabla dlp_auditoria."
            ),
        )


def guardar_politica_dlp(
    admin: dict,
    codigo: str,
    activa=None,
    valor_marker=False,
    valor=None,
    accion: str = "ACTUALIZAR",
):
    codigo = _dlp_codigo(codigo)

    if codigo not in DLP_POLITICAS_DEFAULT:
        raise HTTPException(
            status_code=404,
            detail=f"Política DLP no reconocida: {codigo}",
        )

    actual = obtener_politica_dlp(codigo)
    obligatoria = bool(
        actual.get("obligatoria")
        or codigo in DLP_CODIGOS_OBLIGATORIOS
    )

    if activa is None:
        nueva_activa = bool(actual.get("activa", True))
    else:
        nueva_activa = bool(activa)

    if obligatoria and not nueva_activa:
        raise HTTPException(
            status_code=409,
            detail=(
                f"La política {codigo} es obligatoria "
                "y no puede desactivarse."
            ),
        )

    nuevo_valor = (
        normalizar_valor_politica_dlp(codigo, valor)
        if valor_marker
        else actual.get("valor")
    )

    config = DLP_POLITICAS_DEFAULT[codigo]
    ahora = ahora_iso()

    fila = {
        "codigo": codigo,
        "nombre": config.get("nombre"),
        "descripcion": config.get("descripcion"),
        "tipo": config.get("tipo"),
        "valor_json": nuevo_valor,
        "activa": True if obligatoria else nueva_activa,
        "obligatoria": obligatoria,
        "updated_by": str(admin.get("id") or "") or None,
        "updated_at": ahora,
    }

    try:
        existente = (
            supabase_admin
            .table("dlp_politicas")
            .select("id")
            .eq("codigo", codigo)
            .limit(1)
            .execute()
        )

        if existente.data:
            respuesta = (
                supabase_admin
                .table("dlp_politicas")
                .update(fila)
                .eq("codigo", codigo)
                .execute()
            )
        else:
            fila["created_at"] = ahora
            respuesta = (
                supabase_admin
                .table("dlp_politicas")
                .insert(fila)
                .execute()
            )
    except Exception as error:
        print("[DLP POLICY SAVE ERROR]", error)
        raise HTTPException(
            status_code=500,
            detail=(
                "No se pudo guardar la política DLP en Supabase. "
                "Ejecuta primero la migración de la Entrega B."
            ),
        )

    nuevo = {
        **actual,
        "codigo": codigo,
        "valor": nuevo_valor,
        "activa": True if obligatoria else nueva_activa,
        "obligatoria": obligatoria,
        "updated_at": ahora,
        "updated_by": str(admin.get("id") or "") or None,
    }

    registrar_auditoria_cambio_dlp(
        admin,
        codigo,
        accion,
        {
            "activa": actual.get("activa"),
            "valor": actual.get("valor"),
        },
        {
            "activa": nuevo.get("activa"),
            "valor": nuevo.get("valor"),
        },
    )

    registrar_evento_auditoria(
        evento="POLITICA_DLP_MODIFICADA",
        categoria="DLP",
        actor=admin,
        accion=accion,
        objeto_tipo="POLITICA_DLP",
        objeto_id=codigo,
        estado=ESTADO_APROBADO,
        detalle={
            "anterior": {
                "activa": actual.get("activa"),
                "valor": actual.get("valor"),
            },
            "nuevo": {
                "activa": nuevo.get("activa"),
                "valor": nuevo.get("valor"),
            },
        },
    )

    return nuevo


# ============================================================
# ALMACENAMIENTO TEMPORAL PERSISTENTE (SUPABASE STORAGE)
# ============================================================

def _storage_temporal():
    if not SUPABASE_SERVICE_ROLE_KEY:
        raise RuntimeError(
            "SUPABASE_SERVICE_ROLE_KEY es obligatorio para el almacenamiento temporal."
        )
    return supabase_admin.storage.from_(SUPABASE_TEMP_BUCKET)


def construir_ruta_archivo_temporal(
    auditoria_id: str,
    solicitante_id: str,
    nombre_archivo: str,
    lote_id: str | None = None,
) -> str:
    extension = os.path.splitext(str(nombre_archivo or ""))[1].lower()
    if not extension or len(extension) > 12 or not all(
        c.isalnum() or c == "." for c in extension
    ):
        extension = ".bin"

    grupo = str(lote_id or "individual").strip() or "individual"
    return (
        f"pendientes/{str(solicitante_id).strip()}/"
        f"{grupo}/{str(auditoria_id).strip()}{extension}"
    )


def guardar_archivo_temporal(
    auditoria_id: str,
    solicitante_id: str,
    nombre_archivo: str,
    contenido: bytes,
    lote_id: str | None = None,
) -> str:
    if not contenido:
        raise RuntimeError("No se puede guardar un archivo temporal vacío.")

    path = construir_ruta_archivo_temporal(
        auditoria_id,
        solicitante_id,
        nombre_archivo,
        lote_id=lote_id,
    )

    content_type = (
        mimetypes.guess_type(nombre_archivo)[0]
        or "application/octet-stream"
    )

    try:
        _storage_temporal().upload(
            path=path,
            file=contenido,
            file_options={
                "content-type": content_type,
                "upsert": "false",
            },
        )
    except Exception as error:
        raise RuntimeError(
            f"No se pudo guardar el archivo temporal en Supabase Storage: {error}"
        ) from error

    return path


def leer_archivo_temporal(
    path: str,
    bucket: str | None = None,
) -> bytes:
    path = str(path or "").strip()
    if not path:
        raise RuntimeError("El registro no tiene una ruta de almacenamiento temporal.")

    try:
        storage = supabase_admin.storage.from_(
            str(bucket or SUPABASE_TEMP_BUCKET).strip()
            or SUPABASE_TEMP_BUCKET
        )
        contenido = storage.download(path)
    except Exception as error:
        raise RuntimeError(
            f"No se pudo recuperar el archivo temporal de Supabase Storage: {error}"
        ) from error

    if not contenido:
        raise RuntimeError("Supabase Storage devolvió un archivo temporal vacío.")

    return bytes(contenido)


def eliminar_archivo_temporal(
    path: str,
    bucket: str | None = None,
) -> bool:
    path = str(path or "").strip()
    if not path:
        return True

    try:
        storage = supabase_admin.storage.from_(
            str(bucket or SUPABASE_TEMP_BUCKET).strip()
            or SUPABASE_TEMP_BUCKET
        )
        storage.remove([path])
        return True
    except Exception as error:
        print("[TEMP STORAGE DELETE ERROR]", path, error)
        return False


def limpiar_temporal_de_registro(registro: dict) -> bool:
    auditoria_id = str(registro.get("id") or "").strip()
    path = str(registro.get("temp_storage_path") or "").strip()
    bucket = str(
        registro.get("temp_storage_bucket")
        or SUPABASE_TEMP_BUCKET
    ).strip()

    if not path:
        return True

    eliminado = eliminar_archivo_temporal(path, bucket=bucket)

    if eliminado and auditoria_id:
        try:
            (
                supabase_admin
                .table("auditoria_custodia")
                .update({
                    "temp_storage_deleted_at": ahora_iso(),
                    "temp_storage_path": None,
                })
                .eq("id", auditoria_id)
                .execute()
            )
        except Exception as error:
            print("[TEMP STORAGE DB CLEANUP ERROR]", auditoria_id, error)

    return eliminado



# ============================================================
# AUDITORÍA AMPLIADA + LOGIN (PUNTOS 22-24)
# ============================================================

def _ip_request(request: Request | None) -> str | None:
    if request is None:
        return None

    forwarded = str(
        request.headers.get("x-forwarded-for")
        or ""
    ).strip()

    if forwarded:
        return forwarded.split(",")[0].strip() or None

    try:
        return request.client.host if request.client else None
    except Exception:
        return None


def _jwt_payload_sin_verificar(token: str) -> dict:
    """Solo extrae claims de un JWT que YA fue validado contra Supabase."""
    try:
        partes = str(token or "").split(".")
        if len(partes) < 2:
            return {}

        payload = partes[1]
        payload += "=" * (-len(payload) % 4)

        return json.loads(
            base64.urlsafe_b64decode(
                payload.encode("utf-8")
            ).decode("utf-8")
        )
    except Exception:
        return {}


def registrar_evento_auditoria(
    evento: str,
    categoria: str = "SISTEMA",
    actor: dict | None = None,
    accion: str | None = None,
    objeto_tipo: str | None = None,
    objeto_id: str | None = None,
    estado: str | None = None,
    detalle=None,
    request: Request | None = None,
    session_hash: str | None = None,
) -> bool:
    actor = actor or {}

    payload = {
        "evento": str(evento or "EVENTO").upper(),
        "categoria": str(categoria or "SISTEMA").upper(),
        "actor_id": str(actor.get("id") or "") or None,
        "actor_nombre": actor.get("nombre"),
        "actor_correo": actor.get("correo"),
        "actor_rol": actor.get("rol"),
        "accion": str(accion or evento or "EVENTO").upper(),
        "objeto_tipo": str(objeto_tipo or "").upper() or None,
        "objeto_id": str(objeto_id or "") or None,
        "estado": normalizar_estado(estado, "") if estado else None,
        "detalle": detalle if detalle is not None else {},
        "ip": _ip_request(request),
        "user_agent": (
            str(request.headers.get("user-agent") or "")[:500]
            if request is not None
            else None
        ),
        "session_hash": session_hash,
        "created_at": ahora_iso(),
    }

    try:
        (
            supabase_admin
            .table("auditoria_eventos_v2")
            .insert(payload)
            .execute()
        )
        return True

    except Exception as error:
        texto = str(error).lower()

        # El login se deduplica por session_hash. Un duplicado es normal.
        if session_hash and (
            "duplicate" in texto
            or "unique" in texto
            or "23505" in texto
        ):
            return False

        print("[AUDITORIA EVENTO ERROR]", evento, error)
        return False


def notificar_login_subordinado_telegram(
    usuario: dict,
    request: Request | None = None,
):
    if not politica_dlp_activa("ALERTA_LOGIN_SUBORDINADO"):
        return

    if not politica_dlp_activa("ALERTA_TELEGRAM"):
        return

    if not TELEGRAM_TOKEN or not AUTHORIZED_CHAT_IDS:
        return

    texto = (
        "🔐 INICIO DE SESIÓN DATAVAULT\n\n"
        f"👤 {usuario.get('nombre') or 'Usuario'}\n"
        f"📧 {usuario.get('correo') or 'Sin correo'}\n"
        f"🧩 Rol: {usuario.get('rol') or 'subordinado'}\n"
        f"🕒 {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    )

    for chat_id in AUTHORIZED_CHAT_IDS:
        try:
            telegram_request(
                "sendMessage",
                {
                    "chat_id": chat_id,
                    "text": texto,
                },
            )
        except Exception as error:
            print("[TELEGRAM LOGIN ALERT ERROR]", error)


def registrar_login_usuario_si_corresponde(
    usuario: dict,
    access_token: str,
    request: Request,
):
    claims = _jwt_payload_sin_verificar(access_token)

    # Supabase suele mantener session_id durante los refresh del mismo login.
    # Si no está presente usamos el hash del access token como fallback.
    referencia_sesion = str(
        claims.get("session_id")
        or claims.get("sid")
        or ""
    ).strip()

    if referencia_sesion:
        session_hash = hashlib.sha256(
            referencia_sesion.encode("utf-8")
        ).hexdigest()
    else:
        session_hash = hashlib.sha256(
            str(access_token).encode("utf-8")
        ).hexdigest()

    insertado = registrar_evento_auditoria(
        evento="LOGIN",
        categoria="AUTENTICACION",
        actor=usuario,
        accion="INICIAR_SESION",
        objeto_tipo="SESION",
        objeto_id=referencia_sesion or None,
        estado="APROBADO",
        detalle={
            "proveedor": "SUPABASE_AUTH",
            "rol": usuario.get("rol"),
        },
        request=request,
        session_hash=session_hash,
    )

    if (
        insertado
        and str(usuario.get("rol") or "").lower() == "subordinado"
    ):
        notificar_login_subordinado_telegram(
            usuario,
            request=request,
        )


# ============================================================
# IDENTIDAD REAL DEL USUARIO DESDE SUPABASE AUTH
# ============================================================

def obtener_usuario_supabase_desde_request(request: Request) -> dict:

    authorization = request.headers.get("Authorization", "").strip()

    if not authorization.lower().startswith("bearer "):
        raise HTTPException(
            status_code=401,
            detail="Sesión de Supabase no enviada."
        )

    access_token = authorization.split(" ", 1)[1].strip()

    if not access_token:
        raise HTTPException(
            status_code=401,
            detail="Token de Supabase vacío."
        )

    try:
        respuesta = HTTP_CLIENT.get(
            f"{SUPABASE_URL}/auth/v1/user",
            headers={
                "apikey": SUPABASE_KEY,
                "Authorization": f"Bearer {access_token}"
            },
            timeout=15
        )
    except httpx.HTTPError as error:
        print("[SUPABASE AUTH ERROR]", error)
        raise HTTPException(
            status_code=503,
            detail="No se pudo validar la sesión con Supabase."
        )

    if respuesta.status_code != 200:
        print(
            "[SUPABASE AUTH REJECTED]",
            respuesta.status_code,
            respuesta.text[:500]
        )
        raise HTTPException(
            status_code=401,
            detail="Sesión de Supabase inválida o expirada."
        )

    usuario = respuesta.json()

    solicitante_id = str(usuario.get("id") or "").strip()
    solicitante_correo = str(usuario.get("email") or "").strip()
    metadata = usuario.get("user_metadata") or {}
    app_metadata = usuario.get("app_metadata") or {}

    solicitante_nombre = str(
        metadata.get("full_name")
        or metadata.get("name")
        or solicitante_correo
        or "Usuario desconocido"
    ).strip()

    rol = str(
        app_metadata.get("rol")
        or app_metadata.get("role")
        or metadata.get("rol")
        or metadata.get("role")
        or "subordinado"
    ).strip().lower()

    if rol in ("admin", "administrador"):
        rol = "jefe"

    if not solicitante_id:
        raise HTTPException(
            status_code=401,
            detail="Supabase no devolvió el UID del usuario."
        )

    usuario_normalizado = {
        "id": solicitante_id,
        "nombre": solicitante_nombre,
        "correo": solicitante_correo,
        "rol": rol,
    }

    try:
        registrar_login_usuario_si_corresponde(
            usuario_normalizado,
            access_token,
            request,
        )
    except Exception as error:
        # Un problema de auditoría/Telegram nunca invalida una sesión válida.
        print("[LOGIN AUDIT BACKGROUND ERROR]", error)

    return usuario_normalizado


def obtener_admin_desde_request(request: Request) -> dict:
    usuario = obtener_usuario_supabase_desde_request(request)
    if usuario.get("rol") != "jefe":
        raise HTTPException(
            status_code=403,
            detail="Solo el administrador puede realizar esta operación.",
        )
    return usuario


# ============================================================
# OPERACIONES GOOGLE DRIVE / SINCRONIZACIÓN
# ============================================================

def ahora_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def obtener_archivo_drive(service, file_id: str):
    return (
        service
        .files()
        .get(
            fileId=file_id,
            fields="id,name,mimeType,parents,trashed",
            supportsAllDrives=True,
        )
        .execute()
    )


def _cache_drive_get(cache_key: str):
    try:
        respuesta = (
            supabase_admin
            .table("drive_cache_navegacion")
            .select("payload,expires_at")
            .eq("cache_key", str(cache_key))
            .gt("expires_at", ahora_iso())
            .limit(1)
            .execute()
        )
        if respuesta.data:
            return respuesta.data[0].get("payload")
    except Exception as error:
        print("[DRIVE CACHE GET]", error)
    return None


def _cache_drive_set(cache_key: str, payload):
    ahora = datetime.now(timezone.utc)
    expira = ahora + timedelta(seconds=DRIVE_CACHE_TTL_SECONDS)

    fila = {
        "cache_key": str(cache_key),
        "payload": payload,
        "created_at": ahora.isoformat(),
        "updated_at": ahora.isoformat(),
        "expires_at": expira.isoformat(),
    }

    try:
        existente = (
            supabase_admin
            .table("drive_cache_navegacion")
            .select("id")
            .eq("cache_key", str(cache_key))
            .limit(1)
            .execute()
        )

        if existente.data:
            (
                supabase_admin
                .table("drive_cache_navegacion")
                .update(fila)
                .eq("id", existente.data[0]["id"])
                .execute()
            )
        else:
            (
                supabase_admin
                .table("drive_cache_navegacion")
                .insert(fila)
                .execute()
            )
    except Exception as error:
        print("[DRIVE CACHE SET]", error)


def invalidar_cache_drive():
    """Invalida navegación después de subir, mover o eliminar."""
    try:
        (
            supabase_admin
            .table("drive_cache_navegacion")
            .delete()
            .like("cache_key", "%")
            .execute()
        )
    except Exception as error:
        print("[DRIVE CACHE INVALIDATE]", error)


def obtener_raiz_drive_autorizada(service=None) -> dict:
    cache_key = f"root:{GOOGLE_DRIVE_ROOT_ID}"
    cached = _cache_drive_get(cache_key)
    if cached:
        return cached

    service = service or obtener_servicio_google_drive()

    meta = (
        service
        .files()
        .get(
            fileId=GOOGLE_DRIVE_ROOT_ID,
            fields="id,name,mimeType,parents,trashed",
            supportsAllDrives=True,
        )
        .execute()
    )

    raiz = {
        "id": str(meta.get("id") or GOOGLE_DRIVE_ROOT_ID),
        "name": "Mi unidad" if GOOGLE_DRIVE_ROOT_ID == "root" else (
            meta.get("name") or "Raíz autorizada"
        ),
        "mimeType": meta.get("mimeType"),
        "parents": meta.get("parents") or [],
        "path": "Mi unidad" if GOOGLE_DRIVE_ROOT_ID == "root" else (
            meta.get("name") or "Raíz autorizada"
        ),
    }

    _cache_drive_set(cache_key, raiz)
    return raiz


def _normalizar_drive_id(folder_id: str, service=None) -> str:
    valor = str(folder_id or "").strip()

    if not valor or valor == "root":
        return obtener_raiz_drive_autorizada(service)["id"]

    if valor == GOOGLE_DRIVE_ROOT_ID:
        return obtener_raiz_drive_autorizada(service)["id"]

    return valor


def resolver_ruta_drive(folder_id: str) -> dict:
    """Valida pertenencia a la raíz autorizada sin recorrer todo Drive."""

    service = obtener_servicio_google_drive()
    raiz = obtener_raiz_drive_autorizada(service)
    raiz_id = str(raiz["id"])

    objetivo_id = _normalizar_drive_id(folder_id, service)

    cache_key = f"path:{raiz_id}:{objetivo_id}"
    cached = _cache_drive_get(cache_key)
    if cached:
        return cached

    actual_id = objetivo_id
    visitados = set()
    cadena = []
    autorizado = False

    for _ in range(80):
        if actual_id in visitados:
            break
        visitados.add(actual_id)

        meta = (
            service
            .files()
            .get(
                fileId=actual_id,
                fields="id,name,mimeType,parents,trashed",
                supportsAllDrives=True,
            )
            .execute()
        )

        if meta.get("trashed"):
            raise error_api(
                409,
                "DRIVE_FOLDER_TRASHED",
                "La carpeta seleccionada está en la papelera de Google Drive.",
            )

        cadena.append(meta)

        if str(meta.get("id")) == raiz_id:
            autorizado = True
            break

        padres = meta.get("parents") or []
        if not padres:
            break

        actual_id = str(padres[0])

    if not autorizado:
        raise error_api(
            403,
            "DRIVE_OUTSIDE_SCOPE",
            "La carpeta no pertenece a la raíz de Drive autorizada para DataVault.",
        )

    cadena.reverse()

    nombres = []
    breadcrumb = []
    for indice, meta in enumerate(cadena):
        nombre = (
            raiz.get("name")
            if indice == 0
            else str(meta.get("name") or "Carpeta")
        )
        nombres.append(nombre)
        breadcrumb.append({
            "id": str(meta.get("id")),
            "name": nombre,
        })

    objetivo = cadena[-1]

    if (
        objetivo.get("mimeType")
        != "application/vnd.google-apps.folder"
    ):
        raise error_api(
            400,
            "DRIVE_DESTINATION_NOT_FOLDER",
            "El destino seleccionado no es una carpeta de Google Drive.",
        )

    respuesta = {
        "id": str(objetivo.get("id")),
        "name": (
            raiz.get("name")
            if str(objetivo.get("id")) == raiz_id
            else str(objetivo.get("name") or "Carpeta")
        ),
        "parent_id": (
            str((objetivo.get("parents") or [None])[0])
            if objetivo.get("parents")
            else None
        ),
        "path": " / ".join(nombres),
        "depth": max(0, len(cadena) - 1),
        "ancestors": [
            str(x.get("id"))
            for x in cadena[:-1]
        ],
        "breadcrumb": breadcrumb,
        "root": str(objetivo.get("id")) == raiz_id,
    }

    _cache_drive_set(cache_key, respuesta)
    return respuesta


def listar_hijos_drive(
    folder_id: str,
    include_files: bool = True,
) -> dict:
    """Lista solo los hijos inmediatos. No hace búsqueda recursiva."""

    service = obtener_servicio_google_drive()
    actual = resolver_ruta_drive(folder_id)
    actual_id = str(actual["id"])

    cache_key = (
        f"children:{actual_id}:"
        f"{'all' if include_files else 'folders'}"
    )
    cached = _cache_drive_get(cache_key)
    if cached:
        return cached

    consulta = f"'{actual_id}' in parents and trashed=false"

    page_token = None
    elementos = []

    while True:
        respuesta = (
            service
            .files()
            .list(
                q=consulta,
                fields=(
                    "nextPageToken,"
                    "files(id,name,mimeType,parents,size,modifiedTime)"
                ),
                orderBy="folder,name",
                pageSize=500,
                pageToken=page_token,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            )
            .execute()
        )

        elementos.extend(respuesta.get("files") or [])

        page_token = respuesta.get("nextPageToken")
        if not page_token:
            break

    folders = []
    files = []

    for item in elementos:
        es_carpeta = (
            item.get("mimeType")
            == "application/vnd.google-apps.folder"
        )

        normalizado = {
            "id": str(item.get("id") or ""),
            "name": str(item.get("name") or ""),
            "mimeType": item.get("mimeType"),
            "parent_id": actual_id,
            "modifiedTime": item.get("modifiedTime"),
        }

        if es_carpeta:
            normalizado["type"] = "folder"
            folders.append(normalizado)
        elif include_files:
            normalizado["type"] = "file"
            try:
                normalizado["size"] = int(item.get("size") or 0)
            except Exception:
                normalizado["size"] = 0
            files.append(normalizado)

    salida = {
        "current": actual,
        "breadcrumb": actual.get("breadcrumb") or [],
        "folders": folders,
        "files": files,
    }

    _cache_drive_set(cache_key, salida)
    return salida


def listar_carpetas_drive_raiz():
    raiz = obtener_raiz_drive_autorizada()
    salida = listar_hijos_drive(
        raiz["id"],
        include_files=False,
    )
    return salida.get("folders") or []


def listar_carpetas_drive_recursivas():
    """Compatibilidad con el selector antiguo.

    La navegación nueva usa /drive/browse (hijos inmediatos). Esta función solo
    mantiene el frontend existente y guarda el resultado plano en caché.
    """

    raiz = obtener_raiz_drive_autorizada()
    cache_key = f"flat-folders:{raiz['id']}"

    cached = _cache_drive_get(cache_key)
    if cached:
        return cached

    pendientes = [{
        "id": raiz["id"],
        "path": raiz["path"],
        "depth": 0,
        "ancestors": [],
    }]

    visitados = {str(raiz["id"])}
    carpetas = []

    while pendientes:
        padre = pendientes.pop(0)

        hijos = listar_hijos_drive(
            padre["id"],
            include_files=False,
        ).get("folders") or []

        for carpeta in hijos:
            carpeta_id = str(carpeta.get("id") or "")
            if not carpeta_id or carpeta_id in visitados:
                continue

            visitados.add(carpeta_id)

            nombre = str(carpeta.get("name") or "Carpeta")
            path = f"{padre['path']} / {nombre}"

            item = {
                "id": carpeta_id,
                "name": nombre,
                "parent_id": padre["id"],
                "path": path,
                "depth": int(padre["depth"]) + 1,
                "ancestors": [
                    *padre["ancestors"],
                    padre["id"],
                ],
            }

            carpetas.append(item)
            pendientes.append(item)

            if len(carpetas) >= 5000:
                _cache_drive_set(cache_key, carpetas)
                return carpetas

    _cache_drive_set(cache_key, carpetas)
    return carpetas


def validar_destino_drive(destino_id: str):
    """Valida un destino concreto sin buscar recursivamente todo Drive."""
    return resolver_ruta_drive(destino_id)


def mover_objeto_google_drive(
    file_id: str,
    destino_id: str
):

    service = obtener_servicio_google_drive()

    actual = obtener_archivo_drive(
        service,
        file_id
    )

    if actual.get("trashed"):
        raise RuntimeError(
            "El elemento ya está en la papelera de Google Drive."
        )

    padres = [
        str(x)
        for x in (actual.get("parents") or [])
        if x
    ]

    if destino_id in padres and len(padres) == 1:
        return actual

    remove_parents = ",".join([
        p
        for p in padres
        if p != destino_id
    ])

    kwargs = {
        "fileId": file_id,
        "addParents": destino_id,
        "fields": "id,name,parents,trashed",
        "supportsAllDrives": True,
    }

    if remove_parents:
        kwargs["removeParents"] = remove_parents

    resultado = (
        service
        .files()
        .update(**kwargs)
        .execute()
    )

    invalidar_cache_drive()
    return resultado


def enviar_a_papelera_google_drive(file_id: str):

    service = obtener_servicio_google_drive()

    resultado = (
        service
        .files()
        .update(
            fileId=file_id,
            body={"trashed": True},
            fields="id,name,trashed",
            supportsAllDrives=True,
        )
        .execute()
    )

    invalidar_cache_drive()
    return resultado


def obtener_objeto_operable(
    usuario_id: str,
    objeto_tipo: str,
    auditoria_id=None,
    lote_id=None
):

    objeto_tipo = str(
        objeto_tipo or ""
    ).upper().strip()

    if objeto_tipo == "CARPETA":

        if not lote_id:
            raise HTTPException(
                status_code=400,
                detail="Falta lote_id."
            )

        respuesta = (
            supabase_admin
            .table("auditoria_custodia")
            .select("*")
            .eq("lote_id", str(lote_id))
            .eq("solicitante_id", usuario_id)
            .execute()
        )

        filas = respuesta.data or []

        if not filas:
            raise HTTPException(
                status_code=404,
                detail="No se encontró la carpeta del usuario."
            )

        if any(
            str(f.get("estado") or "").upper() != "APROBADO"
            for f in filas
        ):
            raise HTTPException(
                status_code=409,
                detail="Solo se pueden operar carpetas aprobadas."
            )

        if any(
            str(f.get("estado_archivo") or "").upper().startswith("ELIMINADO")
            for f in filas
        ):
            raise HTTPException(
                status_code=409,
                detail="La carpeta ya fue eliminada."
            )

        drive_id = str(
            filas[0].get("drive_folder_id") or ""
        ).strip()

        if not drive_id:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Esta carpeta es anterior al registro de IDs de Drive. "
                    "Carga una carpeta nueva para operar sobre ella."
                ),
            )

        nombre = (
            obtener_carpeta_desde_ruta(
                filas[0].get("ruta_relativa")
            )
            or "Carpeta"
        )

        return {
            "tipo": "CARPETA",
            "drive_id": drive_id,
            "nombre": nombre,
            "filas": filas,
            "lote_id": str(lote_id),
            "auditoria_id": None,
            "ubicacion": (
                filas[0].get("ubicacion_drive")
                or "DRIVE PROYECTO"
            ),
            "drive_parent_id": str(
                filas[0].get("drive_parent_id")
                or GOOGLE_FOLDER_ID
            ),
        }

    if objeto_tipo == "ARCHIVO":

        if not auditoria_id:
            raise HTTPException(
                status_code=400,
                detail="Falta auditoria_id."
            )

        respuesta = (
            supabase_admin
            .table("auditoria_custodia")
            .select("*")
            .eq("id", str(auditoria_id))
            .eq("solicitante_id", usuario_id)
            .limit(1)
            .execute()
        )

        if not respuesta.data:
            raise HTTPException(
                status_code=404,
                detail="No se encontró el archivo del usuario."
            )

        fila = respuesta.data[0]

        if str(
            fila.get("estado") or ""
        ).upper() != "APROBADO":
            raise HTTPException(
                status_code=409,
                detail="Solo se pueden operar archivos aprobados."
            )

        if str(
            fila.get("estado_archivo") or ""
        ).upper().startswith("ELIMINADO"):
            raise HTTPException(
                status_code=409,
                detail="El archivo ya fue eliminado."
            )

        drive_id = str(
            fila.get("drive_file_id") or ""
        ).strip()

        if not drive_id:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Este archivo es anterior al registro de IDs de Drive. "
                    "Carga un archivo nuevo para operar sobre él."
                ),
            )

        return {
            "tipo": "ARCHIVO",
            "drive_id": drive_id,
            "nombre": (
                fila.get("nombre_archivo")
                or "Archivo"
            ),
            "filas": [fila],
            "lote_id": None,
            "auditoria_id": str(auditoria_id),
            "ubicacion": (
                fila.get("ubicacion_drive")
                or "DRIVE PROYECTO"
            ),
            "drive_parent_id": str(
                fila.get("drive_parent_id")
                or GOOGLE_FOLDER_ID
            ),
        }

    raise HTTPException(
        status_code=400,
        detail="objeto_tipo inválido."
    )


def actualizar_auditoria_operacion(
    objeto: dict,
    cambios: dict
):

    query = (
        supabase_admin
        .table("auditoria_custodia")
        .update(cambios)
    )

    if objeto["tipo"] == "CARPETA":
        query = query.eq(
            "lote_id",
            objeto["lote_id"]
        )

    else:
        query = query.eq(
            "id",
            objeto["auditoria_id"]
        )

    return query.execute()


def normalizar_referencias_telegram_operacion(solicitud: dict):
    """Devuelve referencias chat_id/message_id guardadas para una operación.

    Supabase JSONB normalmente llega como lista. También aceptamos string JSON
    para conservar compatibilidad si PostgREST lo serializa en alguna versión.
    """
    raw = (solicitud or {}).get("telegram_mensajes") or []

    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except Exception:
            raw = []

    if isinstance(raw, dict):
        raw = [raw]

    if not isinstance(raw, list):
        return []

    referencias = []
    vistos = set()

    for item in raw:
        if not isinstance(item, dict):
            continue

        chat_id = item.get("chat_id")
        message_id = item.get("message_id")

        if chat_id is None or message_id is None:
            continue

        clave = (str(chat_id), str(message_id))
        if clave in vistos:
            continue
        vistos.add(clave)

        referencias.append({
            "chat_id": chat_id,
            "message_id": message_id,
        })

    return referencias


def guardar_referencias_telegram_operacion(solicitud_id: str, referencias: list):
    """Guarda en JSONB los mensajes enviados a los custodios.

    Si la columna todavía no fue creada, la solicitud NO falla: queda un log
    claro para que la operación principal siga funcionando.
    """
    if not referencias:
        return

    try:
        supabase_admin.table("solicitudes_operacion").update({
            "telegram_mensajes": referencias,
        }).eq("id", str(solicitud_id)).execute()
    except Exception as error:
        print(
            "[TELEGRAM REFERENCIAS SAVE ERROR]",
            "Ejecuta ALTER TABLE solicitudes_operacion ADD COLUMN telegram_mensajes jsonb...",
            error,
        )


def construir_texto_operacion_resuelta(solicitud: dict, origen_resolucion: str = ""):
    estado = str((solicitud or {}).get("estado") or "").upper()
    tipo = str((solicitud or {}).get("tipo_operacion") or "OPERACIÓN").upper()
    icono = "✅" if estado == "APROBADO" else "❌" if estado == "RECHAZADO" else "⚠️"

    origen = str(origen_resolucion or (solicitud or {}).get("resuelto_por") or "").strip()
    if origen.startswith("WEB:"):
        procesado_por = "Panel web del administrador"
    elif origen:
        procesado_por = f"Telegram ID: {origen}"
    else:
        procesado_por = "Custodio"

    texto = (
        "🛡️ DataVault DLP - GM Ingenieros\n\n"
        f"{icono} {tipo}: {estado or 'PROCESADA'}\n\n"
        f"👤 Usuario: {(solicitud or {}).get('solicitante_nombre') or 'Usuario'}\n"
        f"📄 Objeto: {(solicitud or {}).get('nombre_objeto') or 'Archivo/Carpeta'}\n"
        f"📌 Tipo: {(solicitud or {}).get('objeto_tipo') or '—'}\n"
        f"📂 Origen: {(solicitud or {}).get('carpeta_origen') or 'DRIVE PROYECTO'}\n"
    )

    if (solicitud or {}).get("carpeta_destino_nombre"):
        texto += f"➡️ Destino: {(solicitud or {}).get('carpeta_destino_nombre')}\n"

    if (solicitud or {}).get("resultado"):
        texto += f"\n📝 {(solicitud or {}).get('resultado')}\n"

    texto += f"\n👤 Procesado por: {procesado_por}"
    return texto


def sincronizar_mensajes_telegram_operacion(
    solicitud: dict,
    origen_resolucion: str = "",
    fallback_chat_id=None,
    fallback_message_id=None,
    error_texto: str | None = None,
):
    """Sincroniza una decisión de MOVER/ELIMINAR con Telegram.

    Las solicitudes nuevas se muestran dentro de la bandeja agrupada por
    usuario. Si la decisión se tomó desde el detalle de Telegram, ese mismo
    mensaje se actualiza. Luego se publica una bandeja agrupada actualizada para
    cada custodio, de forma que Web y Telegram reflejen los mismos pendientes.
    """
    actualizar_alertas_operacion_telegram(
        solicitud,
        error_texto=error_texto,
    )

    referencias = normalizar_referencias_telegram_operacion(solicitud)

    if fallback_chat_id is not None and fallback_message_id is not None:
        clave_fallback = (str(fallback_chat_id), str(fallback_message_id))
        existentes = {
            (str(r.get("chat_id")), str(r.get("message_id")))
            for r in referencias
        }
        if clave_fallback not in existentes:
            referencias.append({
                "chat_id": fallback_chat_id,
                "message_id": fallback_message_id,
            })

    if error_texto:
        solicitud_id = str((solicitud or {}).get("id") or "")
        texto = (
            "🛡️ DataVault DLP - GM Ingenieros\n\n"
            "❌ NO SE PUDO COMPLETAR LA OPERACIÓN\n\n"
            f"Motivo: {str(error_texto)[:800]}\n\n"
            "La solicitud continúa pendiente."
        )
        reply_markup = {
            "inline_keyboard": [
                [
                    {"text": "✅ Aprobar", "callback_data": f"opap:{solicitud_id}"},
                    {"text": "❌ Rechazar", "callback_data": f"opre:{solicitud_id}"},
                ],
                [{"text": "🔙 Volver al usuario", "callback_data": f"usr:{(solicitud or {}).get('solicitante_id') or ''}"}],
            ]
        }
    else:
        texto = construir_texto_operacion_resuelta(solicitud, origen_resolucion)
        uid = str((solicitud or {}).get("solicitante_id") or "")
        reply_markup = {
            "inline_keyboard": [[
                {"text": "🔙 Pendientes del usuario", "callback_data": f"usr:{uid}"}
            ]] if uid else []
        }

    # Actualiza el detalle desde el que se tomó la decisión y cualquier mensaje
    # individual antiguo que haya quedado de versiones previas.
    for referencia in referencias:
        try:
            telegram_request(
                "editMessageText",
                {
                    "chat_id": referencia["chat_id"],
                    "message_id": referencia["message_id"],
                    "text": texto,
                    "reply_markup": reply_markup,
                },
            )
        except Exception as error:
            print("[TELEGRAM SYNC MESSAGE ERROR]", error)

    # Publica la bandeja unificada actualizada para todos los custodios.
    if not error_texto:
        for chat_id in AUTHORIZED_CHAT_IDS:
            try:
                mostrar_menu_usuarios(chat_id)
            except Exception as error:
                print("[TELEGRAM SYNC MENU ERROR]", error)

def notificar_solicitud_operacion_telegram(solicitud: dict):
    """Envía una alerta real y actualiza la bandeja persistente."""
    if not TELEGRAM_TOKEN or not AUTHORIZED_CHAT_IDS:
        return

    if not politica_dlp_activa("ALERTA_TELEGRAM"):
        print("[DLP] ALERTA_TELEGRAM desactivada; no se envía aviso.")
        return

    try:
        enviar_alerta_operacion_telegram(solicitud)
    except Exception as error:
        print("[TELEGRAM OPERACION ALERT ERROR]", error)

    for chat_id in AUTHORIZED_CHAT_IDS:
        try:
            resultado = mostrar_menu_usuarios(chat_id)
            print(
                "[TELEGRAM OPERACION MENU]",
                f"chat={chat_id}",
                f"solicitud={solicitud.get('id')}",
                f"ok={(resultado or {}).get('ok')}",
            )
        except Exception as error:
            print("[TELEGRAM OPERACION MENU ERROR]", error)


# ============================================================
# PAPELERA + HISTORIAL DE ELIMINACIONES (PUNTOS 18-21)
# ============================================================

def registrar_historial_eliminacion(
    solicitud: dict,
    objeto: dict,
    resuelto_por: str,
) -> dict | None:
    """Registra solo eliminaciones realmente aplicadas en Google Drive."""

    solicitud_id = str((solicitud or {}).get("id") or "").strip()
    solicitante_id = str((solicitud or {}).get("solicitante_id") or "").strip()

    if not solicitud_id or not solicitante_id:
        print("[TRASH HISTORY] solicitud incompleta; no se registró historial.")
        return None

    filas = list((objeto or {}).get("filas") or [])
    auditoria_ids = [
        str(fila.get("id"))
        for fila in filas
        if fila.get("id")
    ]

    objeto_tipo = str((objeto or {}).get("tipo") or "").upper() or "ARCHIVO"
    nombre_objeto = (
        (solicitud or {}).get("nombre_objeto")
        or (objeto or {}).get("nombre")
        or "Archivo/Carpeta"
    )
    ubicacion_anterior = (
        (solicitud or {}).get("carpeta_origen")
        or (objeto or {}).get("ubicacion")
        or "DRIVE PROYECTO"
    )
    detalle = {
        "cantidad_registros": len(filas),
        "auditoria_ids": auditoria_ids,
        "resultado_drive": "TRASHED",
    }

    # La tabla elementos_eliminados ya existía con nombres de columnas de una
    # versión anterior. Guardamos ambos juegos de campos para mantener
    # compatibilidad y no perder el historial previo.
    payload = {
        # Columnas nuevas
        "solicitud_operacion_id": solicitud_id,
        "objeto_tipo": objeto_tipo,
        "solicitante_id": solicitante_id,
        "solicitante_nombre": (solicitud or {}).get("solicitante_nombre"),
        "solicitante_correo": (solicitud or {}).get("solicitante_correo"),
        "nombre_objeto": nombre_objeto,
        "drive_parent_id": str((objeto or {}).get("drive_parent_id") or "") or None,
        "ubicacion_anterior": ubicacion_anterior,
        "detalle": detalle,
        "resuelto_por_ref": str(resuelto_por or "") or None,

        # Columnas comunes
        "auditoria_id": (
            str((objeto or {}).get("auditoria_id"))
            if (objeto or {}).get("auditoria_id")
            else None
        ),
        "lote_id": (
            str((objeto or {}).get("lote_id"))
            if (objeto or {}).get("lote_id")
            else None
        ),
        "drive_id": str((objeto or {}).get("drive_id") or "") or None,
        "fecha_eliminacion": ahora_iso(),
        "estado": "ELIMINADO",
        "updated_at": ahora_iso(),

        # Compatibilidad con la estructura antigua ya existente en Supabase
        "usuario_id": solicitante_id,
        "usuario_nombre": (solicitud or {}).get("solicitante_nombre"),
        "usuario_correo": (solicitud or {}).get("solicitante_correo"),
        "solicitud_id": solicitud_id,
        "elemento_tipo": objeto_tipo,
        "nombre": nombre_objeto,
        "ruta_anterior": ubicacion_anterior,
        "metadata": detalle,
    }

    try:
        existente = (
            supabase_admin
            .table("elementos_eliminados")
            .select("id")
            .eq("solicitud_operacion_id", solicitud_id)
            .limit(1)
            .execute()
        )

        if existente.data:
            respuesta = (
                supabase_admin
                .table("elementos_eliminados")
                .update(payload)
                .eq("id", existente.data[0]["id"])
                .execute()
            )
        else:
            payload["created_at"] = ahora_iso()
            respuesta = (
                supabase_admin
                .table("elementos_eliminados")
                .insert(payload)
                .execute()
            )

        return (respuesta.data or [None])[0]

    except Exception as error:
        print("[TRASH HISTORY ERROR]", solicitud_id, error)
        return None


def construir_items_papelera(filas: list) -> list:
    """Agrupa carpetas eliminadas como una sola unidad visual."""

    carpetas = {}
    archivos = []

    for fila in filas or []:
        estado_archivo = str(fila.get("estado_archivo") or "").upper()
        if not estado_archivo.startswith("ELIMINADO"):
            continue

        drive_folder_id = str(fila.get("drive_folder_id") or "").strip()
        lote_id = str(fila.get("lote_id") or "").strip()

        if drive_folder_id and lote_id:
            clave = f"{drive_folder_id}:{lote_id}"

            grupo = carpetas.setdefault(
                clave,
                {
                    "tipo": "CARPETA",
                    "id": lote_id,
                    "lote_id": lote_id,
                    "auditoria_id": None,
                    "drive_id": drive_folder_id,
                    "nombre": (
                        obtener_carpeta_desde_ruta(fila.get("ruta_relativa"))
                        or "Carpeta"
                    ),
                    "solicitante_id": fila.get("solicitante_id"),
                    "solicitante_nombre": fila.get("solicitante_nombre"),
                    "solicitante_correo": fila.get("solicitante_correo"),
                    "ubicacion_anterior": (
                        fila.get("ubicacion_drive")
                        or "DRIVE PROYECTO"
                    ),
                    "fecha_eliminacion": fila.get("fecha_eliminacion"),
                    "estado_archivo": estado_archivo,
                    "cantidad_archivos": 0,
                    "tamano_bytes": 0,
                },
            )

            grupo["cantidad_archivos"] += 1
            grupo["tamano_bytes"] += int(fila.get("tamano_bytes") or 0)

            if str(fila.get("fecha_eliminacion") or "") > str(
                grupo.get("fecha_eliminacion") or ""
            ):
                grupo["fecha_eliminacion"] = fila.get("fecha_eliminacion")

            continue

        archivos.append({
            "tipo": "ARCHIVO",
            "id": str(fila.get("id") or ""),
            "auditoria_id": str(fila.get("id") or ""),
            "lote_id": None,
            "drive_id": str(fila.get("drive_file_id") or "") or None,
            "nombre": fila.get("nombre_archivo") or "Archivo",
            "solicitante_id": fila.get("solicitante_id"),
            "solicitante_nombre": fila.get("solicitante_nombre"),
            "solicitante_correo": fila.get("solicitante_correo"),
            "ubicacion_anterior": (
                fila.get("ubicacion_drive")
                or "DRIVE PROYECTO"
            ),
            "fecha_eliminacion": fila.get("fecha_eliminacion"),
            "estado_archivo": estado_archivo,
            "cantidad_archivos": 1,
            "tamano_bytes": int(fila.get("tamano_bytes") or 0),
        })

    items = [*carpetas.values(), *archivos]
    items.sort(
        key=lambda item: str(item.get("fecha_eliminacion") or ""),
        reverse=True,
    )

    return items


def consultar_papelera_backend(
    usuario: dict,
    solo_usuario: bool = False,
) -> dict:
    """Admin ve todo; subordinado solo sus propios eliminados."""

    es_admin = str(usuario.get("rol") or "").lower() == "jefe"

    query = (
        supabase_admin
        .table("auditoria_custodia")
        .select(
            "id,lote_id,nombre_archivo,ruta_relativa,tamano_bytes,"
            "solicitante_id,solicitante_nombre,solicitante_correo,"
            "estado,estado_archivo,fecha_eliminacion,"
            "drive_file_id,drive_folder_id,drive_parent_id,ubicacion_drive"
        )
        .like("estado_archivo", "ELIMINADO%")
        .order("fecha_eliminacion", desc=True)
    )

    if solo_usuario or not es_admin:
        query = query.eq(
            "solicitante_id",
            str(usuario.get("id") or ""),
        )

    respuesta = query.execute()
    filas = respuesta.data or []
    items = construir_items_papelera(filas)

    return {
        "status": "ok",
        "rol": usuario.get("rol"),
        "solo_usuario": bool(solo_usuario or not es_admin),
        "total_items": len(items),
        "total_registros": len(filas),
        "items": items,
    }


def consultar_historial_eliminaciones(
    usuario: dict,
    solo_usuario: bool = False,
) -> dict:
    """Historial persistente de eliminaciones aprobadas."""

    es_admin = str(usuario.get("rol") or "").lower() == "jefe"

    query = (
        supabase_admin
        .table("elementos_eliminados")
        .select("*")
        .order("fecha_eliminacion", desc=True)
    )

    if solo_usuario or not es_admin:
        query = query.eq(
            "solicitante_id",
            str(usuario.get("id") or ""),
        )

    respuesta = query.execute()
    filas = respuesta.data or []

    return {
        "status": "ok",
        "rol": usuario.get("rol"),
        "solo_usuario": bool(solo_usuario or not es_admin),
        "total": len(filas),
        "historial": filas,
    }


def procesar_solicitud_operacion(
    solicitud_id: str,
    aprobar: bool,
    resuelto_por: str
):
    """Procesa una solicitud MOVER/ELIMINAR.

    El destino de un movimiento ya fue validado cuando el usuario creó la
    solicitud. En la aprobación reutilizamos ese ID guardado para evitar
    recorrer nuevamente todo el árbol de Google Drive.
    """

    with DECISION_LOCK:
        respuesta = (
            supabase_admin
            .table("solicitudes_operacion")
            .select("*")
            .eq("id", solicitud_id)
            .limit(1)
            .execute()
        )

        if not respuesta.data:
            raise RuntimeError("La solicitud no existe.")

        solicitud = respuesta.data[0]

        if str(solicitud.get("estado") or "").upper() != "PENDIENTE":
            return solicitud, False

        if not aprobar:
            tipo_rechazado = str(
                solicitud.get("tipo_operacion") or ""
            ).upper()

            resultado_rechazo = (
                "Eliminación rechazada por el custodio. "
                "El elemento permanece activo y no fue enviado a la papelera."
                if tipo_rechazado == "ELIMINAR"
                else
                "Operación rechazada por el custodio."
            )

            # PUNTO 20:
            # Si se rechaza ELIMINAR, no se toca Drive, no se cambia
            # estado_archivo y no se crea historial de papelera.
            actualizado = (
                supabase_admin
                .table("solicitudes_operacion")
                .update({
                    "estado": ESTADO_RECHAZADO,
                    "fecha_resolucion": ahora_iso(),
                    "resuelto_por": str(resuelto_por),
                    "resultado": resultado_rechazo,
                })
                .eq("id", solicitud_id)
                .eq("estado", ESTADO_PENDIENTE)
                .execute()
            )

            solicitud_final = (actualizado.data or [solicitud])[0]

            registrar_evento_auditoria(
                evento="OPERACION_RECHAZADA",
                categoria="OPERACIONES",
                actor={
                    "id": str(resuelto_por),
                    "nombre": str(resuelto_por),
                    "rol": "custodio",
                },
                accion=tipo_rechazado,
                objeto_tipo=solicitud.get("objeto_tipo"),
                objeto_id=solicitud_id,
                estado=ESTADO_RECHAZADO,
                detalle={
                    "nombre_objeto": solicitud.get("nombre_objeto"),
                    "resultado": resultado_rechazo,
                },
            )

            return solicitud_final, True

        objeto = obtener_objeto_operable(
            str(solicitud.get("solicitante_id") or ""),
            solicitud.get("objeto_tipo"),
            auditoria_id=solicitud.get("auditoria_id"),
            lote_id=solicitud.get("lote_id"),
        )

        tipo = str(solicitud.get("tipo_operacion") or "").upper()
        cambios = {"fecha_ultima_operacion": ahora_iso()}

        if tipo == "ELIMINAR":
            resultado_drive = enviar_a_papelera_google_drive(
                objeto["drive_id"]
            )

            if not resultado_drive or not resultado_drive.get("trashed"):
                raise RuntimeError(
                    "Google Drive no confirmó el envío a la papelera."
                )

            fecha_eliminacion = ahora_iso()

            cambios.update({
                "en_drive": False,
                "estado_archivo": "ELIMINADO",
                "fecha_eliminacion": fecha_eliminacion,
            })

            # PUNTO 21:
            # Solo una eliminación aprobada y confirmada por Drive
            # entra al historial independiente.
            historial = registrar_historial_eliminacion(
                solicitud,
                objeto,
                str(resuelto_por),
            )

            resultado_texto = (
                "Elemento enviado a la papelera de Google Drive."
                if historial
                else
                "Elemento enviado a la papelera de Google Drive. "
                "Advertencia: el historial independiente no pudo registrarse; "
                "revisa los logs del backend."
            )

        elif tipo == "MOVER":
            destino_id = str(
                solicitud.get("carpeta_destino_id") or ""
            ).strip()

            if not destino_id:
                raise RuntimeError("La solicitud no tiene carpeta de destino.")

            if destino_id == objeto["drive_parent_id"]:
                raise RuntimeError("El elemento ya se encuentra en esa carpeta.")

            if objeto["tipo"] == "CARPETA" and destino_id == objeto["drive_id"]:
                raise RuntimeError("Una carpeta no puede moverse dentro de sí misma.")

            mover_objeto_google_drive(
                objeto["drive_id"],
                destino_id
            )

            ubicacion_destino = str(
                solicitud.get("carpeta_destino_nombre")
                or "DRIVE PROYECTO"
            ).strip()

            cambios.update({
                "en_drive": True,
                "estado_archivo": "ACTIVO",
                "drive_parent_id": destino_id,
                "ubicacion_drive": ubicacion_destino,
            })
            resultado_texto = f"Elemento movido a {ubicacion_destino}."

        else:
            raise RuntimeError("Tipo de operación no soportado.")

        actualizar_auditoria_operacion(objeto, cambios)

        destino_final_id = None
        destino_final_ruta = None

        if tipo == "MOVER":
            destino_final_id = str(
                solicitud.get("carpeta_destino_id") or ""
            ) or None
            destino_final_ruta = (
                solicitud.get("carpeta_destino_nombre")
                or None
            )
        elif tipo == "ELIMINAR":
            destino_final_id = str(objeto.get("drive_id") or "") or None
            destino_final_ruta = "PAPELERA_GOOGLE_DRIVE"

        actualizado = (
            supabase_admin
            .table("solicitudes_operacion")
            .update({
                "estado": ESTADO_APROBADO,
                "fecha_resolucion": ahora_iso(),
                "resuelto_por": str(resuelto_por),
                "resultado": resultado_texto,
                "destino_final_id": destino_final_id,
                "destino_final_ruta": destino_final_ruta,
            })
            .eq("id", solicitud_id)
            .eq("estado", ESTADO_PENDIENTE)
            .execute()
        )

        solicitud_final = (actualizado.data or [solicitud])[0]

        registrar_evento_auditoria(
            evento="OPERACION_APROBADA",
            categoria="OPERACIONES",
            actor={
                "id": str(resuelto_por),
                "nombre": str(resuelto_por),
                "rol": "custodio",
            },
            accion=tipo,
            objeto_tipo=objeto.get("tipo"),
            objeto_id=solicitud_id,
            estado=ESTADO_APROBADO,
            detalle={
                "nombre_objeto": objeto.get("nombre"),
                "resultado": resultado_texto,
                "destino_final_id": destino_final_id,
                "destino_final_ruta": destino_final_ruta,
            },
        )

        return solicitud_final, True


def procesar_operacion_telegram_background(
    solicitud_id: str,
    aprobar_operacion: bool,
    user_id: str,
    chat_id,
    message_id
):
    """Completa MOVER/ELIMINAR después de responder el webhook a Telegram."""

    try:
        solicitud, cambio = procesar_solicitud_operacion(
            solicitud_id,
            aprobar_operacion,
            str(user_id),
        )

        sincronizar_mensajes_telegram_operacion(
            solicitud,
            str(user_id),
            fallback_chat_id=chat_id,
            fallback_message_id=message_id,
        )

    except Exception as error:
        print("[OPERACION BACKGROUND ERROR]", error)

        solicitud = {"id": solicitud_id}
        try:
            respuesta = (
                supabase_admin
                .table("solicitudes_operacion")
                .select("*")
                .eq("id", solicitud_id)
                .limit(1)
                .execute()
            )
            if respuesta.data:
                solicitud = respuesta.data[0]
        except Exception as consulta_error:
            print("[OPERACION BACKGROUND RECOVERY ERROR]", consulta_error)

        sincronizar_mensajes_telegram_operacion(
            solicitud,
            str(user_id),
            fallback_chat_id=chat_id,
            fallback_message_id=message_id,
            error_texto=str(error),
        )


def notificar_resultado_custodia_web_telegram(
    payload: dict,
    resultado: dict,
    usuario_admin: dict,
):
    """Actualiza la alerta individual y la bandeja después de decidir en web."""
    if not TELEGRAM_TOKEN or not AUTHORIZED_CHAT_IDS:
        return

    objeto_tipo = str(payload.get("objeto_tipo") or "").upper()
    estado = str(resultado.get("estado") or "").upper()

    try:
        if objeto_tipo == "CARPETA":
            actualizar_alertas_subida_telegram(
                lote_id=payload.get("lote_id") or resultado.get("lote_id"),
                estado=estado,
            )
        else:
            actualizar_alertas_subida_telegram(
                auditoria_id=payload.get("auditoria_id") or resultado.get("auditoria_id"),
                estado=estado,
            )
    except Exception as error:
        print("[TELEGRAM CUSTODY WEB ALERT ERROR]", error)

    for chat_id in AUTHORIZED_CHAT_IDS:
        try:
            mostrar_menu_usuarios(chat_id)
        except Exception as error:
            print("[TELEGRAM CUSTODY WEB MENU ERROR]", error)


# ============================================================
# TELEGRAM
# ============================================================

def telegram_request(
    metodo: str,
    payload: dict
):

    if not TELEGRAM_TOKEN:

        return {
            "ok": False,
            "description":
                "TELEGRAM_TOKEN no configurado"
        }


    url = (
        f"https://api.telegram.org/"
        f"bot{TELEGRAM_TOKEN}/{metodo}"
    )


    try:

        response = HTTP_CLIENT.post(
            url,
            json=payload,
            timeout=15
        )


        try:

            data = response.json()

        except Exception:

            data = {
                "ok": False,
                "description": response.text
            }


        if (
            metodo == "editMessageText"
            and response.status_code == 400
            and "message is not modified"
            in str(
                data.get(
                    "description",
                    ""
                )
            ).lower()
        ):

            print(
                "[TELEGRAM] editMessageText "
                "sin cambios; se considera OK"
            )

            return {
                "ok": True,
                "unchanged": True,
                "description": "Sin cambios",
            }


        print(
            f"[TELEGRAM] {metodo} "
            f"HTTP={response.status_code} "
            f"respuesta={data}"
        )


        return data


    except Exception as error:

        print(
            f"[TELEGRAM ERROR] "
            f"{metodo}: {error}"
        )


        return {
            "ok": False,
            "description": str(error)
        }


def formatear_bytes_telegram(valor) -> str:
    try:
        numero = float(valor or 0)
    except (TypeError, ValueError):
        numero = 0.0

    unidades = ["B", "KB", "MB", "GB", "TB"]
    indice = 0
    while numero >= 1024 and indice < len(unidades) - 1:
        numero /= 1024.0
        indice += 1

    if indice == 0:
        return f"{int(numero)} {unidades[indice]}"
    return f"{numero:.1f} {unidades[indice]}"


def registrar_notificacion_telegram(
    chat_id,
    message_id,
    tipo_solicitud: str,
    solicitud_id=None,
    auditoria_id=None,
    lote_id=None,
    estado: str = "PENDIENTE",
):
    """Persiste cada alerta individual enviada a Telegram."""
    if not message_id:
        return None

    payload = {
        "chat_id": str(chat_id),
        "message_id": int(message_id),
        "tipo_solicitud": str(tipo_solicitud or "SOLICITUD").upper(),
        "solicitud_id": str(solicitud_id) if solicitud_id else None,
        "auditoria_id": str(auditoria_id) if auditoria_id else None,
        "lote_id": str(lote_id) if lote_id else None,
        "estado": str(estado or "PENDIENTE").upper(),
        "created_at": ahora_iso(),
        "updated_at": ahora_iso(),
    }

    try:
        respuesta = (
            supabase_admin
            .table("telegram_notificaciones")
            .insert(payload)
            .execute()
        )
        return (respuesta.data or [None])[0]
    except Exception as error:
        print("[TELEGRAM NOTIFICACION SAVE ERROR]", error)
        return None


def buscar_notificaciones_telegram(
    solicitud_id=None,
    auditoria_id=None,
    lote_id=None,
    chat_id=None,
    estado=None,
):
    try:
        query = (
            supabase_admin
            .table("telegram_notificaciones")
            .select("*")
            .order("created_at", desc=True)
        )

        if solicitud_id:
            query = query.eq("solicitud_id", str(solicitud_id))
        elif lote_id:
            query = query.eq("lote_id", str(lote_id))
        elif auditoria_id:
            query = query.eq("auditoria_id", str(auditoria_id))
        else:
            return []

        if chat_id is not None:
            query = query.eq("chat_id", str(chat_id))
        if estado:
            query = query.eq("estado", str(estado).upper())

        respuesta = query.execute()
        return respuesta.data or []
    except Exception as error:
        print("[TELEGRAM NOTIFICACION LOAD ERROR]", error)
        return []


def actualizar_estado_notificacion_db(notificacion_id, estado: str):
    try:
        (
            supabase_admin
            .table("telegram_notificaciones")
            .update({
                "estado": str(estado).upper(),
                "updated_at": ahora_iso(),
            })
            .eq("id", str(notificacion_id))
            .execute()
        )
    except Exception as error:
        print("[TELEGRAM NOTIFICACION STATE ERROR]", error)


def construir_texto_alerta_subida(documentos: list, estado: str = "PENDIENTE") -> str:
    documentos = list(documentos or [])
    if not documentos:
        return "🔔 SOLICITUD DLP\n\nNo se encontró información de la subida."

    primero = documentos[0]
    estado = str(estado or primero.get("estado") or "PENDIENTE").upper()
    usuario = primero.get("solicitante_nombre") or primero.get("solicitante_correo") or "Usuario"
    destino = primero.get("ubicacion_drive") or "DRIVE PROYECTO"
    total = sum(int(doc.get("tamano_bytes") or 0) for doc in documentos)

    nombre_carpeta = obtener_carpeta_desde_ruta(primero.get("ruta_relativa"))
    if nombre_carpeta:
        accion = "📁 SUBIR CARPETA"
        objeto = nombre_carpeta
        detalle = f"📄 Archivos: {len(documentos)}\n📦 Tamaño total: {formatear_bytes_telegram(total)}"
    elif len(documentos) > 1:
        accion = "📤 SUBIR ARCHIVOS"
        objeto = f"{len(documentos)} archivos"
        detalle = f"📦 Tamaño total: {formatear_bytes_telegram(total)}"
    else:
        accion = "📤 SUBIR ARCHIVO"
        objeto = primero.get("nombre_archivo") or "Archivo"
        detalle = f"📦 {formatear_bytes_telegram(total)}"

    iconos = {
        "PENDIENTE": "⏳",
        "APROBADO": "✅",
        "RECHAZADO": "❌",
        "ERROR": "⚠️",
    }
    icono = iconos.get(estado, "ℹ️")

    return (
        "🔔 SOLICITUD DLP\n\n"
        f"👤 {usuario}\n"
        f"{accion}\n"
        f"📄 {objeto}\n"
        f"📁 Destino: {destino}\n"
        f"{detalle}\n\n"
        f"{icono} {estado}"
    )


def construir_texto_alerta_operacion(solicitud: dict, estado: str | None = None) -> str:
    solicitud = solicitud or {}
    tipo = str(solicitud.get("tipo_operacion") or "OPERACION").upper()
    objeto_tipo = str(solicitud.get("objeto_tipo") or "ARCHIVO").upper()
    estado = str(estado or solicitud.get("estado") or "PENDIENTE").upper()

    icono_tipo = "↔️" if tipo == "MOVER" else "🗑️" if tipo == "ELIMINAR" else "📌"
    icono_estado = {
        "PENDIENTE": "⏳",
        "APROBADO": "✅",
        "RECHAZADO": "❌",
        "ERROR": "⚠️",
    }.get(estado, "ℹ️")

    lineas = [
        "🔔 SOLICITUD DLP",
        "",
        f"👤 {solicitud.get('solicitante_nombre') or solicitud.get('solicitante_correo') or 'Usuario'}",
        f"{icono_tipo} {tipo} {objeto_tipo}",
        f"📄 {solicitud.get('nombre_objeto') or 'Archivo/Carpeta'}",
        f"📂 Origen: {solicitud.get('carpeta_origen') or 'DRIVE PROYECTO'}",
    ]

    if tipo == "MOVER":
        lineas.append(
            f"➡️ Destino: {solicitud.get('carpeta_destino_nombre') or 'Destino seleccionado'}"
        )

    lineas.extend(["", f"{icono_estado} {estado}"])

    if solicitud.get("resultado") and estado != "PENDIENTE":
        lineas.extend(["", f"📝 {solicitud.get('resultado')}"])

    return "\n".join(lineas)


def _boton_bandeja_para_usuario(solicitante_id: str):
    # No reutilizamos el mensaje de alerta como bandeja. Este callback solo
    # refresca el mensaje principal persistido en telegram_bandejas.
    return {
        "inline_keyboard": [[
            {"text": "📥 Ver en bandeja", "callback_data": "tray:refresh"}
        ]]
    }


def enviar_alerta_subida_telegram(documentos: list):
    documentos = list(documentos or [])
    if not documentos or not TELEGRAM_TOKEN or not AUTHORIZED_CHAT_IDS:
        return []

    primero = documentos[0]
    auditoria_id = primero.get("id") if len(documentos) == 1 else None
    lote_id = primero.get("lote_id") if len(documentos) > 1 or primero.get("lote_id") else None
    solicitante_id = primero.get("solicitante_id")
    texto = construir_texto_alerta_subida(documentos, "PENDIENTE")
    resultados = []

    for chat_id in AUTHORIZED_CHAT_IDS:
        existentes = buscar_notificaciones_telegram(
            auditoria_id=auditoria_id,
            lote_id=lote_id,
            chat_id=chat_id,
            estado="PENDIENTE",
        )
        if existentes:
            resultados.append({"chat_id": chat_id, "ok": True, "deduplicated": True})
            continue

        resultado = telegram_request(
            "sendMessage",
            {
                "chat_id": chat_id,
                "text": texto,
                "reply_markup": _boton_bandeja_para_usuario(solicitante_id),
            },
        )
        message_id = extraer_message_id_telegram(resultado)
        if (resultado or {}).get("ok") and message_id:
            registrar_notificacion_telegram(
                chat_id,
                message_id,
                "SUBIDA",
                auditoria_id=auditoria_id,
                lote_id=lote_id,
            )
        resultados.append({
            "chat_id": chat_id,
            "ok": bool((resultado or {}).get("ok")),
            "description": (resultado or {}).get("description", "OK"),
        })

    return resultados


def enviar_alerta_operacion_telegram(solicitud: dict):
    if not solicitud or not TELEGRAM_TOKEN or not AUTHORIZED_CHAT_IDS:
        return []

    solicitud_id = solicitud.get("id")
    texto = construir_texto_alerta_operacion(solicitud, "PENDIENTE")
    resultados = []

    for chat_id in AUTHORIZED_CHAT_IDS:
        existentes = buscar_notificaciones_telegram(
            solicitud_id=solicitud_id,
            chat_id=chat_id,
            estado="PENDIENTE",
        )
        if existentes:
            resultados.append({"chat_id": chat_id, "ok": True, "deduplicated": True})
            continue

        resultado = telegram_request(
            "sendMessage",
            {
                "chat_id": chat_id,
                "text": texto,
                "reply_markup": _boton_bandeja_para_usuario(solicitud.get("solicitante_id")),
            },
        )
        message_id = extraer_message_id_telegram(resultado)
        if (resultado or {}).get("ok") and message_id:
            registrar_notificacion_telegram(
                chat_id,
                message_id,
                str(solicitud.get("tipo_operacion") or "OPERACION").upper(),
                solicitud_id=solicitud_id,
            )
        resultados.append({
            "chat_id": chat_id,
            "ok": bool((resultado or {}).get("ok")),
            "description": (resultado or {}).get("description", "OK"),
        })

    return resultados


def obtener_documentos_notificacion_subida(auditoria_id=None, lote_id=None):
    try:
        query = supabase_admin.table("auditoria_custodia").select("*")
        if lote_id:
            query = query.eq("lote_id", str(lote_id)).order("ruta_relativa")
        elif auditoria_id:
            query = query.eq("id", str(auditoria_id)).limit(1)
        else:
            return []
        return query.execute().data or []
    except Exception as error:
        print("[TELEGRAM SUBIDA LOAD ERROR]", error)
        return []


def actualizar_alertas_subida_telegram(auditoria_id=None, lote_id=None, estado=None):
    documentos = obtener_documentos_notificacion_subida(
        auditoria_id=auditoria_id,
        lote_id=lote_id,
    )
    if not documentos:
        return

    estado_final = str(
        estado
        or documentos[0].get("estado")
        or "PENDIENTE"
    ).upper()
    texto = construir_texto_alerta_subida(documentos, estado_final)
    notificaciones = buscar_notificaciones_telegram(
        auditoria_id=auditoria_id,
        lote_id=lote_id,
    )

    for notificacion in notificaciones:
        resultado = telegram_request(
            "editMessageText",
            {
                "chat_id": notificacion.get("chat_id"),
                "message_id": notificacion.get("message_id"),
                "text": texto,
                "reply_markup": _boton_bandeja_para_usuario(
                    documentos[0].get("solicitante_id")
                ),
            },
        )
        if not (resultado or {}).get("ok"):
            print("[TELEGRAM SUBIDA UPDATE ERROR]", resultado)
        actualizar_estado_notificacion_db(notificacion.get("id"), estado_final)


def actualizar_alertas_operacion_telegram(solicitud: dict, error_texto: str | None = None):
    if not solicitud:
        return

    estado = "PENDIENTE" if error_texto else str(solicitud.get("estado") or "PENDIENTE").upper()
    solicitud_texto = dict(solicitud)
    if error_texto:
        solicitud_texto["resultado"] = f"No se pudo completar: {str(error_texto)[:500]}"
    texto = construir_texto_alerta_operacion(solicitud_texto, estado)

    for notificacion in buscar_notificaciones_telegram(
        solicitud_id=solicitud.get("id")
    ):
        resultado = telegram_request(
            "editMessageText",
            {
                "chat_id": notificacion.get("chat_id"),
                "message_id": notificacion.get("message_id"),
                "text": texto,
                "reply_markup": _boton_bandeja_para_usuario(
                    solicitud.get("solicitante_id")
                ),
            },
        )
        if not (resultado or {}).get("ok"):
            print("[TELEGRAM OPERACION ALERT UPDATE ERROR]", resultado)
        actualizar_estado_notificacion_db(notificacion.get("id"), estado)


def refrescar_bandejas_telegram(excluir_chat_id=None):
    for destino_chat_id in AUTHORIZED_CHAT_IDS:
        if (
            excluir_chat_id is not None
            and str(destino_chat_id) == str(excluir_chat_id)
        ):
            continue
        try:
            mostrar_menu_usuarios(destino_chat_id)
        except Exception as error:
            print("[TELEGRAM BANDEJA REFRESH ERROR]", destino_chat_id, error)


def notificar_nueva_subida_background(auditoria_id=None, lote_id=None):
    """Genera una alerta real y refresca la bandeja usando Supabase como fuente."""
    documentos = obtener_documentos_notificacion_subida(
        auditoria_id=auditoria_id,
        lote_id=lote_id,
    )
    if documentos:
        enviar_alerta_subida_telegram(documentos)

    for chat_id in AUTHORIZED_CHAT_IDS:
        try:
            mostrar_menu_usuarios(chat_id)
        except Exception as error:
            print("[TELEGRAM SUBIDA MENU ERROR]", error)



# ============================================================
# PANEL PRINCIPAL INLINE DE TELEGRAM
# ============================================================

def quitar_teclado_inferior(chat_id):
    """Retira cualquier ReplyKeyboard antiguo que haya quedado en Telegram.

    Versiones anteriores del bot mostraban un teclado persistente abajo.
    El panel actual usa exclusivamente botones inline dentro del mensaje.
    """

    resultado = telegram_request(
        "sendMessage",
        {
            "chat_id": chat_id,
            "text": "Actualizando panel de custodia…",
            "reply_markup": {
                "remove_keyboard": True
            },
            "disable_notification": True,
        },
    )

    try:

        message_id = (
            resultado.get("result")
            or {}
        ).get("message_id")

        if resultado.get("ok") and message_id:

            telegram_request(
                "deleteMessage",
                {
                    "chat_id": chat_id,
                    "message_id": message_id
                },
            )

    except Exception as error:

        print(
            "[TELEGRAM REMOVE KEYBOARD]",
            error
        )

    return resultado


def obtener_resumen_panel():
    """Resumen unificado de SUBIDAS, MOVIMIENTOS y ELIMINACIONES pendientes."""
    try:
        respuesta = (
            supabase_admin
            .table("auditoria_custodia")
            .select("estado,solicitante_id,lote_id,ruta_relativa")
            .execute()
        )
        registros = respuesta.data or []
    except Exception as error:
        print("[PANEL RESUMEN ERROR]", error)
        registros = []

    try:
        respuesta_ops = (
            supabase_admin
            .table("solicitudes_operacion")
            .select("estado,tipo_operacion,solicitante_id")
            .execute()
        )
        operaciones = respuesta_ops.data or []
    except Exception as error:
        print("[PANEL OPERACIONES ERROR]", error)
        operaciones = []

    pendientes = [r for r in registros if str(r.get("estado") or "").upper() == "PENDIENTE"]
    aprobados = sum(1 for r in registros if str(r.get("estado") or "").upper() == "APROBADO")
    rechazados = sum(1 for r in registros if str(r.get("estado") or "").upper() == "RECHAZADO")

    operaciones_pendientes = [
        o for o in operaciones
        if str(o.get("estado") or "").upper() == "PENDIENTE"
    ]
    movimientos = sum(
        1 for o in operaciones_pendientes
        if str(o.get("tipo_operacion") or "").upper() == "MOVER"
    )
    eliminaciones = sum(
        1 for o in operaciones_pendientes
        if str(o.get("tipo_operacion") or "").upper() == "ELIMINAR"
    )

    usuarios = {
        str(r.get("solicitante_id"))
        for r in pendientes
        if r.get("solicitante_id")
    }
    usuarios.update({
        str(o.get("solicitante_id"))
        for o in operaciones_pendientes
        if o.get("solicitante_id")
    })

    carpetas = set()
    archivos_sueltos = 0
    for registro in pendientes:
        lote_id = str(registro.get("lote_id") or "").strip()
        ruta = str(registro.get("ruta_relativa") or "").replace("\\", "/").strip("/")
        if lote_id and "/" in ruta:
            carpetas.add(lote_id)
        else:
            archivos_sueltos += 1

    subidas_unidades = len(carpetas) + archivos_sueltos

    return {
        "pendientes": len(pendientes),
        "aprobados": aprobados,
        "rechazados": rechazados,
        "usuarios_pendientes": len(usuarios),
        "carpetas_pendientes": len(carpetas),
        "archivos_sueltos_pendientes": archivos_sueltos,
        "subidas_pendientes": subidas_unidades,
        "movimientos_pendientes": movimientos,
        "eliminaciones_pendientes": eliminaciones,
        "operaciones_pendientes": len(operaciones_pendientes),
        "solicitudes_pendientes": subidas_unidades + len(operaciones_pendientes),
    }

def mostrar_panel_principal(chat_id, message_id=None):
    resumen = obtener_resumen_panel()

    texto = (
        "🛡️ DataVault DLP - GM Ingenieros\n\n"
        "Panel de custodia\n\n"
        f"👥 Usuarios con pendientes: {resumen['usuarios_pendientes']}\n"
        f"📤 Subidas pendientes: {resumen['subidas_pendientes']}\n"
        f"↔️ Movimientos pendientes: {resumen['movimientos_pendientes']}\n"
        f"🗑️ Eliminaciones pendientes: {resumen['eliminaciones_pendientes']}\n"
        f"🟡 Solicitudes pendientes: {resumen['solicitudes_pendientes']}"
    )

    payload = {
        "chat_id": chat_id,
        "text": texto,
        "reply_markup": {
            "inline_keyboard": [
                [
                    {"text": "👥 Usuarios", "callback_data": "panel:usuarios"},
                    {"text": "🔄 Actualizar", "callback_data": "panel:actualizar"},
                ],
                [{"text": "📊 Estado", "callback_data": "panel:estado"}],
            ]
        },
    }

    if message_id:
        payload["message_id"] = message_id
        return telegram_request("editMessageText", payload)

    return telegram_request("sendMessage", payload)

def mostrar_estado_panel(chat_id, message_id=None):
    resumen = obtener_resumen_panel()

    texto = (
        "📊 ESTADO DATAVAULT\n\n"
        f"👥 Usuarios con pendientes: {resumen['usuarios_pendientes']}\n"
        f"📤 Subidas: {resumen['subidas_pendientes']}\n"
        f"↔️ Movimientos: {resumen['movimientos_pendientes']}\n"
        f"🗑️ Eliminaciones: {resumen['eliminaciones_pendientes']}\n"
        f"🟡 Total solicitudes: {resumen['solicitudes_pendientes']}\n\n"
        f"🟢 Documentos aprobados: {resumen['aprobados']}\n"
        f"🔴 Documentos rechazados: {resumen['rechazados']}"
    )

    payload = {
        "chat_id": chat_id,
        "text": texto,
        "reply_markup": {
            "inline_keyboard": [
                [
                    {"text": "👥 Usuarios", "callback_data": "panel:usuarios"},
                    {"text": "🔄 Actualizar", "callback_data": "panel:estado"},
                ],
                [{"text": "🏠 Panel", "callback_data": "panel:inicio"}],
            ]
        },
    }

    if message_id:
        payload["message_id"] = message_id
        return telegram_request("editMessageText", payload)

    return telegram_request("sendMessage", payload)

# ============================================================
# MENÚ TELEGRAM - USUARIOS -> ARCHIVOS -> DECISIÓN
# ============================================================

def obtener_documentos_pendientes():
    """Obtiene documentos PENDIENTES con datos para agrupar carpetas."""

    try:

        respuesta = (
            supabase_admin
            .table("auditoria_custodia")
            .select(
                "id,nombre_archivo,hash_sha256,"
                "tamano_bytes,estado,fecha_solicitud,"
                "solicitante_id,solicitante_nombre,"
                "solicitante_correo,lote_id,ruta_relativa"
            )
            .eq(
                "estado",
                "PENDIENTE"
            )
            .execute()
        )

        return respuesta.data or []

    except Exception as error:

        print(
            "[MENU TELEGRAM ERROR]",
            error
        )

        return []


def info_carpeta_documento(
    documento: dict
):

    lote_id = str(
        documento.get("lote_id")
        or ""
    ).strip()

    ruta = str(
        documento.get("ruta_relativa")
        or ""
    ).replace("\\", "/").strip("/")


    if not lote_id or "/" not in ruta:
        return None


    nombre_carpeta = obtener_carpeta_desde_ruta(
        ruta
    )


    if not nombre_carpeta:
        return None


    return {
        "lote_id":
            lote_id,

        "nombre":
            nombre_carpeta,
    }


def obtener_operaciones_pendientes(solicitante_id: str | None = None):
    """Obtiene solicitudes MOVER/ELIMINAR pendientes desde el backend admin."""
    try:
        query = (
            supabase_admin
            .table("solicitudes_operacion")
            .select("*")
            .eq("estado", "PENDIENTE")
            .order("fecha_solicitud", desc=True)
        )
        if solicitante_id:
            query = query.eq("solicitante_id", str(solicitante_id))
        respuesta = query.execute()
        return respuesta.data or []
    except Exception as error:
        print("[MENU TELEGRAM OPERACIONES ERROR]", error)
        return []



def obtener_message_id_bandeja_telegram(chat_id):
    """Obtiene desde Supabase el mensaje que actúa como bandeja principal."""
    try:
        respuesta = (
            supabase_admin
            .table("telegram_bandejas")
            .select("message_id,estado")
            .eq("chat_id", str(chat_id))
            .eq("estado", "ACTIVA")
            .limit(1)
            .execute()
        )
        if not respuesta.data:
            return None
        return int(respuesta.data[0].get("message_id"))
    except Exception as error:
        print("[TELEGRAM BANDEJA LOAD ERROR]", chat_id, error)
        return None


def guardar_message_id_bandeja_telegram(chat_id, message_id):
    """Persiste la bandeja en Supabase para sobrevivir a reinicios de Railway."""
    if not message_id:
        return

    chat_id_texto = str(chat_id)
    ahora = ahora_iso()

    try:
        existente = (
            supabase_admin
            .table("telegram_bandejas")
            .select("id")
            .eq("chat_id", chat_id_texto)
            .limit(1)
            .execute()
        )

        valores = {
            "message_id": int(message_id),
            "administrador_id": None,
            "administrador_nombre": None,
            "estado": "ACTIVA",
            "updated_at": ahora,
        }

        if existente.data:
            (
                supabase_admin
                .table("telegram_bandejas")
                .update(valores)
                .eq("id", existente.data[0]["id"])
                .execute()
            )
        else:
            valores.update({
                "chat_id": chat_id_texto,
                "created_at": ahora,
            })
            (
                supabase_admin
                .table("telegram_bandejas")
                .insert(valores)
                .execute()
            )

    except Exception as error:
        print("[TELEGRAM BANDEJA SAVE ERROR]", chat_id, error)


def extraer_message_id_telegram(resultado):
    """Extrae de forma segura el message_id devuelto por sendMessage."""
    try:
        return int(
            ((resultado or {}).get("result") or {}).get("message_id")
        )
    except (TypeError, ValueError):
        return None


def mostrar_menu_usuarios(chat_id, message_id=None, forzar_nuevo=False):
    """Bandeja Telegram unificada y agrupada por usuario."""
    documentos = obtener_documentos_pendientes()
    operaciones = obtener_operaciones_pendientes()
    usuarios = {}

    def asegurar_usuario(uid, nombre, correo):
        if uid not in usuarios:
            usuarios[uid] = {
                "nombre": nombre or correo or "Usuario",
                "correo": correo or "",
                "carpetas": set(),
                "archivos_sueltos": 0,
                "movimientos": 0,
                "eliminaciones": 0,
            }
        else:
            if not usuarios[uid].get("nombre") and nombre:
                usuarios[uid]["nombre"] = nombre
            if not usuarios[uid].get("correo") and correo:
                usuarios[uid]["correo"] = correo

    for documento in documentos:
        uid = str(documento.get("solicitante_id") or "").strip()
        if not uid:
            continue
        asegurar_usuario(
            uid,
            documento.get("solicitante_nombre"),
            documento.get("solicitante_correo"),
        )
        info = info_carpeta_documento(documento)
        if info:
            usuarios[uid]["carpetas"].add(info["lote_id"])
        else:
            usuarios[uid]["archivos_sueltos"] += 1

    for operacion in operaciones:
        uid = str(operacion.get("solicitante_id") or "").strip()
        if not uid:
            continue
        asegurar_usuario(
            uid,
            operacion.get("solicitante_nombre"),
            operacion.get("solicitante_correo"),
        )
        tipo = str(operacion.get("tipo_operacion") or "").upper()
        if tipo == "MOVER":
            usuarios[uid]["movimientos"] += 1
        elif tipo == "ELIMINAR":
            usuarios[uid]["eliminaciones"] += 1

    usuarios_ordenados = sorted(
        usuarios.items(),
        key=lambda item: str(item[1]["nombre"]).lower(),
    )

    botones = []
    total_subidas = 0
    total_movimientos = 0
    total_eliminaciones = 0

    for uid, usuario in usuarios_ordenados:
        subidas = len(usuario["carpetas"]) + usuario["archivos_sueltos"]
        movimientos = usuario["movimientos"]
        eliminaciones = usuario["eliminaciones"]
        total_subidas += subidas
        total_movimientos += movimientos
        total_eliminaciones += eliminaciones

        nombre = str(usuario["nombre"] or "Usuario")
        if len(nombre) > 24:
            nombre = nombre[:21] + "..."

        botones.append([{
            "text": f"👤 {nombre} · 📤{subidas} ↔️{movimientos} 🗑️{eliminaciones}",
            "callback_data": f"usr:{uid}",
        }])

    botones.append([{
        "text": "🔄 Actualizar",
        "callback_data": "menu:usuarios",
    }])

    if usuarios:
        texto = (
            "🛡️ DataVault DLP - GM Ingenieros\n\n"
            "📥 SOLICITUDES PENDIENTES\n\n"
            f"📤 Subidas: {total_subidas}\n"
            f"↔️ Movimientos: {total_movimientos}\n"
            f"🗑️ Eliminaciones: {total_eliminaciones}\n\n"
            "Seleccione un usuario:"
        )
    else:
        texto = (
            "🛡️ DataVault DLP - GM Ingenieros\n\n"
            "✅ No existen solicitudes pendientes."
        )

    payload = {
        "chat_id": chat_id,
        "text": texto,
        "reply_markup": {"inline_keyboard": botones},
    }

    # Si la acción viene de un botón inline, editamos ese mismo mensaje.
    # Para comandos o el teclado antiguo podemos forzar un mensaje nuevo para
    # que el usuario vea inmediatamente la respuesta, en vez de editar una
    # bandeja vieja que puede estar más arriba en el chat.
    objetivo_message_id = None

    if not forzar_nuevo:
        objetivo_message_id = (
            message_id
            or obtener_message_id_bandeja_telegram(chat_id)
        )

    if objetivo_message_id:
        payload_edicion = dict(payload)
        payload_edicion["message_id"] = objetivo_message_id

        resultado = telegram_request(
            "editMessageText",
            payload_edicion
        )

        if (resultado or {}).get("ok"):
            guardar_message_id_bandeja_telegram(
                chat_id,
                objetivo_message_id
            )
            return resultado

        # El mensaje pudo ser eliminado o quedar inaccesible. Solo en ese caso
        # creamos una nueva bandeja y desde ahí volvemos a reutilizarla.
        print(
            "[TELEGRAM BANDEJA] No se pudo editar; "
            "se creará una nueva bandeja.",
            chat_id,
            (resultado or {}).get("description"),
        )

    resultado = telegram_request(
        "sendMessage",
        payload
    )

    nuevo_message_id = extraer_message_id_telegram(
        resultado
    )

    if (resultado or {}).get("ok") and nuevo_message_id:
        guardar_message_id_bandeja_telegram(
            chat_id,
            nuevo_message_id
        )

    return resultado

def mostrar_archivos_usuario(chat_id, message_id, solicitante_id):
    """Muestra en una sola bandeja las 3 clases de solicitud del usuario."""
    try:
        respuesta = (
            supabase_admin
            .table("auditoria_custodia")
            .select(
                "id,nombre_archivo,tamano_bytes,fecha_solicitud,"
                "solicitante_id,solicitante_nombre,solicitante_correo,"
                "lote_id,ruta_relativa"
            )
            .eq("solicitante_id", solicitante_id)
            .eq("estado", "PENDIENTE")
            .order("fecha_solicitud", desc=True)
            .execute()
        )
        archivos = respuesta.data or []
    except Exception as error:
        print("[ARCHIVOS USUARIO ERROR]", error)
        archivos = []

    operaciones = obtener_operaciones_pendientes(solicitante_id)

    if not archivos and not operaciones:
        return mostrar_menu_usuarios(chat_id, message_id)

    fuente = archivos[0] if archivos else operaciones[0]
    nombre_usuario = (
        fuente.get("solicitante_nombre")
        or fuente.get("solicitante_correo")
        or "Usuario"
    )
    correo = fuente.get("solicitante_correo") or ""

    carpetas = {}
    sueltos = []
    for archivo in archivos:
        info = info_carpeta_documento(archivo)
        if info:
            lote_id = info["lote_id"]
            if lote_id not in carpetas:
                carpetas[lote_id] = {
                    "nombre": info["nombre"],
                    "cantidad": 0,
                    "tamano": 0,
                }
            carpetas[lote_id]["cantidad"] += 1
            carpetas[lote_id]["tamano"] += int(archivo.get("tamano_bytes") or 0)
        else:
            sueltos.append(archivo)

    movimientos = [
        o for o in operaciones
        if str(o.get("tipo_operacion") or "").upper() == "MOVER"
    ]
    eliminaciones = [
        o for o in operaciones
        if str(o.get("tipo_operacion") or "").upper() == "ELIMINAR"
    ]

    botones = []

    for lote_id, carpeta in carpetas.items():
        nombre = str(carpeta["nombre"] or "Carpeta")
        if len(nombre) > 28:
            nombre = nombre[:25] + "..."
        botones.append([{
            "text": f"📤 📁 {nombre} · {carpeta['cantidad']} archivos",
            "callback_data": f"lot:{lote_id}",
        }])

    for archivo in sueltos:
        nombre = str(archivo.get("nombre_archivo") or "Archivo")
        if len(nombre) > 36:
            nombre = nombre[:33] + "..."
        botones.append([{
            "text": f"📤 📄 {nombre}",
            "callback_data": f"doc:{archivo['id']}",
        }])

    for operacion in movimientos:
        nombre = str(operacion.get("nombre_objeto") or "Archivo/Carpeta")
        destino = str(operacion.get("carpeta_destino_nombre") or "Destino")
        if len(nombre) > 25:
            nombre = nombre[:22] + "..."
        if len(destino) > 18:
            destino = ".../" + destino.split("/")[-1].strip()
        botones.append([{
            "text": f"↔️ {nombre} → {destino}",
            "callback_data": f"opdet:{operacion['id']}",
        }])

    for operacion in eliminaciones:
        nombre = str(operacion.get("nombre_objeto") or "Archivo/Carpeta")
        if len(nombre) > 36:
            nombre = nombre[:33] + "..."
        botones.append([{
            "text": f"🗑️ {nombre}",
            "callback_data": f"opdet:{operacion['id']}",
        }])

    botones.append([{
        "text": "🔙 Usuarios",
        "callback_data": "menu:usuarios",
    }])

    subidas = len(carpetas) + len(sueltos)
    texto = (
        f"👤 {nombre_usuario}\n"
        f"📧 {correo}\n\n"
        f"📤 Subidas: {subidas}\n"
        f"↔️ Movimientos: {len(movimientos)}\n"
        f"🗑️ Eliminaciones: {len(eliminaciones)}\n\n"
        "Seleccione una solicitud:"
    )

    return telegram_request(
        "editMessageText",
        {
            "chat_id": chat_id,
            "message_id": message_id,
            "text": texto,
            "reply_markup": {"inline_keyboard": botones},
        },
    )


def mostrar_detalle_operacion(chat_id, message_id, solicitud_id):
    """Detalle y decisión de una solicitud MOVER/ELIMINAR desde la bandeja."""
    try:
        respuesta = (
            supabase_admin
            .table("solicitudes_operacion")
            .select("*")
            .eq("id", str(solicitud_id))
            .limit(1)
            .execute()
        )
        if not respuesta.data:
            raise ValueError("La solicitud no existe.")
        solicitud = respuesta.data[0]
    except Exception as error:
        print("[DETALLE OPERACION TELEGRAM ERROR]", error)
        return telegram_request(
            "editMessageText",
            {
                "chat_id": chat_id,
                "message_id": message_id,
                "text": "❌ No se pudo cargar la solicitud seleccionada.",
                "reply_markup": {
                    "inline_keyboard": [[
                        {"text": "👥 Usuarios", "callback_data": "menu:usuarios"}
                    ]]
                },
            },
        )

    tipo = str(solicitud.get("tipo_operacion") or "").upper()
    estado = str(solicitud.get("estado") or "").upper()
    icono = "↔️" if tipo == "MOVER" else "🗑️"
    etiqueta = "MOVIMIENTO" if tipo == "MOVER" else "ELIMINACIÓN"
    uid = str(solicitud.get("solicitante_id") or "")

    texto = (
        "🛡️ DataVault DLP - GM Ingenieros\n\n"
        f"{icono} SOLICITUD DE {etiqueta}\n\n"
        f"👤 Usuario: {solicitud.get('solicitante_nombre') or 'Usuario'}\n"
        f"📧 Correo: {solicitud.get('solicitante_correo') or ''}\n"
        f"📄 Objeto: {solicitud.get('nombre_objeto') or 'Archivo/Carpeta'}\n"
        f"📌 Tipo: {solicitud.get('objeto_tipo') or '—'}\n"
        f"📂 Origen: {solicitud.get('carpeta_origen') or 'DRIVE PROYECTO'}\n"
    )
    if tipo == "MOVER":
        texto += f"➡️ Destino: {solicitud.get('carpeta_destino_nombre') or '—'}\n"
    texto += f"\n🟡 Estado: {estado or 'PENDIENTE'}"

    botones = []
    if estado == "PENDIENTE":
        botones.append([
            {"text": "✅ Aprobar", "callback_data": f"opap:{solicitud_id}"},
            {"text": "❌ Rechazar", "callback_data": f"opre:{solicitud_id}"},
        ])
    if uid:
        botones.append([{
            "text": "🔙 Solicitudes del usuario",
            "callback_data": f"usr:{uid}",
        }])
    botones.append([{
        "text": "👥 Usuarios",
        "callback_data": "menu:usuarios",
    }])

    # Este detalle vive dentro de la misma bandeja unificada. Guardamos tanto
    # el message_id activo del chat como la referencia de la operación para
    # que una resolución hecha desde la web pueda actualizar ese mismo mensaje.
    guardar_message_id_bandeja_telegram(
        chat_id,
        message_id
    )

    try:
        referencias = normalizar_referencias_telegram(solicitud)
        referencia_actual = {
            "chat_id": int(chat_id),
            "message_id": int(message_id),
        }
        if referencia_actual not in referencias:
            referencias.append(referencia_actual)
            guardar_referencias_telegram_operacion(
                solicitud_id,
                referencias,
            )
    except Exception as error:
        print("[TELEGRAM DETALLE REF ERROR]", error)

    return telegram_request(
        "editMessageText",
        {
            "chat_id": chat_id,
            "message_id": message_id,
            "text": texto,
            "reply_markup": {"inline_keyboard": botones},
        },
    )

def mostrar_detalle_lote(
    chat_id,
    message_id,
    lote_id
):

    try:

        respuesta = (
            supabase_admin
            .table("auditoria_custodia")
            .select(
                "id,nombre_archivo,hash_sha256,"
                "tamano_bytes,estado,solicitante_id,"
                "solicitante_nombre,solicitante_correo,"
                "lote_id,ruta_relativa,drive_parent_id,"
                "ubicacion_drive,temp_storage_path,"
                "temp_storage_bucket"
            )
            .eq(
                "lote_id",
                lote_id
            )
            .eq(
                "estado",
                "PENDIENTE"
            )
            .order(
                "fecha_solicitud",
                desc=False
            )
            .execute()
        )

        documentos = respuesta.data or []

    except Exception as error:

        print(
            "[DETALLE LOTE ERROR]",
            error
        )

        documentos = []


    if not documentos:

        return telegram_request(

            "editMessageText",

            {

                "chat_id":
                    chat_id,

                "message_id":
                    message_id,

                "text":
                    (
                        "🛡️ DataVault DLP - GM Ingenieros\n\n"
                        "⚠️ Esta carpeta ya no tiene "
                        "documentos pendientes."
                    ),

                "reply_markup": {

                    "inline_keyboard": [[
                        {
                            "text":
                                "👥 Usuarios",

                            "callback_data":
                                "menu:usuarios",
                        }
                    ]]

                },

            },

        )


    primer = documentos[0]


    solicitante_id = str(
        primer.get(
            "solicitante_id"
        )
        or ""
    ).strip()


    nombre_usuario = (
        primer.get(
            "solicitante_nombre"
        )
        or
        "Usuario"
    )


    correo = (
        primer.get(
            "solicitante_correo"
        )
        or ""
    )


    destino_drive = str(
        primer.get("ubicacion_drive")
        or "DRIVE PROYECTO"
    )


    nombre_carpeta = (
        obtener_carpeta_desde_ruta(
            primer.get(
                "ruta_relativa"
            )
        )
        or
        f"LOTE_{lote_id[:8]}"
    )


    total_bytes = sum(
        int(
            d.get(
                "tamano_bytes"
            )
            or 0
        )
        for d in documentos
    )


    def formato_bytes(valor):

        unidades = [
            "B",
            "KB",
            "MB",
            "GB"
        ]

        numero = float(
            valor or 0
        )

        indice = 0


        while (
            numero >= 1024
            and
            indice < len(unidades) - 1
        ):

            numero /= 1024
            indice += 1


        return (
            f"{numero:.2f} "
            f"{unidades[indice]}"
            if indice
            else
            f"{int(numero)} B"
        )


    nombres = []


    for doc in documentos[:8]:

        ruta = str(
            doc.get(
                "ruta_relativa"
            )
            or
            doc.get(
                "nombre_archivo"
            )
            or
            "Archivo"
        ).replace("\\", "/")


        partes = [
            p
            for p in ruta.split("/")
            if p
        ]


        if (
            partes
            and
            partes[0]
            == nombre_carpeta
        ):

            partes = partes[1:]


        nombres.append(
            "/".join(partes)
            if partes
            else
            (
                doc.get(
                    "nombre_archivo"
                )
                or
                "Archivo"
            )
        )


    contenido = "\n".join(
        f"• {nombre}"
        for nombre in nombres
    )


    if len(documentos) > len(nombres):

        contenido += (
            f"\n• ... y "
            f"{len(documentos) - len(nombres)} "
            f"archivo(s) más"
        )


    texto = (
        "🛡️ DataVault DLP - GM Ingenieros\n\n"
        f"📁 Carpeta: {nombre_carpeta}\n"
        f"👤 Solicitante: {nombre_usuario}\n"
        f"📧 Correo: {correo}\n"
        f"📄 Archivos: {len(documentos)}\n"
        f"📦 Tamaño total: {formato_bytes(total_bytes)}\n"
        f"📂 Destino: {destino_drive}\n\n"
        f"Contenido:\n{contenido}\n\n"
        "¿Autoriza la transferencia de TODA "
        "la carpeta a Google Drive?"
    )


    botones = [

        [
            {
                "text":
                    "✅ Aprobar carpeta",

                "callback_data":
                    f"aplot:{lote_id}",
            },

            {
                "text":
                    "❌ Rechazar carpeta",

                "callback_data":
                    f"relot:{lote_id}",
            },
        ],

        [
            {
                "text":
                    "🔙 Volver",

                "callback_data":
                    f"usr:{solicitante_id}",
            },

            {
                "text":
                    "👥 Usuarios",

                "callback_data":
                    "menu:usuarios",
            },
        ],

    ]


    return telegram_request(

        "editMessageText",

        {

            "chat_id":
                chat_id,

            "message_id":
                message_id,

            "text":
                texto,

            "reply_markup": {
                "inline_keyboard":
                    botones
            },

        },

    )


def mostrar_detalle_documento(
    chat_id,
    message_id,
    auditoria_id
):

    try:

        respuesta = (
            supabase_admin
            .table(
                "auditoria_custodia"
            )
            .select(
                "id,nombre_archivo,hash_sha256,"
                "estado,solicitante_id,"
                "solicitante_nombre,solicitante_correo,"
                "drive_parent_id,ubicacion_drive"
            )
            .eq(
                "id",
                auditoria_id
            )
            .limit(1)
            .execute()
        )


        if not respuesta.data:

            raise ValueError(
                "No existe el documento seleccionado."
            )


        documento = respuesta.data[0]


    except Exception as error:

        print(
            "[DETALLE DOCUMENTO ERROR]",
            error
        )


        return telegram_request(

            "editMessageText",

            {

                "chat_id":
                    chat_id,

                "message_id":
                    message_id,

                "text":
                    (
                        "🛡️ DataVault DLP - GM Ingenieros\n\n"
                        "❌ No se pudo cargar el documento seleccionado."
                    ),

                "reply_markup": {

                    "inline_keyboard": [[
                        {
                            "text":
                                "👥 Usuarios",

                            "callback_data":
                                "menu:usuarios",
                        }
                    ]]

                },

            },

        )


    solicitante_id = str(
        documento.get(
            "solicitante_id"
        )
        or ""
    ).strip()


    estado = (
        documento.get(
            "estado"
        )
        or
        "DESCONOCIDO"
    )


    if estado != "PENDIENTE":

        texto = (
            "🛡️ DataVault DLP - GM Ingenieros\n\n"
            f"📁 Archivo: "
            f"{documento.get('nombre_archivo') or 'Archivo'}\n"
            f"👤 Solicitante: "
            f"{documento.get('solicitante_nombre') or 'Usuario'}\n\n"
            f"⚠️ Este documento ya fue procesado.\n"
            f"Estado actual: {estado}"
        )


        botones = [[
            {
                "text":
                    "👥 Usuarios",

                "callback_data":
                    "menu:usuarios",
            }
        ]]


    else:

        texto = (
            "🛡️ DataVault DLP - GM Ingenieros\n\n"
            f"📁 Archivo: "
            f"{documento.get('nombre_archivo') or 'Archivo'}\n"
            f"👤 Solicitante: "
            f"{documento.get('solicitante_nombre') or 'Usuario'}\n"
            f"📧 Correo: "
            f"{documento.get('solicitante_correo') or ''}\n"
            f"📂 Destino: "
            f"{documento.get('ubicacion_drive') or 'DRIVE PROYECTO'}\n\n"
            f"🔑 Hash SHA-256:\n"
            f"{documento.get('hash_sha256') or ''}\n\n"
            f"🆔 Auditoría:\n"
            f"{documento.get('id')}\n\n"
            "¿Autoriza su transferencia "
            "a la custodia corporativa?"
        )


        botones = [

            [
                {
                    "text":
                        "✅ Aprobar",

                    "callback_data":
                        f"aprobar:{auditoria_id}",
                },

                {
                    "text":
                        "❌ Rechazar",

                    "callback_data":
                        f"rechazar:{auditoria_id}",
                },
            ],

            [
                {
                    "text":
                        "🔙 Archivos",

                    "callback_data":
                        f"usr:{solicitante_id}",
                },

                {
                    "text":
                        "👥 Usuarios",

                    "callback_data":
                        "menu:usuarios",
                },
            ],

        ]


    return telegram_request(

        "editMessageText",

        {

            "chat_id":
                chat_id,

            "message_id":
                message_id,

            "text":
                texto,

            "reply_markup": {
                "inline_keyboard":
                    botones
            },

        },

    )


# ============================================================
# OBTENER DOMINIO ACTUAL DE RAILWAY
# ============================================================

def obtener_base_url_request(
    request: Request
):

    forwarded_host = request.headers.get(
        "x-forwarded-host",
        ""
    ).strip()


    host = (
        forwarded_host
        or request.headers.get(
            "host",
            ""
        ).strip()
    )


    forwarded_proto = request.headers.get(
        "x-forwarded-proto",
        ""
    ).strip()


    proto = (
        forwarded_proto
        or request.url.scheme
        or "https"
    )


    if (
        host.endswith(".railway.app")
        or host.endswith(".up.railway.app")
    ):

        proto = "https"


    if host:

        return (
            f"{proto}://{host}"
        ).rstrip("/")


    return str(
        request.base_url
    ).rstrip("/")


# ============================================================
# CONFIGURAR WEBHOOK
# ============================================================

def configurar_webhook_url(
    base_url: str
):
    if not TELEGRAM_TOKEN:
        return {
            "ok": False,
            "description": "TELEGRAM_TOKEN no configurado",
        }

    if not TELEGRAM_WEBHOOK_SECRET:
        return {
            "ok": False,
            "description": "TELEGRAM_WEBHOOK_SECRET no configurado",
        }

    base_url = str(base_url or "").strip().rstrip("/")
    if not base_url:
        return {
            "ok": False,
            "description": "No se pudo determinar la URL pública",
        }

    webhook_url = f"{base_url}/telegram-webhook"
    info = telegram_request("getWebhookInfo", {})
    url_actual = ""
    if (info or {}).get("ok"):
        url_actual = str(((info or {}).get("result") or {}).get("url") or "")

    # Telegram no expone el secret_token actual en getWebhookInfo. Por eso
    # reconfiguramos el webhook de forma idempotente para garantizar que el
    # secreto guardado en Railway quede aplicado.
    resultado = telegram_request(
        "setWebhook",
        {
            "url": webhook_url,
            "secret_token": TELEGRAM_WEBHOOK_SECRET,
            "allowed_updates": ["message", "callback_query"],
            "drop_pending_updates": False,
        },
    )

    return {
        "ok": bool((resultado or {}).get("ok")),
        "webhook": webhook_url,
        "changed": url_actual != webhook_url,
        "secret_configured": True,
        "telegram": resultado,
    }


# ============================================================
# INICIO DE SERVIDOR
# ============================================================

@app.on_event("startup")
def startup_event():

    print("=" * 60)

    print(
        "DATAVAULT DLP INICIADO"
    )

    print(
        "Google Drive:",
        (
            "CONFIGURADO"
            if google_drive_configurado()
            else
            "NO CONFIGURADO"
        )
    )

    print(
        "Telegram autorizados:",
        AUTHORIZED_CHAT_IDS
    )

    print(
        "PUBLIC_BASE_URL:",
        PUBLIC_BASE_URL
        or
        "(se detectará al subir)"
    )

    print("=" * 60)


    if PUBLIC_BASE_URL:

        print(

            "[WEBHOOK STARTUP]",

            configurar_webhook_url(
                PUBLIC_BASE_URL
            )

        )


# ============================================================
# WEB
# ============================================================

@app.get(
    "/",
    response_class=HTMLResponse
)
async def read_index():

    if os.path.exists(
        "index.html"
    ):

        with open(
            "index.html",
            "r",
            encoding="utf-8"
        ) as archivo:

            return archivo.read()


    return """
    <h1>DataVault DLP API activa</h1>
    """


# ============================================================
# HEALTH
# ============================================================

@app.get("/health")
def health_check():

    return {

        "status":
            "ok",

        "service":
            (
                "DataVault DLP API | "
                "GM Ingenieros y Consultores"
            ),

        "telegram_configurado":
            bool(
                TELEGRAM_TOKEN
            ),

        "usuarios_telegram":
            len(
                AUTHORIZED_CHAT_IDS
            ),

        "supabase_configurado":
            bool(
                SUPABASE_URL
                and
                SUPABASE_KEY
            ),

        "google_drive_configurado":
            google_drive_configurado(),

        "google_folder_id":
            GOOGLE_FOLDER_ID
            or None,

        "public_url":
            PUBLIC_BASE_URL
            or None
    }


# ============================================================
# INFORMACIÓN DEL WEBHOOK
# ============================================================

@app.get("/telegram-info")
def telegram_info(request: Request):

    obtener_admin_desde_request(request)

    if not TELEGRAM_TOKEN:

        return {
            "ok": False,
            "error":
                "TELEGRAM_TOKEN no configurado"
        }


    return telegram_request(
        "getWebhookInfo",
        {}
    )


# ============================================================
# FORZAR WEBHOOK MANUALMENTE
# ============================================================

@app.get("/set-webhook")
def set_webhook_manual(
    request: Request
):

    obtener_admin_desde_request(request)

    base_url = obtener_base_url_request(
        request
    )


    return configurar_webhook_url(
        base_url
    )


# ============================================================
# API DE POLÍTICAS DLP (PUNTOS 13, 15, 16 Y 17)
# ============================================================

@app.get("/dlp/policies")
def listar_politicas_dlp_api(request: Request):
    admin = obtener_admin_desde_request(request)
    politicas = obtener_politicas_dlp()

    return {
        "status": "ok",
        "usuario": {
            "id": admin.get("id"),
            "nombre": admin.get("nombre"),
            "correo": admin.get("correo"),
        },
        "politicas": list(politicas.values()),
    }


@app.post("/dlp/policies/update")
async def actualizar_politica_dlp_api(request: Request):
    admin = obtener_admin_desde_request(request)

    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="JSON inválido.",
        )

    codigo = _dlp_codigo(payload.get("codigo"))
    if not codigo:
        raise HTTPException(
            status_code=400,
            detail="Falta codigo.",
        )

    tiene_valor = "valor" in payload
    activa = payload.get("activa") if "activa" in payload else None

    politica = guardar_politica_dlp(
        admin=admin,
        codigo=codigo,
        activa=activa,
        valor_marker=tiene_valor,
        valor=payload.get("valor"),
        accion="ACTUALIZAR",
    )

    return {
        "status": "ok",
        "politica": politica,
    }


@app.post("/dlp/policies/reset")
async def restaurar_politicas_dlp_api(request: Request):
    admin = obtener_admin_desde_request(request)

    restauradas = []

    for codigo, config in DLP_POLITICAS_DEFAULT.items():
        politica = guardar_politica_dlp(
            admin=admin,
            codigo=codigo,
            activa=True,
            valor_marker=True,
            valor=config.get("valor"),
            accion="RESTAURAR",
        )
        restauradas.append(politica)

    return {
        "status": "ok",
        "total": len(restauradas),
        "politicas": restauradas,
    }


# ============================================================
# SUBIR ARCHIVO
# ============================================================

@app.post("/upload")
async def registrar_y_solicitar_custodia(

    request: Request,

    background_tasks: BackgroundTasks,

    file: UploadFile = File(...),

    carpeta: str = Form(
        "PLANOS"
    )

):

    usuario_auth = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    solicitante_id = (
        usuario_auth["id"]
    )

    solicitante_nombre = (
        usuario_auth["nombre"]
    )

    solicitante_correo = (
        usuario_auth["correo"]
    )


    usuario = solicitante_nombre
    base_url_actual = obtener_base_url_request(request)
    estado_webhook = {
        "ok": True,
        "source": "startup",
        "base_url": PUBLIC_BASE_URL or base_url_actual,
    }

    contenido = await file.read()


    nombre_original = (
        file.filename
        or
        "archivo_sin_nombre"
    )


    # DLP REAL DEL BACKEND:
    # el navegador ya no decide si el archivo puede ingresar.
    validar_archivo_dlp(
        nombre_original,
        len(contenido) if contenido else 0,
    )


    # SHA-256 es una política obligatoria: aunque una fila de Supabase se
    # manipule, el backend siempre mantiene esta validación activa.
    hash_sha256 = hashlib.sha256(
        contenido
    ).hexdigest()


    nombre_base, extension = os.path.splitext(
        nombre_original
    )


    try:

        res_existentes = (

            supabase_admin

            .table(
                "auditoria_custodia"
            )

            .select(
                "nombre_archivo"
            )

            .ilike(
                "nombre_archivo",
                f"{nombre_base}%{extension}"
            )

            .execute()

        )


        archivos_existentes = (
            res_existentes.data
            or []
        )


        if archivos_existentes:

            contador = (
                len(archivos_existentes)
                + 1
            )

            nombre_final = (
                f"{nombre_base} "
                f"({contador})"
                f"{extension}"
            )

        else:

            nombre_final = (
                nombre_original
            )


    except Exception as e_nombre:

        print(
            "[CORRELATIVO ERROR]",
            e_nombre
        )

        nombre_final = (
            nombre_original
        )


    registro = {

        "nombre_archivo":
            nombre_final,

        "hash_sha256":
            hash_sha256,

        "tamano_bytes":
            len(contenido),

        "estado":
            "PENDIENTE",

        "usuario_solicitante":
            solicitante_nombre,

        "solicitante_id":
            solicitante_id,

        "solicitante_nombre":
            solicitante_nombre,

        "solicitante_correo":
            solicitante_correo,

        "drive_parent_id":
            (GOOGLE_FOLDER_ID or GOOGLE_DRIVE_ROOT_ID),

        "ubicacion_drive":
            "DRIVE PROYECTO",

        "destino_solicitado_id":
            (GOOGLE_FOLDER_ID or GOOGLE_DRIVE_ROOT_ID),

        "destino_solicitado_ruta":
            "DRIVE PROYECTO"

    }


    try:

        respuesta_db = (

            supabase_admin

            .table(
                "auditoria_custodia"
            )

            .insert(
                registro
            )

            .execute()

        )


    except Exception as error:

        print(
            "[SUPABASE INSERT ERROR]",
            error
        )


        raise HTTPException(

            status_code=500,

            detail=(
                "No se pudo registrar "
                "en Supabase: "
                f"{error}"
            )

        )


    if not respuesta_db.data:

        raise HTTPException(

            status_code=500,

            detail=(
                "Supabase no devolvió "
                "el registro insertado."
            )

        )


    id_auditoria = (
        respuesta_db
        .data[0]
        .get("id")
    )


    temp_path = None

    try:
        temp_path = guardar_archivo_temporal(
            str(id_auditoria),
            solicitante_id,
            nombre_final,
            contenido,
        )

        temp_update = (
            supabase_admin
            .table("auditoria_custodia")
            .update({
                "temp_storage_path": temp_path,
                "temp_storage_bucket": SUPABASE_TEMP_BUCKET,
                "temp_storage_saved_at": ahora_iso(),
                "temp_storage_deleted_at": None,
            })
            .eq("id", id_auditoria)
            .execute()
        )

        if not temp_update.data:
            raise RuntimeError(
                "El temporal se guardó, pero no se pudo vincular a la auditoría."
            )

    except Exception as error:
        if temp_path:
            eliminar_archivo_temporal(temp_path)

        try:
            (
                supabase_admin
                .table("auditoria_custodia")
                .delete()
                .eq("id", id_auditoria)
                .execute()
            )
        except Exception as cleanup_error:
            print("[TEMP STORAGE ROLLBACK DB ERROR]", cleanup_error)

        raise HTTPException(
            status_code=500,
            detail=str(error),
        )

    finally:
        del contenido


    registrar_evento_auditoria(
        evento="SOLICITUD_SUBIDA",
        categoria="CUSTODIA",
        actor=usuario_auth,
        accion="SUBIR_ARCHIVO",
        objeto_tipo="ARCHIVO",
        objeto_id=str(id_auditoria),
        estado=ESTADO_PENDIENTE,
        detalle={
            "nombre_archivo": nombre_final,
            "sha256": hash_sha256,
            "destino_solicitado_id": (
                GOOGLE_FOLDER_ID or GOOGLE_DRIVE_ROOT_ID
            ),
            "destino_solicitado_ruta": "DRIVE PROYECTO",
        },
        request=request,
    )

    # La solicitud ya está persistida. Telegram se procesa en segundo plano:
    # una caída del bot no invalida ni elimina la solicitud.
    background_tasks.add_task(
        notificar_menu_pendientes_background,
        base_url_actual,
        id_auditoria,
        None,
    )

    return {
        "status": "ok",
        "mensaje": (
            "Documento registrado correctamente. "
            "La notificación de Telegram se procesará en segundo plano."
        ),
        "id_auditoria": id_auditoria,
        "sha256": hash_sha256,
        "telegram_notificacion": "en_cola",
        "webhook": estado_webhook,
    }


# ============================================================
# APOYO PARA CARGA POR LOTES
# ============================================================

def normalizar_ruta_relativa(
    ruta: str,
    nombre_archivo: str
) -> str:

    ruta_limpia = str(
        ruta
        or
        nombre_archivo
        or
        "archivo_sin_nombre"
    )

    ruta_limpia = (
        ruta_limpia
        .replace("\\", "/")
        .lstrip("/")
    )


    partes = [

        parte

        for parte
        in ruta_limpia.split("/")

        if parte not in (
            "",
            ".",
            ".."
        )

    ]


    if not partes:

        return (
            nombre_archivo
            or
            "archivo_sin_nombre"
        )


    return "/".join(
        partes
    )


def obtener_nombre_correlativo(
    nombre_original: str
) -> str:

    nombre_original = (
        nombre_original
        or
        "archivo_sin_nombre"
    )


    nombre_base, extension = (
        os.path.splitext(
            nombre_original
        )
    )


    try:

        res_existentes = (

            supabase_admin

            .table(
                "auditoria_custodia"
            )

            .select(
                "nombre_archivo"
            )

            .ilike(
                "nombre_archivo",
                f"{nombre_base}%{extension}"
            )

            .execute()

        )


        archivos_existentes = (
            res_existentes.data
            or []
        )


        if archivos_existentes:

            contador = (
                len(
                    archivos_existentes
                )
                + 1
            )

            return (
                f"{nombre_base} "
                f"({contador})"
                f"{extension}"
            )


        return nombre_original


    except Exception as error:

        print(
            "[CORRELATIVO BATCH ERROR]",
            error
        )

        return nombre_original


# ============================================================
# NOTIFICACIÓN TELEGRAM EN SEGUNDO PLANO
# ============================================================

def notificar_menu_pendientes_background(
    base_url: str = "",
    auditoria_id=None,
    lote_id=None,
):
    """Notifica una subida persistida y refresca la bandeja principal."""
    if not politica_dlp_activa("ALERTA_TELEGRAM"):
        print("[DLP] ALERTA_TELEGRAM desactivada; solicitud guardada sin aviso.")
        return
    if base_url and not PUBLIC_BASE_URL:
        try:
            estado_webhook = configurar_webhook_url(base_url)
            print(f"[WEBHOOK BG] {estado_webhook}")
        except Exception as error:
            print(f"[WEBHOOK BG ERROR] {error}")

    try:
        notificar_nueva_subida_background(
            auditoria_id=auditoria_id,
            lote_id=lote_id,
        )
    except Exception as error:
        # La solicitud ya está persistida: Telegram no puede invalidarla.
        print("[TELEGRAM SUBIDA BACKGROUND ERROR]", error)


# ============================================================
# SUBIR VARIOS ARCHIVOS / CARPETA COMPLETA
# ============================================================

@app.post("/upload-batch")
async def registrar_lote_custodia(

    request: Request,

    background_tasks:
        BackgroundTasks,

    files:
        List[UploadFile]
        = File(...),

    relative_paths:
        List[str]
        = Form(default=[]),

    carpeta:
        str
        = Form("PLANOS"),

    lote_id:
        str
        = Form(""),

    carpeta_destino_id:
        str
        = Form("")

):

    usuario_auth = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    solicitante_id = (
        usuario_auth["id"]
    )

    solicitante_nombre = (
        usuario_auth["nombre"]
    )

    solicitante_correo = (
        usuario_auth["correo"]
    )


    if not files:

        raise HTTPException(
            status_code=400,
            detail=
                "No se recibieron archivos."
        )


    if len(files) > MAX_BATCH_FILES:

        raise HTTPException(

            status_code=400,

            detail=(
                f"El lote contiene "
                f"{len(files)} archivos. "
                f"El máximo permitido actualmente "
                f"es {MAX_BATCH_FILES}."
            )

        )


    if (
        relative_paths
        and
        len(relative_paths) != len(files)
    ):

        raise HTTPException(

            status_code=400,

            detail=(
                "La cantidad de rutas relativas "
                "no coincide con la cantidad "
                "de archivos."
            )

        )
    carpeta_destino_id = str(
        carpeta_destino_id
        or ""
    ).strip()


    if not carpeta_destino_id:

        raise HTTPException(
            status_code=400,
            detail=(
                "Selecciona una carpeta destino de Google Drive. "
                "Si deseas dejar archivos sueltos, selecciona "
                "explícitamente DRIVE PROYECTO (raíz)."
            )
        )


    destino_upload = validar_destino_drive(
        carpeta_destino_id
    )


    destino_upload_id = str(
        destino_upload.get("id")
        or GOOGLE_FOLDER_ID
    )


    destino_upload_path = str(
        destino_upload.get("path")
        or destino_upload.get("name")
        or "DRIVE PROYECTO"
    )


    if lote_id.strip():

        try:

            lote_uuid = str(
                UUID(
                    lote_id.strip()
                )
            )

        except (
            ValueError,
            TypeError,
            AttributeError
        ):

            raise HTTPException(
                status_code=400,
                detail="lote_id inválido."
            )

    else:

        lote_uuid = str(
            uuid4()
        )


    base_url_actual = (
        obtener_base_url_request(
            request
        )
    )


    estado_webhook = {

        "ok":
            True,

        "changed":
            False,

        "source":
            "background",

        "base_url":
            base_url_actual
    }


    procesados = []
    errores = []
    cancelado = False


    for indice, archivo in enumerate(
        files
    ):

        try:

            if await request.is_disconnected():

                cancelado = True

                print(
                    f"[BATCH CANCELLED] "
                    f"lote={lote_uuid} "
                    f"indice={indice}"
                )

                break

        except Exception:

            pass


        nombre_original = (
            archivo.filename
            or
            f"archivo_{indice + 1}"
        )


        ruta_original = (

            relative_paths[indice]

            if indice < len(relative_paths)

            else nombre_original

        )


        ruta_relativa = (
            normalizar_ruta_relativa(
                ruta_original,
                nombre_original
            )
        )


        try:

            contenido = await archivo.read()
            tamano_archivo = len(contenido) if contenido else 0


            try:
                validar_archivo_dlp(
                    nombre_original,
                    tamano_archivo,
                )
            except HTTPException as dlp_error:
                errores.append({
                    "archivo": nombre_original,
                    "ruta_relativa": ruta_relativa,
                    "error": str(dlp_error.detail),
                })
                del contenido
                continue


            hash_sha256 = hashlib.sha256(
                contenido
            ).hexdigest()


            nombre_final = obtener_nombre_correlativo(
                nombre_original
            )


            registro = {

                "nombre_archivo":
                    nombre_final,

                "hash_sha256":
                    hash_sha256,

                "tamano_bytes":
                    tamano_archivo,

                "estado":
                    "PENDIENTE",

                "usuario_solicitante":
                    solicitante_nombre,

                "solicitante_id":
                    solicitante_id,

                "solicitante_nombre":
                    solicitante_nombre,

                "solicitante_correo":
                    solicitante_correo,

                "lote_id":
                    lote_uuid,

                "ruta_relativa":
                    ruta_relativa,

                "drive_parent_id":
                    destino_upload_id,

                "ubicacion_drive":
                    destino_upload_path,

                "destino_solicitado_id":
                    destino_upload_id,

                "destino_solicitado_ruta":
                    destino_upload_path,

            }


            respuesta_db = (

                supabase_admin

                .table(
                    "auditoria_custodia"
                )

                .insert(
                    registro
                )

                .execute()

            )


            if not respuesta_db.data:

                raise RuntimeError(
                    "Supabase no devolvió "
                    "el registro insertado."
                )


            id_auditoria = (
                respuesta_db
                .data[0]
                .get("id")
            )


            temp_path = None

            try:
                temp_path = guardar_archivo_temporal(
                    str(id_auditoria),
                    solicitante_id,
                    nombre_final,
                    contenido,
                    lote_id=lote_uuid,
                )

                temp_update = (
                    supabase_admin
                    .table("auditoria_custodia")
                    .update({
                        "temp_storage_path": temp_path,
                        "temp_storage_bucket": SUPABASE_TEMP_BUCKET,
                        "temp_storage_saved_at": ahora_iso(),
                        "temp_storage_deleted_at": None,
                    })
                    .eq("id", id_auditoria)
                    .execute()
                )

                if not temp_update.data:
                    raise RuntimeError(
                        "El temporal se guardó, pero no se pudo vincular a la auditoría."
                    )

            except Exception:
                if temp_path:
                    eliminar_archivo_temporal(temp_path)

                try:
                    (
                        supabase_admin
                        .table("auditoria_custodia")
                        .delete()
                        .eq("id", id_auditoria)
                        .execute()
                    )
                except Exception as cleanup_error:
                    print("[TEMP STORAGE BATCH ROLLBACK DB ERROR]", cleanup_error)
                raise

            finally:
                del contenido


            procesados.append({

                "id_auditoria":
                    id_auditoria,

                "nombre_archivo":
                    nombre_final,

                "ruta_relativa":
                    ruta_relativa,

                "sha256":
                    hash_sha256,

                "tamano_bytes":
                    tamano_archivo,

                "drive_parent_id":
                    destino_upload_id,

                "ubicacion_drive":
                    destino_upload_path,

                "temp_storage_path":
                    temp_path,

            })


        except Exception as error:

            print(
                f"[BATCH FILE ERROR] "
                f"{nombre_original}: {error}"
            )


            errores.append({

                "archivo":
                    nombre_original,

                "ruta_relativa":
                    ruta_relativa,

                "error":
                    str(error),

            })


    if procesados:
        registrar_evento_auditoria(
            evento="SOLICITUD_SUBIDA_LOTE",
            categoria="CUSTODIA",
            actor=usuario_auth,
            accion="SUBIR_LOTE",
            objeto_tipo="LOTE",
            objeto_id=lote_uuid,
            estado=ESTADO_PENDIENTE,
            detalle={
                "total_registrados": len(procesados),
                "total_errores": len(errores),
                "destino_solicitado_id": destino_upload_id,
                "destino_solicitado_ruta": destino_upload_path,
            },
            request=request,
        )

        background_tasks.add_task(
            notificar_menu_pendientes_background,
            base_url_actual,
            None,
            lote_uuid,
        )


    if not procesados and errores:

        raise HTTPException(

            status_code=400,

            detail={

                "mensaje":
                    "Ningún archivo del lote "
                    "pudo registrarse.",

                "lote_id":
                    lote_uuid,

                "errores":
                    errores,

            }

        )


    estado = (

        "CANCELADO"

        if cancelado

        else (

            "PARCIAL"

            if errores

            else "REGISTRADO"

        )

    )


    return {

        "status":
            estado,

        "mensaje":
            (
                "Lote cancelado parcialmente."

                if cancelado

                else (

                    "Lote registrado "
                    "con algunas observaciones."

                    if errores

                    else
                    "Lote registrado correctamente."

                )
            ),

        "lote_id":
            lote_uuid,

        "destino_drive": {
            "id": destino_upload_id,
            "path": destino_upload_path,
        },

        "total_recibidos":
            len(files),

        "total_registrados":
            len(procesados),

        "total_errores":
            len(errores),

        "cancelado":
            cancelado,

        "archivos":
            procesados,

        "errores":
            errores,

        "telegram_notificacion":
            (
                "en_cola"
                if procesados
                else
                "no_aplica"
            ),

        "webhook":
            estado_webhook,

    }


# ============================================================
# CREAR USUARIOS DESDE EL PANEL DEL ADMINISTRADOR
# ============================================================

@app.post("/admin/users")
async def crear_usuario_desde_web(
    request: Request
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    if usuario.get("rol") != "jefe":

        raise HTTPException(
            status_code=403,
            detail=(
                "Solo el administrador "
                "puede crear usuarios."
            )
        )


    service_key = (
        SUPABASE_SERVICE_ROLE_KEY
    )


    if not service_key:

        raise HTTPException(
            status_code=500,
            detail=(
                "Falta SUPABASE_SERVICE_ROLE_KEY "
                "en la configuración del backend."
            )
        )


    try:

        payload = await request.json()

    except Exception:

        raise HTTPException(
            status_code=400,
            detail=(
                "El cuerpo de la solicitud "
                "no es JSON válido."
            )
        )


    nombre = str(
        payload.get("nombre")
        or ""
    ).strip()


    email = str(
        payload.get("email")
        or ""
    ).strip().lower()


    rol = str(
        payload.get("rol")
        or
        "subordinado"
    ).strip().lower()


    password = str(
        payload.get("password")
        or ""
    )


    if not nombre:

        raise HTTPException(
            status_code=400,
            detail=(
                "El nombre completo "
                "es obligatorio."
            )
        )


    if not email or "@" not in email:

        raise HTTPException(
            status_code=400,
            detail=(
                "Ingresa un correo electrónico válido."
            )
        )


    if rol not in (
        "jefe",
        "subordinado"
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "El rol debe ser jefe "
                "o subordinado."
            )
        )


    if len(password) < 6:

        raise HTTPException(
            status_code=400,
            detail=(
                "La contraseña debe tener "
                "al menos 6 caracteres."
            )
        )


    if len(password) > 128:

        raise HTTPException(
            status_code=400,
            detail=(
                "La contraseña no puede "
                "superar 128 caracteres."
            )
        )


    supabase_auth_url = (
        f"{SUPABASE_URL}"
        f"/auth/v1/admin/users"
    )


    headers = {

        "apikey":
            service_key,

        "Authorization":
            f"Bearer {service_key}",

        "Content-Type":
            "application/json",

    }


    body = {

        "email":
            email,

        "password":
            password,

        "email_confirm":
            True,

        "user_metadata": {

            "full_name":
                nombre,

            "rol":
                rol,

        },

        "app_metadata": {

            "rol":
                rol,

        },

    }


    try:

        respuesta = HTTP_CLIENT.post(

            supabase_auth_url,

            headers=headers,

            json=body,

            timeout=20,

        )


    except httpx.HTTPError as error:

        print(
            "[SUPABASE CREATE USER ERROR]",
            error
        )

        raise HTTPException(
            status_code=503,
            detail=(
                "No se pudo conectar "
                "con Supabase Auth."
            )
        )


    if respuesta.status_code not in (
        200,
        201
    ):

        try:

            detalle = respuesta.json()

        except ValueError:

            detalle = {}


        mensaje = str(
            detalle.get("msg")
            or
            detalle.get("message")
            or
            detalle.get("error_description")
            or
            ""
        )


        mensaje_lower = mensaje.lower()


        if (
            "already" in mensaje_lower
            or
            "exist" in mensaje_lower
            or
            "duplicate" in mensaje_lower
        ):

            raise HTTPException(
                status_code=409,
                detail=(
                    "Ya existe un usuario "
                    "con ese correo."
                )
            )


        print(
            "[SUPABASE CREATE USER REJECTED]",
            respuesta.status_code,
            respuesta.text[:500]
        )


        raise HTTPException(
            status_code=502,
            detail=(
                "Supabase no permitió crear el usuario."
            )
        )


    creado = (
        respuesta.json()
        or {}
    )


    return {

        "status":
            "ok",

        "mensaje":
            "Usuario creado correctamente.",

        "usuario": {

            "id":
                creado.get("id"),

            "email":
                creado.get("email")
                or email,

            "nombre":
                nombre,

            "rol":
                rol,

        },

    }


@app.get("/admin/users")
async def listar_usuarios_desde_web(
    request: Request
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    if usuario.get("rol") != "jefe":

        raise HTTPException(
            status_code=403,
            detail=(
                "Solo el administrador "
                "puede consultar los usuarios."
            )
        )


    service_key = (
        SUPABASE_SERVICE_ROLE_KEY
    )


    if not service_key:

        raise HTTPException(
            status_code=500,
            detail=(
                "Falta SUPABASE_SERVICE_ROLE_KEY "
                "en la configuración del backend."
            )
        )


    headers = {

        "apikey":
            service_key,

        "Authorization":
            f"Bearer {service_key}",

    }


    try:

        respuesta = HTTP_CLIENT.get(

            f"{SUPABASE_URL}"
            f"/auth/v1/admin/users",

            headers=headers,

            params={
                "page": 1,
                "per_page": 1000
            },

            timeout=20,

        )


    except httpx.HTTPError as error:

        print(
            "[SUPABASE LIST USERS ERROR]",
            error
        )

        raise HTTPException(
            status_code=503,
            detail=(
                "No se pudo conectar "
                "con Supabase Auth."
            )
        )


    if respuesta.status_code != 200:

        print(
            "[SUPABASE LIST USERS REJECTED]",
            respuesta.status_code,
            respuesta.text[:500]
        )

        raise HTTPException(
            status_code=502,
            detail=(
                "Supabase no permitió "
                "consultar los usuarios."
            )
        )


    try:

        data = (
            respuesta.json()
            or {}
        )

    except ValueError:

        raise HTTPException(
            status_code=502,
            detail=(
                "Supabase devolvió "
                "una respuesta inválida."
            )
        )


    usuarios = []


    for item in data.get(
        "users",
        []
    ):

        metadata = (
            item.get("user_metadata")
            or {}
        )

        app_metadata = (
            item.get("app_metadata")
            or {}
        )


        usuarios.append({

            "id":
                item.get("id"),

            "email":
                item.get("email")
                or "",

            "nombre":
                (
                    metadata.get("full_name")
                    or
                    metadata.get("name")
                    or
                    item.get("email")
                    or
                    "Usuario"
                ),

            "rol":
                (
                    app_metadata.get("rol")
                    or
                    metadata.get("rol")
                    or
                    "subordinado"
                ),

            "created_at":
                item.get("created_at"),

            "last_sign_in_at":
                item.get("last_sign_in_at"),

        })


    return {
        "status": "ok",
        "usuarios": usuarios
    }


@app.delete("/admin/users/{user_id}")
async def eliminar_usuario_desde_web(
    user_id: str,
    request: Request
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    if usuario.get("rol") != "jefe":

        raise HTTPException(
            status_code=403,
            detail=(
                "Solo el administrador "
                "puede eliminar usuarios."
            )
        )


    try:

        target_id = str(
            UUID(
                user_id.strip()
            )
        )

    except (
        ValueError,
        AttributeError
    ):

        raise HTTPException(
            status_code=400,
            detail="ID de usuario inválido."
        )


    current_id = str(
        usuario.get("id")
        or
        usuario.get("sub")
        or
        ""
    ).strip()


    if target_id == current_id:

        raise HTTPException(
            status_code=400,
            detail=(
                "No puedes eliminar "
                "tu propia cuenta."
            )
        )


    service_key = (
        SUPABASE_SERVICE_ROLE_KEY
    )


    if not service_key:

        raise HTTPException(
            status_code=500,
            detail=(
                "Falta SUPABASE_SERVICE_ROLE_KEY "
                "en la configuración del backend."
            )
        )


    headers = {

        "apikey":
            service_key,

        "Authorization":
            f"Bearer {service_key}",

    }


    try:

        respuesta = requests.delete(

            f"{SUPABASE_URL}"
            f"/auth/v1/admin/users/"
            f"{target_id}",

            headers=headers,

            timeout=20,

        )


    except httpx.HTTPError as error:

        print(
            "[SUPABASE DELETE USER ERROR]",
            error
        )

        raise HTTPException(
            status_code=503,
            detail=(
                "No se pudo conectar "
                "con Supabase Auth."
            )
        )


    if respuesta.status_code not in (
        200,
        204
    ):

        try:

            detalle = respuesta.json()

        except ValueError:

            detalle = {}


        mensaje = str(
            detalle.get("msg")
            or
            detalle.get("message")
            or
            detalle.get("error_description")
            or
            ""
        )


        print(
            "[SUPABASE DELETE USER REJECTED]",
            respuesta.status_code,
            respuesta.text[:500]
        )


        raise HTTPException(
            status_code=502,
            detail=(
                mensaje
                or
                "Supabase no permitió "
                "eliminar el usuario."
            )
        )


    return {

        "status":
            "ok",

        "mensaje":
            "Usuario eliminado correctamente.",

        "usuario_id":
            target_id

    }


# ============================================================
# DECISIONES DE CUSTODIA (WEB + TELEGRAM)
# ============================================================

def resolver_custodia_archivo(
    auditoria_id: str,
    aprobar: bool
):

    auditoria_id = str(
        UUID(
            str(
                auditoria_id
            ).strip()
        )
    )


    with DECISION_LOCK:

        consulta = (
            supabase_admin
            .table(
                "auditoria_custodia"
            )
            .select(
                "id,estado,nombre_archivo,hash_sha256,"
                "solicitante_id,solicitante_nombre,"
                "solicitante_correo,drive_parent_id,"
                "ubicacion_drive,temp_storage_path,"
                "temp_storage_bucket"
            )
            .eq(
                "id",
                auditoria_id
            )
            .limit(1)
            .execute()
        )


        if not consulta.data:

            raise ValueError(
                f"No existe la auditoría "
                f"{auditoria_id}."
            )


        registro = (
            consulta.data[0]
        )


        estado_actual = str(
            registro.get("estado")
            or ""
        ).upper()


        if estado_actual != "PENDIENTE":

            return {

                "procesado":
                    False,

                "estado":
                    (
                        estado_actual
                        or
                        "DESCONOCIDO"
                    ),

                "registro":
                    registro,

                "drive_id":
                    registro.get(
                        "drive_file_id"
                    ),

                "resultado_update":
                    None,

            }


        drive_id = None
        resultado_update = None


        if aprobar:

            temp_path = str(
                registro.get("temp_storage_path")
                or ""
            ).strip()

            if not temp_path:
                raise RuntimeError(
                    "El archivo pendiente no tiene almacenamiento temporal persistente. "
                    "No se modificó el estado en Supabase."
                )

            contenido_temporal = leer_archivo_temporal(
                temp_path,
                bucket=(
                    registro.get("temp_storage_bucket")
                    or SUPABASE_TEMP_BUCKET
                ),
            )

            drive_id = subir_a_google_drive(

                registro.get("nombre_archivo")
                or "archivo",

                contenido_temporal,

                auditoria_id=
                    auditoria_id,

                parent_id=(
                    registro.get("drive_parent_id")
                    or GOOGLE_FOLDER_ID
                ),

            )

            del contenido_temporal


            if not drive_id:

                raise RuntimeError(
                    "Google Drive no pudo completar "
                    "la transferencia. "
                    "El documento sigue PENDIENTE."
                )


            resultado_update = (

                supabase_admin

                .table(
                    "auditoria_custodia"
                )

                .update({

                    "estado":
                        "APROBADO",

                    "drive_file_id":
                        drive_id,

                    "drive_parent_id":
                        (
                            registro.get("drive_parent_id")
                            or GOOGLE_FOLDER_ID
                        ),

                    "ubicacion_drive":
                        (
                            registro.get("ubicacion_drive")
                            or "DRIVE PROYECTO"
                        ),

                    "destino_final_id":
                        (
                            registro.get("drive_parent_id")
                            or GOOGLE_FOLDER_ID
                            or GOOGLE_DRIVE_ROOT_ID
                        ),

                    "destino_final_ruta":
                        (
                            registro.get("ubicacion_drive")
                            or "DRIVE PROYECTO"
                        ),

                    "en_drive":
                        True,

                    "estado_archivo":
                        "ACTIVO",

                    "fecha_transferencia":
                        ahora_iso(),

                    "fecha_ultima_operacion":
                        ahora_iso(),

                })

                .eq(
                    "id",
                    auditoria_id
                )

                .eq(
                    "estado",
                    "PENDIENTE"
                )

                .execute()

            )


            if not resultado_update.data:

                try:

                    service = (
                        obtener_servicio_google_drive()
                    )


                    (
                        service
                        .files()
                        .delete(
                            fileId=drive_id,
                            supportsAllDrives=True
                        )
                        .execute()
                    )


                except Exception as rollback_error:

                    print(
                        "[DRIVE/SUPABASE FILE "
                        "ROLLBACK ERROR]",
                        rollback_error
                    )


                raise RuntimeError(
                    "El archivo llegó a Drive, "
                    "pero Supabase no pudo confirmar "
                    "el estado APROBADO. "
                    "Se intentó revertir la transferencia."
                )


            nuevo_estado = (
                "APROBADO"
            )


        else:

            resultado_update = (

                supabase_admin

                .table(
                    "auditoria_custodia"
                )

                .update({
                    "estado":
                        "RECHAZADO"
                })

                .eq(
                    "id",
                    auditoria_id
                )

                .eq(
                    "estado",
                    "PENDIENTE"
                )

                .execute()

            )


            if not resultado_update.data:

                raise RuntimeError(
                    "No se pudo cambiar "
                    "el documento a RECHAZADO."
                )


            nuevo_estado = (
                "RECHAZADO"
            )


        # El temporal solo se elimina cuando la decisión quedó confirmada.
        # Si Drive falla, la solicitud continúa PENDIENTE y el archivo permanece
        # disponible para volver a intentar.
        limpiar_temporal_de_registro(registro)


        return {

            "procesado":
                True,

            "estado":
                nuevo_estado,

            "registro":
                registro,

            "drive_id":
                drive_id,

            "resultado_update":
                resultado_update,

        }


def resolver_custodia_carpeta(
    lote_id: str,
    aprobar: bool
):

    lote_id = str(
        UUID(
            str(
                lote_id
            ).strip()
        )
    )


    with DECISION_LOCK:

        consulta_lote = (

            supabase_admin

            .table(
                "auditoria_custodia"
            )

            .select(
                "id,nombre_archivo,hash_sha256,"
                "tamano_bytes,estado,solicitante_id,"
                "solicitante_nombre,solicitante_correo,"
                "lote_id,ruta_relativa,drive_parent_id,"
                "ubicacion_drive,temp_storage_path,temp_storage_bucket"
            )

            .eq(
                "lote_id",
                lote_id
            )

            .eq(
                "estado",
                "PENDIENTE"
            )

            .execute()

        )


        documentos = (
            consulta_lote.data
            or []
        )


        if not documentos:

            return {

                "procesado":
                    False,

                "estado":
                    "SIN_PENDIENTES",

                "documentos":
                    [],

                "drive_lote":
                    None,

                "resultado_update":
                    None,

            }


        nombre_carpeta = (
            obtener_carpeta_desde_ruta(
                documentos[0].get(
                    "ruta_relativa"
                )
            )
        )


        if not nombre_carpeta:

            raise RuntimeError(
                "El lote no corresponde "
                "a una carpeta completa."
            )


        drive_lote = None
        resultado_update = None


        if aprobar:

            faltantes_temporales = [
                str(doc.get("id"))
                for doc in documentos
                if not str(doc.get("temp_storage_path") or "").strip()
            ]

            if faltantes_temporales:
                raise RuntimeError(
                    f"{len(faltantes_temporales)} archivo(s) no tienen "
                    "almacenamiento temporal persistente. "
                    "La carpeta no fue transferida."
                )


            drive_lote = (
                subir_lote_carpeta_a_drive(
                    documentos,
                    parent_id=(
                        documentos[0].get("drive_parent_id")
                        or GOOGLE_FOLDER_ID
                    )
                )
            )


            if not drive_lote:

                raise RuntimeError(
                    "Google Drive no pudo completar "
                    "toda la carpeta. "
                    "Los documentos siguen PENDIENTES."
                )


            resultado_update = (

                supabase_admin

                .table(
                    "auditoria_custodia"
                )

                .update({

                    "estado":
                        "APROBADO",

                    "drive_folder_id":
                        drive_lote[
                            "folder_id"
                        ],

                    "drive_parent_id":
                        (
                            documentos[0].get("drive_parent_id")
                            or GOOGLE_FOLDER_ID
                        ),

                    "ubicacion_drive":
                        (
                            documentos[0].get("ubicacion_drive")
                            or "DRIVE PROYECTO"
                        ),

                    "destino_final_id":
                        (
                            documentos[0].get("drive_parent_id")
                            or GOOGLE_FOLDER_ID
                            or GOOGLE_DRIVE_ROOT_ID
                        ),

                    "destino_final_ruta":
                        (
                            documentos[0].get("ubicacion_drive")
                            or "DRIVE PROYECTO"
                        ),

                    "en_drive":
                        True,

                    "estado_archivo":
                        "ACTIVO",

                    "fecha_transferencia":
                        ahora_iso(),

                    "fecha_ultima_operacion":
                        ahora_iso(),

                })

                .eq(
                    "lote_id",
                    lote_id
                )

                .eq(
                    "estado",
                    "PENDIENTE"
                )

                .execute()

            )


            if not resultado_update.data:

                try:

                    service = (
                        obtener_servicio_google_drive()
                    )


                    (
                        service
                        .files()
                        .delete(
                            fileId=
                                drive_lote[
                                    "folder_id"
                                ],
                            supportsAllDrives=True,
                        )
                        .execute()
                    )


                except Exception as rollback_error:

                    print(
                        "[DRIVE/SUPABASE "
                        "ROLLBACK ERROR]",
                        rollback_error
                    )


                raise RuntimeError(
                    "La carpeta llegó a Drive, "
                    "pero Supabase no pudo "
                    "confirmar APROBADO."
                )


            nuevo_estado = (
                "APROBADO"
            )


        else:

            resultado_update = (

                supabase_admin

                .table(
                    "auditoria_custodia"
                )

                .update({
                    "estado":
                        "RECHAZADO"
                })

                .eq(
                    "lote_id",
                    lote_id
                )

                .eq(
                    "estado",
                    "PENDIENTE"
                )

                .execute()

            )


            if not resultado_update.data:

                raise RuntimeError(
                    "No se pudo cambiar "
                    "la carpeta a RECHAZADO."
                )


            nuevo_estado = (
                "RECHAZADO"
            )


        for doc in documentos:
            limpiar_temporal_de_registro(doc)


        return {

            "procesado":
                True,

            "estado":
                nuevo_estado,

            "documentos":
                documentos,

            "drive_lote":
                drive_lote,

            "resultado_update":
                resultado_update,

        }


def resolver_payload_custodia_web(
    payload: dict,
    aprobar: bool
):

    objeto_tipo = str(
        payload.get(
            "objeto_tipo"
        )
        or ""
    ).upper().strip()


    if objeto_tipo not in (
        "ARCHIVO",
        "CARPETA"
    ):

        raise HTTPException(
            status_code=400,
            detail="objeto_tipo inválido."
        )


    try:

        if objeto_tipo == "CARPETA":

            lote_id = str(
                UUID(
                    str(
                        payload.get(
                            "lote_id"
                        )
                        or ""
                    ).strip()
                )
            )


            resultado = (
                resolver_custodia_carpeta(
                    lote_id,
                    aprobar
                )
            )


            if not resultado[
                "procesado"
            ]:

                raise HTTPException(
                    status_code=409,
                    detail=(
                        "La carpeta ya fue procesada "
                        "o no tiene archivos pendientes."
                    ),
                )


            return {

                "status":
                    "ok",

                "objeto_tipo":
                    "CARPETA",

                "lote_id":
                    lote_id,

                "estado":
                    resultado[
                        "estado"
                    ],

                "procesados":
                    len(
                        resultado[
                            "documentos"
                        ]
                    ),

                "drive_folder_id":
                    (
                        resultado[
                            "drive_lote"
                        ][
                            "folder_id"
                        ]

                        if resultado[
                            "drive_lote"
                        ]

                        else None
                    ),

            }


        auditoria_id = str(
            UUID(
                str(
                    payload.get(
                        "auditoria_id"
                    )
                    or ""
                ).strip()
            )
        )


        resultado = (
            resolver_custodia_archivo(
                auditoria_id,
                aprobar
            )
        )


        if not resultado[
            "procesado"
        ]:

            raise HTTPException(

                status_code=409,

                detail=(
                    "El archivo ya fue procesado. "
                    f"Estado actual: "
                    f"{resultado['estado']}."
                ),

            )


        return {

            "status":
                "ok",

            "objeto_tipo":
                "ARCHIVO",

            "auditoria_id":
                auditoria_id,

            "estado":
                resultado[
                    "estado"
                ],

            "procesados":
                1,

            "drive_file_id":
                resultado.get(
                    "drive_id"
                ),

        }


    except HTTPException:

        raise


    except (
        ValueError,
        TypeError,
        AttributeError
    ):

        campo = (
            "lote_id"

            if objeto_tipo
            == "CARPETA"

            else
            "auditoria_id"
        )


        raise HTTPException(
            status_code=400,
            detail=f"{campo} inválido."
        )


    except RuntimeError as error:

        raise HTTPException(
            status_code=409,
            detail=str(error)
        )


# ============================================================
# OPERACIONES SOLICITADAS DESDE LA WEB
# ============================================================

@app.post("/custody/decision")
async def decidir_custodia_desde_web(
    request: Request,
    background_tasks: BackgroundTasks
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    if usuario.get("rol") != "jefe":

        raise HTTPException(

            status_code=403,

            detail=(
                "Solo el administrador puede "
                "aprobar o rechazar cargas en cola."
            ),

        )


    payload = (
        await request.json()
    )


    decision = str(
        payload.get(
            "decision"
        )
        or ""
    ).upper().strip()


    if decision not in (
        "APROBAR",
        "RECHAZAR"
    ):

        raise HTTPException(
            status_code=400,
            detail="decision inválida."
        )


    resultado = (
        resolver_payload_custodia_web(
            payload,
            aprobar=(
                decision
                == "APROBAR"
            )
        )
    )

    background_tasks.add_task(
        notificar_resultado_custodia_web_telegram,
        payload,
        resultado,
        usuario,
    )

    return resultado


@app.post("/custody/cancel")
async def cancelar_custodia_desde_web(
    request: Request,
    background_tasks: BackgroundTasks
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    if usuario.get("rol") != "jefe":

        raise HTTPException(

            status_code=403,

            detail=(
                "Solo el administrador puede "
                "rechazar cargas en cola."
            ),

        )


    payload = (
        await request.json()
    )


    resultado = resolver_payload_custodia_web(
        payload,
        aprobar=False
    )

    background_tasks.add_task(
        notificar_resultado_custodia_web_telegram,
        payload,
        resultado,
        usuario,
    )

    return resultado



# ============================================================
# REPORTES EXCEL XLSX (PUNTOS 25-26)
# ============================================================

def _safe_select_tabla(tabla: str, columnas: str = "*") -> list:
    try:
        respuesta = (
            supabase_admin
            .table(tabla)
            .select(columnas)
            .execute()
        )
        return respuesta.data or []
    except Exception as error:
        print(f"[REPORT TABLE ERROR] {tabla}", error)
        return []


def _valor_excel(valor):
    if valor is None:
        return ""
    if isinstance(valor, (dict, list, tuple, set)):
        return json.dumps(
            valor,
            ensure_ascii=False,
            default=str,
        )
    return valor


def _formatear_hoja_excel(ws):
    if ws.max_row < 1:
        return

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )
    borde = Border(
        bottom=Side(
            style="thin",
            color="B7B7B7",
        )
    )

    for celda in ws[1]:
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        celda.border = borde

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for columna in ws.columns:
        letra = get_column_letter(columna[0].column)
        max_len = 0

        for celda in columna:
            texto = str(celda.value or "")
            max_len = max(
                max_len,
                min(len(texto), 60),
            )

            if isinstance(celda.value, datetime):
                celda.number_format = "yyyy-mm-dd hh:mm:ss"

            celda.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        ws.column_dimensions[letra].width = max(
            12,
            min(max_len + 2, 55),
        )


def _crear_hoja_datos(
    wb: Workbook,
    titulo: str,
    columnas: list,
    filas: list,
):
    ws = wb.create_sheet(title=titulo[:31])
    ws.append([col[0] for col in columnas])

    for fila in filas:
        ws.append([
            _valor_excel(fila.get(col[1]))
            for col in columnas
        ])

    _formatear_hoja_excel(ws)
    return ws


def generar_reporte_excel_datavault() -> bytes:
    auditoria = _safe_select_tabla("auditoria_custodia")
    operaciones = _safe_select_tabla("solicitudes_operacion")
    eliminados = _safe_select_tabla("elementos_eliminados")
    eventos = _safe_select_tabla("auditoria_eventos_v2")
    dlp = _safe_select_tabla("dlp_auditoria")

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    total = len(auditoria)
    aprobados = sum(
        1 for x in auditoria
        if normalizar_estado(x.get("estado")) == ESTADO_APROBADO
    )
    pendientes = sum(
        1 for x in auditoria
        if normalizar_estado(x.get("estado")) == ESTADO_PENDIENTE
    )
    rechazados = sum(
        1 for x in auditoria
        if normalizar_estado(x.get("estado")) == ESTADO_RECHAZADO
    )

    ws.append(["INDICADOR", "VALOR"])
    ws.append(["Generado UTC", ahora_iso()])
    ws.append(["Documentos auditados", total])
    ws.append(["Aprobados", aprobados])
    ws.append(["Pendientes", pendientes])
    ws.append(["Rechazados", rechazados])
    ws.append(["Operaciones registradas", len(operaciones)])
    ws.append(["Elementos eliminados", len(eliminados)])
    ws.append(["Eventos de auditoría", len(eventos)])
    ws.append(["Cambios DLP", len(dlp)])
    _formatear_hoja_excel(ws)

    _crear_hoja_datos(
        wb,
        "Custodia",
        [
            ("ID", "id"),
            ("Archivo", "nombre_archivo"),
            ("Usuario", "solicitante_nombre"),
            ("Correo", "solicitante_correo"),
            ("Estado", "estado"),
            ("Estado archivo", "estado_archivo"),
            ("Tamaño bytes", "tamano_bytes"),
            ("SHA256", "hash_sha256"),
            ("Ruta relativa", "ruta_relativa"),
            ("Destino solicitado", "destino_solicitado_ruta"),
            ("Destino final", "destino_final_ruta"),
            ("Ubicación Drive", "ubicacion_drive"),
            ("Fecha solicitud", "fecha_solicitud"),
            ("Fecha transferencia", "fecha_transferencia"),
            ("Fecha eliminación", "fecha_eliminacion"),
        ],
        auditoria,
    )

    _crear_hoja_datos(
        wb,
        "Operaciones",
        [
            ("ID", "id"),
            ("Tipo", "tipo_operacion"),
            ("Objeto", "objeto_tipo"),
            ("Nombre", "nombre_objeto"),
            ("Usuario", "solicitante_nombre"),
            ("Origen", "carpeta_origen"),
            ("Destino solicitado", "destino_solicitado_ruta"),
            ("Destino final", "destino_final_ruta"),
            ("Estado", "estado"),
            ("Fecha solicitud", "fecha_solicitud"),
            ("Fecha resolución", "fecha_resolucion"),
            ("Resultado", "resultado"),
        ],
        operaciones,
    )

    _crear_hoja_datos(
        wb,
        "Eliminaciones",
        [
            ("ID", "id"),
            ("Tipo", "objeto_tipo"),
            ("Nombre", "nombre_objeto"),
            ("Usuario", "solicitante_nombre"),
            ("Correo", "solicitante_correo"),
            ("Ruta anterior", "ubicacion_anterior"),
            ("Estado", "estado"),
            ("Fecha eliminación", "fecha_eliminacion"),
            ("Detalle", "detalle"),
        ],
        eliminados,
    )

    _crear_hoja_datos(
        wb,
        "Auditoria",
        [
            ("Evento", "evento"),
            ("Categoría", "categoria"),
            ("Acción", "accion"),
            ("Actor", "actor_nombre"),
            ("Correo", "actor_correo"),
            ("Rol", "actor_rol"),
            ("Objeto", "objeto_tipo"),
            ("Objeto ID", "objeto_id"),
            ("Estado", "estado"),
            ("IP", "ip"),
            ("Fecha", "created_at"),
            ("Detalle", "detalle"),
        ],
        eventos,
    )

    _crear_hoja_datos(
        wb,
        "Cambios DLP",
        [
            ("Política", "politica_codigo"),
            ("Acción", "accion"),
            ("Valor anterior", "valor_anterior"),
            ("Valor nuevo", "valor_nuevo"),
            ("Actor", "actor_nombre"),
            ("Correo", "actor_correo"),
            ("Fecha", "fecha"),
        ],
        dlp,
    )

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida.getvalue()


@app.get("/reports/audit.xlsx")
def descargar_reporte_excel(
    request: Request,
):
    admin = obtener_admin_desde_request(request)

    contenido = generar_reporte_excel_datavault()

    registrar_evento_auditoria(
        evento="REPORTE_EXCEL_GENERADO",
        categoria="REPORTES",
        actor=admin,
        accion="EXPORTAR_XLSX",
        objeto_tipo="REPORTE",
        objeto_id=None,
        estado=ESTADO_APROBADO,
        detalle={
            "formato": "xlsx",
            "bytes": len(contenido),
        },
        request=request,
    )

    nombre = (
        "datavault_reporte_"
        + datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    return StreamingResponse(
        io.BytesIO(contenido),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{nombre}"'
        },
    )


@app.get("/drive/browse")
def navegar_drive(
    request: Request,
    folder_id: str = "root",
    include_files: bool = True,
):
    obtener_usuario_supabase_desde_request(request)

    if not google_drive_configurado():
        raise error_api(
            503,
            "DRIVE_NOT_CONFIGURED",
            "Google Drive no está configurado.",
        )

    try:
        return {
            "status": "ok",
            "root": obtener_raiz_drive_autorizada(),
            **listar_hijos_drive(
                folder_id,
                include_files=bool(include_files),
            ),
            "cache_ttl_seconds": DRIVE_CACHE_TTL_SECONDS,
        }
    except HTTPException:
        raise
    except Exception as error:
        print("[DRIVE BROWSE ERROR]", error)
        raise error_api(
            502,
            "DRIVE_BROWSE_FAILED",
            "No se pudo consultar la carpeta de Google Drive.",
        )


@app.get("/drive/breadcrumb")
def breadcrumb_drive(
    request: Request,
    folder_id: str = "root",
):
    obtener_usuario_supabase_desde_request(request)

    try:
        info = resolver_ruta_drive(folder_id)
        return {
            "status": "ok",
            "folder": info,
            "breadcrumb": info.get("breadcrumb") or [],
        }
    except HTTPException:
        raise
    except Exception as error:
        print("[DRIVE BREADCRUMB ERROR]", error)
        raise error_api(
            502,
            "DRIVE_BREADCRUMB_FAILED",
            "No se pudo obtener la ruta de Google Drive.",
        )


@app.post("/drive/cache/invalidate")
def invalidar_cache_drive_api(
    request: Request,
):
    obtener_admin_desde_request(request)
    invalidar_cache_drive()
    return {
        "status": "ok",
        "message": "Caché de navegación de Drive invalidada.",
    }


@app.get("/drive-folders")
def obtener_carpetas_drive(
    request: Request
):
    """Compatibilidad con el selector antiguo del frontend."""
    obtener_usuario_supabase_desde_request(request)

    if not google_drive_configurado():
        raise error_api(
            503,
            "DRIVE_NOT_CONFIGURED",
            "Google Drive no está configurado.",
        )

    raiz = obtener_raiz_drive_autorizada()
    carpetas = listar_carpetas_drive_recursivas()

    return {
        "folders": [
            {
                "id": raiz.get("id"),
                "name": raiz.get("name"),
                "parent_id": None,
                "path": raiz.get("path"),
                "depth": 0,
                "ancestors": [],
                "root": True,
            },
            *[
                {
                    "id": c.get("id"),
                    "name": c.get("name"),
                    "parent_id": c.get("parent_id"),
                    "path": c.get("path"),
                    "depth": c.get("depth", 1),
                    "ancestors": c.get("ancestors") or [],
                    "root": False,
                }
                for c in carpetas
            ],
        ]
    }


# ============================================================
# API PAPELERA / HISTORIAL (PUNTOS 18-21)
# ============================================================

@app.get("/trash")
def listar_papelera(
    request: Request,
):
    """Admin ve todo; subordinado ve únicamente sus propios elementos."""
    usuario = obtener_usuario_supabase_desde_request(request)
    return consultar_papelera_backend(usuario)


@app.get("/trash/mine")
def listar_mi_papelera(
    request: Request,
):
    usuario = obtener_usuario_supabase_desde_request(request)
    return consultar_papelera_backend(
        usuario,
        solo_usuario=True,
    )


@app.get("/admin/trash")
def listar_papelera_admin(
    request: Request,
):
    admin = obtener_admin_desde_request(request)
    return consultar_papelera_backend(admin)


@app.get("/trash/history")
def listar_historial_papelera(
    request: Request,
):
    usuario = obtener_usuario_supabase_desde_request(request)
    return consultar_historial_eliminaciones(usuario)


@app.get("/admin/trash/history")
def listar_historial_papelera_admin(
    request: Request,
):
    admin = obtener_admin_desde_request(request)
    return consultar_historial_eliminaciones(admin)


@app.post("/operations/request")
async def solicitar_operacion(
    request: Request,
    background_tasks: BackgroundTasks
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    payload = (
        await request.json()
    )


    tipo = str(
        payload.get(
            "tipo_operacion"
        )
        or ""
    ).upper().strip()


    objeto_tipo = str(
        payload.get(
            "objeto_tipo"
        )
        or ""
    ).upper().strip()


    if tipo not in (
        "MOVER",
        "ELIMINAR"
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "tipo_operacion inválido."
            )
        )


    # Las operaciones sensibles mantienen aprobación obligatoria desde backend.
    politica_aprobacion = (
        "APROBACION_MOVIMIENTO"
        if tipo == "MOVER"
        else "APROBACION_ELIMINACION"
    )

    if not politica_dlp_activa(politica_aprobacion):
        # Defensa en profundidad: estas políticas son obligatorias y no deben
        # quedar desactivadas ni por una modificación directa en la BD.
        raise HTTPException(
            status_code=503,
            detail=(
                f"La política obligatoria {politica_aprobacion} "
                "no está disponible."
            ),
        )


    if objeto_tipo not in (
        "ARCHIVO",
        "CARPETA"
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "objeto_tipo inválido."
            )
        )


    objeto = obtener_objeto_operable(

        usuario["id"],

        objeto_tipo,

        auditoria_id=
            payload.get(
                "auditoria_id"
            ),

        lote_id=
            payload.get(
                "lote_id"
            ),

    )


    destino = None


    if tipo == "MOVER":

        destino = validar_destino_drive(
            payload.get(
                "carpeta_destino_id"
            )
        )


        if (
            destino["id"]
            ==
            objeto["drive_parent_id"]
        ):

            raise HTTPException(
                status_code=400,
                detail=(
                    "El elemento ya se encuentra "
                    "en esa carpeta."
                )
            )


        if objeto["tipo"] == "CARPETA":

            if (
                destino["id"]
                ==
                objeto["drive_id"]
            ):

                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Una carpeta no puede moverse "
                        "dentro de sí misma."
                    )
                )


            if (
                objeto["drive_id"]
                in
                (
                    destino.get(
                        "ancestors"
                    )
                    or []
                )
            ):

                raise HTTPException(

                    status_code=400,

                    detail=(
                        "Una carpeta no puede moverse "
                        "dentro de una de sus subcarpetas."
                    ),

                )


    pendientes = (

        supabase_admin

        .table("solicitudes_operacion")

        .select(
            "id,tipo_operacion,objeto_tipo,"
            "auditoria_id,lote_id,estado"
        )

        .eq(
            "solicitante_id",
            usuario["id"]
        )

        .eq(
            "estado",
            "PENDIENTE"
        )

        .execute()

    ).data or []


    for item in pendientes:

        mismo = (

            (
                objeto_tipo
                == "CARPETA"

                and

                str(
                    item.get(
                        "lote_id"
                    )
                    or ""
                )
                ==
                str(
                    objeto.get(
                        "lote_id"
                    )
                    or ""
                )
            )

            or

            (
                objeto_tipo
                == "ARCHIVO"

                and

                str(
                    item.get(
                        "auditoria_id"
                    )
                    or ""
                )
                ==
                str(
                    objeto.get(
                        "auditoria_id"
                    )
                    or ""
                )
            )

        )


        if mismo:

            raise HTTPException(

                status_code=409,

                detail=(
                    "Ya existe una operación "
                    "pendiente para este elemento."
                ),

            )


    registro = {

        "tipo_operacion":
            tipo,

        "objeto_tipo":
            objeto_tipo,

        "auditoria_id":
            objeto.get(
                "auditoria_id"
            ),

        "lote_id":
            objeto.get(
                "lote_id"
            ),

        "solicitante_id":
            usuario["id"],

        "solicitante_nombre":
            usuario["nombre"],

        "solicitante_correo":
            usuario["correo"],

        "nombre_objeto":
            objeto["nombre"],

        "carpeta_origen":
            (
                objeto.get(
                    "ubicacion"
                )
                or
                "DRIVE PROYECTO"
            ),

        "carpeta_destino_id":
            (
                destino["id"]
                if destino
                else None
            ),

        "carpeta_destino_nombre":
            (
                (
                    destino.get(
                        "path"
                    )
                    or
                    destino["name"]
                )
                if destino
                else None
            ),

        "destino_solicitado_id":
            (
                destino["id"]
                if destino
                else None
            ),

        "destino_solicitado_ruta":
            (
                (
                    destino.get("path")
                    or destino.get("name")
                )
                if destino
                else None
            ),

        "estado":
            ESTADO_PENDIENTE,

    }


    respuesta = (

        supabase_admin

        .table("solicitudes_operacion")

        .insert(
            registro
        )

        .execute()

    )


    if not respuesta.data:

        raise error_api(
            500,
            "OPERATION_REQUEST_NOT_SAVED",
            "No se pudo registrar la solicitud de operación.",
        )

    registrar_evento_auditoria(
        evento="SOLICITUD_OPERACION",
        categoria="OPERACIONES",
        actor=usuario,
        accion=tipo,
        objeto_tipo=objeto_tipo,
        objeto_id=str(respuesta.data[0].get("id") or ""),
        estado=ESTADO_PENDIENTE,
        detalle={
            "nombre_objeto": objeto.get("nombre"),
            "origen": objeto.get("ubicacion"),
            "destino_solicitado_id": (
                destino.get("id") if destino else None
            ),
            "destino_solicitado_ruta": (
                destino.get("path") if destino else None
            ),
        },
        request=request,
    )


    solicitud = (
        respuesta.data[0]
    )


    background_tasks.add_task(
        notificar_solicitud_operacion_telegram,
        solicitud
    )


    return {

        "status":
            "ok",

        "mensaje":
            (
                "Solicitud registrada y enviada "
                "al custodio."
            ),

        "solicitud":
            solicitud,

    }


@app.post("/operations/decision")
async def decidir_operacion_desde_web(
    request: Request,
    background_tasks: BackgroundTasks
):
    """Aprueba o rechaza MOVER/ELIMINAR desde la cuenta web del administrador.

    Telegram sigue siendo una segunda vía de resolución. La función central
    procesar_solicitud_operacion evita que una misma solicitud se ejecute dos veces.
    """

    usuario = obtener_usuario_supabase_desde_request(request)

    if usuario.get("rol") != "jefe":
        raise HTTPException(
            status_code=403,
            detail="Solo el administrador puede aprobar o rechazar operaciones."
        )

    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="El cuerpo de la solicitud no es JSON válido."
        )

    solicitud_raw = str(payload.get("solicitud_id") or "").strip()
    decision = str(payload.get("decision") or "").upper().strip()

    try:
        solicitud_id = str(UUID(solicitud_raw))
    except (ValueError, TypeError, AttributeError):
        raise HTTPException(status_code=400, detail="solicitud_id inválido.")

    if decision not in ("APROBAR", "RECHAZAR"):
        raise HTTPException(status_code=400, detail="decision inválida.")

    try:
        solicitud, cambio = procesar_solicitud_operacion(
            solicitud_id,
            decision == "APROBAR",
            f"WEB:{usuario.get('id')}",
        )
    except RuntimeError as error:
        raise HTTPException(status_code=409, detail=str(error))
    except HTTPException:
        raise
    except Exception as error:
        print("[OPERATIONS WEB DECISION ERROR]", error)
        raise HTTPException(
            status_code=500,
            detail=f"No se pudo procesar la operación: {error}"
        )

    estado = str(solicitud.get("estado") or "").upper()

    background_tasks.add_task(
        sincronizar_mensajes_telegram_operacion,
        solicitud,
        f"WEB:{usuario.get('id')}",
    )

    return {
        "status": "ok" if cambio else "already_processed",
        "cambio": bool(cambio),
        "estado": estado,
        "solicitud": solicitud,
    }


@app.get("/operations/requests")
def listar_solicitudes_operacion_web(request: Request):
    """Lista solicitudes de mover/eliminar respetando el rol autenticado.

    El administrador ve todas; un subordinado únicamente sus propias
    solicitudes. La consulta se hace en el backend para no depender de acceso
    directo del navegador a solicitudes_operacion.
    """

    usuario = obtener_usuario_supabase_desde_request(request)

    query = (
        supabase_admin
        .table("solicitudes_operacion")
        .select("*")
        .order("fecha_solicitud", desc=True)
        .limit(1000)
    )

    if usuario.get("rol") != "jefe":
        query = query.eq("solicitante_id", usuario["id"])

    try:
        respuesta = query.execute()
    except Exception as error:
        print("[OPERATIONS LIST ERROR]", error)
        raise HTTPException(
            status_code=500,
            detail="No se pudieron consultar las solicitudes de operaciones."
        )

    return {
        "status": "ok",
        "solicitudes": respuesta.data or [],
    }


@app.post("/drive/backfill-legacy")
def vincular_drive_legacy(
    request: Request
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    service = (
        obtener_servicio_google_drive()
    )


    consulta = (

        supabase_admin

        .table(
            "auditoria_custodia"
        )

        .select(
            "id,lote_id,nombre_archivo,ruta_relativa,"
            "estado,estado_archivo,drive_file_id,"
            "drive_folder_id,solicitante_id"
        )

        .eq(
            "solicitante_id",
            usuario["id"]
        )

        .eq(
            "estado",
            "APROBADO"
        )

        .execute()

    )


    filas = (
        consulta.data
        or []
    )


    vinculados = 0
    omitidos = 0


    def esc_drive(
        valor: str
    ) -> str:

        return str(
            valor
            or ""
        ).replace(
            "\\",
            "\\\\"
        ).replace(
            "'",
            "\\'"
        )


    lotes = {}
    sueltos = []


    for fila in filas:

        if (
            fila.get("lote_id")
            and
            obtener_carpeta_desde_ruta(
                fila.get(
                    "ruta_relativa"
                )
            )
        ):

            lotes.setdefault(
                str(
                    fila.get(
                        "lote_id"
                    )
                ),
                []
            ).append(
                fila
            )

        else:

            sueltos.append(
                fila
            )


    for lote_id, grupo in lotes.items():

        if str(
            grupo[0].get(
                "drive_folder_id"
            )
            or ""
        ).strip():

            continue


        nombre = (
            obtener_carpeta_desde_ruta(
                grupo[0].get(
                    "ruta_relativa"
                )
            )
        )


        q = (

            f"'{GOOGLE_FOLDER_ID}' in parents and "

            f"name='{esc_drive(nombre)}' and "

            "mimeType='application/vnd.google-apps.folder'"

        )


        encontrados = (

            service
            .files()
            .list(
                q=q,
                fields=
                    "files(id,name,trashed,parents)",
                pageSize=20,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            )
            .execute()
            .get(
                "files"
            )
            or []

        )


        if len(encontrados) != 1:

            omitidos += 1
            continue


        meta = encontrados[0]


        cambios = {

            "drive_folder_id":
                meta["id"],

            "drive_parent_id":
                GOOGLE_FOLDER_ID,

            "ubicacion_drive":
                nombre,

            "en_drive":
                not bool(
                    meta.get(
                        "trashed"
                    )
                ),

            "estado_archivo":
                (
                    "ELIMINADO_EXTERNAMENTE"

                    if meta.get(
                        "trashed"
                    )

                    else
                    "ACTIVO"
                ),

            "fecha_ultima_operacion":
                ahora_iso(),

        }


        if meta.get("trashed"):

            cambios[
                "fecha_eliminacion"
            ] = ahora_iso()


        (
            supabase_admin
            .table(
                "auditoria_custodia"
            )
            .update(
                cambios
            )
            .eq(
                "lote_id",
                lote_id
            )
            .eq(
                "solicitante_id",
                usuario["id"]
            )
            .execute()
        )


        vinculados += 1


    for fila in sueltos:

        if str(
            fila.get(
                "drive_file_id"
            )
            or ""
        ).strip():

            continue


        nombre = str(
            fila.get(
                "nombre_archivo"
            )
            or ""
        ).strip()


        if not nombre:

            omitidos += 1
            continue


        q = (
            f"'{GOOGLE_FOLDER_ID}' in parents and "
            f"name='{esc_drive(nombre)}'"
        )


        encontrados = (

            service
            .files()
            .list(
                q=q,
                fields=(
                    "files(id,name,mimeType,"
                    "trashed,parents)"
                ),
                pageSize=20,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            )
            .execute()
            .get(
                "files"
            )
            or []

        )


        encontrados = [

            x

            for x in encontrados

            if x.get(
                "mimeType"
            )
            !=
            "application/vnd.google-apps.folder"

        ]


        if len(encontrados) != 1:

            omitidos += 1
            continue


        meta = encontrados[0]


        cambios = {

            "drive_file_id":
                meta["id"],

            "drive_parent_id":
                GOOGLE_FOLDER_ID,

            "ubicacion_drive":
                "DRIVE PROYECTO",

            "en_drive":
                not bool(
                    meta.get(
                        "trashed"
                    )
                ),

            "estado_archivo":
                (
                    "ELIMINADO_EXTERNAMENTE"

                    if meta.get(
                        "trashed"
                    )

                    else
                    "ACTIVO"
                ),

            "fecha_ultima_operacion":
                ahora_iso(),

        }


        if meta.get("trashed"):

            cambios[
                "fecha_eliminacion"
            ] = ahora_iso()


        (
            supabase_admin
            .table(
                "auditoria_custodia"
            )
            .update(
                cambios
            )
            .eq(
                "id",
                fila["id"]
            )
            .eq(
                "solicitante_id",
                usuario["id"]
            )
            .execute()
        )


        vinculados += 1


    return {

        "status":
            "ok",

        "vinculados":
            vinculados,

        "omitidos":
            omitidos

    }


@app.post("/drive/reconcile")
def reconciliar_drive(
    request: Request
):

    usuario = (
        obtener_usuario_supabase_desde_request(
            request
        )
    )


    service = (
        obtener_servicio_google_drive()
    )


    es_admin = (
        str(
            usuario.get("rol")
            or ""
        ).lower()
        ==
        "jefe"
    )


    query = (

        supabase_admin

        .table(
            "auditoria_custodia"
        )

        .select(
            "id,lote_id,estado,estado_archivo,en_drive,"
            "drive_file_id,drive_folder_id,drive_parent_id,"
            "ubicacion_drive,solicitante_id"
        )

        .eq(
            "estado",
            "APROBADO"
        )

    )


    if not es_admin:

        query = query.eq(
            "solicitante_id",
            usuario["id"]
        )


    consulta = query.execute()


    filas = (
        consulta.data
        or []
    )


    cambios = 0
    revisados = 0


    carpetas = {}
    sueltos = []


    for fila in filas:

        folder_id = str(
            fila.get(
                "drive_folder_id"
            )
            or ""
        ).strip()


        file_id = str(
            fila.get(
                "drive_file_id"
            )
            or ""
        ).strip()


        if folder_id:

            carpetas.setdefault(
                folder_id,
                []
            ).append(
                fila
            )


        elif file_id:

            sueltos.append(
                fila
            )


    def comprobar(
        file_id: str
    ):

        nonlocal revisados

        revisados += 1


        try:

            return (
                obtener_archivo_drive(
                    service,
                    file_id
                ),
                None
            )


        except Exception as error:

            texto = str(error)


            if (
                "404" in texto
                or
                "File not found"
                in texto
            ):

                return (
                    None,
                    "missing"
                )


            print(
                "[DRIVE RECONCILE ERROR]",
                file_id,
                error
            )


            return (
                None,
                "error"
            )


    def update_folder(
        folder_id: str,
        values: dict
    ):

        q = (

            supabase_admin

            .table(
                "auditoria_custodia"
            )

            .update(
                values
            )

            .eq(
                "drive_folder_id",
                folder_id
            )

        )


        if not es_admin:

            q = q.eq(
                "solicitante_id",
                usuario["id"]
            )


        return q.execute()


    def update_row(
        row_id: str,
        values: dict
    ):

        q = (

            supabase_admin

            .table(
                "auditoria_custodia"
            )

            .update(
                values
            )

            .eq(
                "id",
                row_id
            )

        )


        if not es_admin:

            q = q.eq(
                "solicitante_id",
                usuario["id"]
            )


        return q.execute()


    for folder_id, grupo in carpetas.items():

        meta, err = comprobar(
            folder_id
        )


        if err == "error":

            continue


        if (
            err == "missing"
            or
            (
                meta
                and
                meta.get(
                    "trashed"
                )
            )
        ):

            update_folder(

                folder_id,

                {
                    "en_drive":
                        False,

                    "estado_archivo":
                        "ELIMINADO_EXTERNAMENTE",

                    "fecha_eliminacion":
                        ahora_iso(),

                    "fecha_ultima_operacion":
                        ahora_iso(),
                }

            )

            cambios += 1

            continue


        padres = (
            meta.get(
                "parents"
            )
            or []
        )


        parent_id = (

            str(
                padres[0]
            )

            if padres

            else ""

        )


        anterior = str(
            grupo[0].get(
                "drive_parent_id"
            )
            or ""
        )


        values = {}


        if (
            grupo[0].get(
                "en_drive"
            )
            is not True
        ):

            values[
                "en_drive"
            ] = True


        if str(
            grupo[0].get(
                "estado_archivo"
            )
            or ""
        ).upper() != "ACTIVO":

            values[
                "estado_archivo"
            ] = "ACTIVO"


        if (
            parent_id
            and
            parent_id != anterior
        ):

            parent_name = (
                "DRIVE PROYECTO"
            )


            if parent_id != GOOGLE_FOLDER_ID:

                try:

                    parent_name = (

                        obtener_archivo_drive(
                            service,
                            parent_id
                        ).get(
                            "name"
                        )

                        or

                        parent_name

                    )

                except Exception:

                    pass


            values.update({

                "drive_parent_id":
                    parent_id,

                "ubicacion_drive":
                    parent_name,

            })


        if values:

            values[
                "fecha_ultima_operacion"
            ] = ahora_iso()


            update_folder(
                folder_id,
                values
            )


            cambios += 1


    for fila in sueltos:

        file_id = str(
            fila.get(
                "drive_file_id"
            )
            or ""
        ).strip()


        meta, err = comprobar(
            file_id
        )


        if err == "error":

            continue


        if (
            err == "missing"
            or
            (
                meta
                and
                meta.get(
                    "trashed"
                )
            )
        ):

            update_row(

                fila["id"],

                {
                    "en_drive":
                        False,

                    "estado_archivo":
                        "ELIMINADO_EXTERNAMENTE",

                    "fecha_eliminacion":
                        ahora_iso(),

                    "fecha_ultima_operacion":
                        ahora_iso(),
                }

            )


            cambios += 1
            continue


        padres = (
            meta.get(
                "parents"
            )
            or []
        )


        parent_id = (

            str(
                padres[0]
            )

            if padres

            else ""

        )


        anterior = str(
            fila.get(
                "drive_parent_id"
            )
            or ""
        )


        values = {}


        if (
            fila.get(
                "en_drive"
            )
            is not True
        ):

            values[
                "en_drive"
            ] = True


        if str(
            fila.get(
                "estado_archivo"
            )
            or ""
        ).upper() != "ACTIVO":

            values[
                "estado_archivo"
            ] = "ACTIVO"


        if (
            parent_id
            and
            parent_id != anterior
        ):

            parent_name = (
                "DRIVE PROYECTO"
            )


            if parent_id != GOOGLE_FOLDER_ID:

                try:

                    parent_name = (

                        obtener_archivo_drive(
                            service,
                            parent_id
                        ).get(
                            "name"
                        )

                        or
                        parent_name

                    )

                except Exception:

                    pass


            values.update({

                "drive_parent_id":
                    parent_id,

                "ubicacion_drive":
                    parent_name,

            })


        if values:

            values[
                "fecha_ultima_operacion"
            ] = ahora_iso()


            update_row(
                fila["id"],
                values
            )


            cambios += 1


    return {

        "status":
            "ok",

        "scope":
            (
                "GLOBAL"
                if es_admin
                else
                "USUARIO"
            ),

        "revisados":
            revisados,

        "cambios":
            cambios,

    }


# ============================================================
# WEBHOOK TELEGRAM
# ============================================================

@app.post("/telegram-webhook")
async def recibir_respuesta_telegram(
    request: Request,
    background_tasks: BackgroundTasks
):

    recibido = request.headers.get(
        "x-telegram-bot-api-secret-token",
        ""
    )

    if (
        not TELEGRAM_WEBHOOK_SECRET
        or not recibido
        or not hmac.compare_digest(
            str(recibido),
            str(TELEGRAM_WEBHOOK_SECRET),
        )
    ):
        raise HTTPException(
            status_code=403,
            detail="Webhook de Telegram no autorizado.",
        )

    data = await request.json()


    print(
        "[TELEGRAM UPDATE]",
        data
    )


    # ========================================================
    # MENSAJES NORMALES / COMANDOS
    # ========================================================

    if "message" in data:

        message = data[
            "message"
        ]


        chat_id = message[
            "chat"
        ][
            "id"
        ]


        user_id = message[
            "from"
        ][
            "id"
        ]


        texto = (
            message.get(
                "text",
                ""
            )
            .strip()
        )


        if user_id not in AUTHORIZED_CHAT_IDS:

            telegram_request(

                "sendMessage",

                {

                    "chat_id":
                        chat_id,

                    "text":
                        (
                            "⛔ Acceso no autorizado.\n\n"
                            f"Tu Telegram ID es: {user_id}\n\n"
                            "Agrega este ID en Railway "
                            "en la variable AUTHORIZED_CHAT_IDS."
                        ),

                },

            )


            return {

                "status":
                    "unauthorized",

                "user_id":
                    user_id,

            }


        texto_lower = (
            texto.lower()
        )


        if (
            texto_lower.startswith(
                "/start"
            )
            or
            texto_lower.startswith(
                "/menu"
            )
        ):

            quitar_teclado_inferior(
                chat_id
            )


            mostrar_panel_principal(
                chat_id
            )


            return {
                "status":
                    "panel_principal"
            }


        if texto_lower.startswith(
            "/pendientes"
        ):

            mostrar_menu_usuarios(
                chat_id,
                forzar_nuevo=True
            )


            return {
                "status":
                    "menu_usuarios"
            }


        # Compatibilidad con la barra/ReplyKeyboard de versiones anteriores.
        # Esos botones envían texto normal, no callback_query. Forzamos un
        # mensaje nuevo para que la respuesta sea visible inmediatamente.
        if texto == "👥 Usuarios":

            mostrar_menu_usuarios(
                chat_id,
                forzar_nuevo=True
            )


            return {
                "status":
                    "menu_usuarios_legacy"
            }


        if texto == "🔄 Actualizar":

            mostrar_panel_principal(
                chat_id
            )


            return {
                "status":
                    "panel_actualizado_legacy"
            }


        if texto == "📊 Estado":

            mostrar_estado_panel(
                chat_id
            )


            return {
                "status":
                    "panel_estado_legacy"
            }


        if texto_lower.startswith(
            "/id"
        ):

            respuesta = (
                "🆔 Tu Telegram ID:\n\n"
                f"{user_id}"
            )


        elif texto_lower.startswith(
            "/estado"
        ):

            resumen = (
                obtener_resumen_panel()
            )


            respuesta = (
                "📊 ESTADO DATAVAULT\n\n"
                f"🟡 Pendientes: "
                f"{resumen['pendientes']}\n"
                f"🟢 Aprobados: "
                f"{resumen['aprobados']}\n"
                f"🔴 Rechazados: "
                f"{resumen['rechazados']}"
            )


        else:

            respuesta = (
                "🛡️ DataVault DLP activo.\n\n"
                "Comandos:\n"
                "/pendientes - Abrir documentos pendientes\n"
                "/menu - Abrir menú\n"
                "/id - Ver tu Telegram ID\n"
                "/estado - Verificar conexión"
            )


        telegram_request(

            "sendMessage",

            {
                "chat_id":
                    chat_id,

                "text":
                    respuesta,
            },

        )


        return {
            "status":
                "message_processed"
        }


    # ========================================================
    # CALLBACKS DE BOTONES
    # ========================================================

    if "callback_query" in data:

        callback = (
            data[
                "callback_query"
            ]
        )


        callback_id = callback[
            "id"
        ]


        user_id = callback[
            "from"
        ][
            "id"
        ]


        action_data = callback.get(
            "data",
            ""
        )


        message = callback.get(
            "message",
            {}
        )


        chat_id = (
            message.get(
                "chat",
                {}
            )
            .get(
                "id"
            )
        )


        message_id = (
            message.get(
                "message_id"
            )
        )


        if user_id not in AUTHORIZED_CHAT_IDS:

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,

                    "text":
                        (
                            "❌ Acceso denegado: "
                            "usuario no autorizado."
                        ),

                    "show_alert":
                        True,
                },

            )


            return {
                "status":
                    "unauthorized"
            }


        print(
            f"[TELEGRAM CALLBACK] "
            f"user={user_id} "
            f"data={action_data}"
        )


        # ----------------------------------------------------
        # ALERTA INDIVIDUAL -> REFRESCAR BANDEJA PERSISTENTE
        # ----------------------------------------------------

        if action_data == "tray:refresh":
            telegram_request(
                "answerCallbackQuery",
                {
                    "callback_query_id": callback_id,
                    "text": "Abriendo bandeja...",
                },
            )

            # Convertimos la alerta tocada en la bandeja. Antes se actualizaba
            # una bandeja persistente vieja y parecía que el botón no hacía nada.
            mostrar_menu_usuarios(
                chat_id,
                message_id
            )

            return {"status": "tray_opened"}


        # ----------------------------------------------------
        # PANEL PRINCIPAL INLINE
        # ----------------------------------------------------

        if action_data in (
            "panel:inicio",
            "panel:actualizar"
        ):

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,
                    "text": (
                        "Panel actualizado."
                        if action_data == "panel:actualizar"
                        else "Panel principal."
                    )
                },

            )


            mostrar_panel_principal(
                chat_id,
                message_id
            )


            return {
                "status":
                    "panel_principal"
            }


        if action_data == "panel:usuarios":

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,
                    "text":
                        "Abriendo solicitudes por usuario..."
                },

            )


            mostrar_menu_usuarios(
                chat_id,
                message_id
            )


            return {
                "status":
                    "menu_usuarios"
            }


        if action_data == "panel:estado":

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id
                },

            )


            mostrar_estado_panel(
                chat_id,
                message_id
            )


            return {
                "status":
                    "panel_estado"
            }


        # ----------------------------------------------------
        # NAVEGACIÓN: MENÚ DE USUARIOS
        # ----------------------------------------------------

        if action_data == "menu:usuarios":

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,
                    "text":
                        "Bandeja actualizada."
                },

            )


            mostrar_menu_usuarios(
                chat_id,
                message_id
            )


            return {
                "status":
                    "menu_usuarios"
            }


        # ----------------------------------------------------
        # NAVEGACIÓN: ARCHIVOS DE UN USUARIO
        # ----------------------------------------------------

        if action_data.startswith(
            "usr:"
        ):

            solicitante_id = (
                action_data
                .split(
                    ":",
                    1
                )[1]
                .strip()
            )


            try:

                solicitante_id = str(
                    UUID(
                        solicitante_id
                    )
                )


            except (
                ValueError,
                TypeError,
                AttributeError
            ):

                telegram_request(

                    "answerCallbackQuery",

                    {
                        "callback_query_id":
                            callback_id,

                        "text":
                            "❌ UID de usuario inválido.",

                        "show_alert":
                            True,
                    },

                )


                return {
                    "status":
                        "callback_error"
                }


            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id
                },

            )


            mostrar_archivos_usuario(

                chat_id,

                message_id,

                solicitante_id,

            )


            return {
                "status":
                    "menu_archivos"
            }


        # ----------------------------------------------------
        # NAVEGACIÓN: DETALLE DEL DOCUMENTO
        # ----------------------------------------------------

        if action_data.startswith(
            "doc:"
        ):

            auditoria_id_raw = (
                action_data
                .split(
                    ":",
                    1
                )[1]
                .strip()
            )


            try:

                auditoria_id = str(
                    UUID(
                        auditoria_id_raw
                    )
                )


            except (
                ValueError,
                TypeError,
                AttributeError
            ):

                telegram_request(

                    "answerCallbackQuery",

                    {
                        "callback_query_id":
                            callback_id,

                        "text":
                            "❌ ID de auditoría inválido.",

                        "show_alert":
                            True,
                    },

                )


                return {
                    "status":
                        "callback_error"
                }


            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id
                },

            )


            mostrar_detalle_documento(

                chat_id,

                message_id,

                auditoria_id,

            )


            return {
                "status":
                    "detalle_documento"
            }


        # ----------------------------------------------------
        # NAVEGACIÓN: DETALLE DE MOVER / ELIMINAR
        # ----------------------------------------------------
        if action_data.startswith("opdet:"):
            solicitud_raw = action_data.split(":", 1)[1].strip()
            try:
                solicitud_id = str(UUID(solicitud_raw))
            except (ValueError, TypeError, AttributeError):
                return {
                    "method": "answerCallbackQuery",
                    "callback_query_id": callback_id,
                    "text": "❌ ID de solicitud inválido.",
                    "show_alert": True,
                }

            # Abrir el detalle es liviano; cortamos primero el spinner.
            telegram_request(
                "answerCallbackQuery",
                {"callback_query_id": callback_id},
            )
            mostrar_detalle_operacion(chat_id, message_id, solicitud_id)
            return {"status": "detalle_operacion"}

        # ----------------------------------------------------
        # DECISIÓN: MOVER / ELIMINAR SOLICITADO DESDE LA WEB
        # ----------------------------------------------------

        if (
            action_data.startswith("opap:")
            or action_data.startswith("opre:")
        ):
            prefijo, solicitud_raw = action_data.split(":", 1)

            try:
                solicitud_id = str(UUID(solicitud_raw.strip()))
            except (ValueError, TypeError, AttributeError):
                telegram_request(
                    "answerCallbackQuery",
                    {
                        "callback_query_id": callback_id,
                        "text": "❌ ID de solicitud inválido.",
                        "show_alert": True,
                    },
                )
                return {"status": "callback_error"}

            aprobar_operacion = prefijo == "opap"

            # Programamos el trabajo pesado y respondemos AL MISMO webhook con
            # answerCallbackQuery. Así Telegram corta el spinner sin esperar una
            # segunda conexión HTTP saliente desde Railway.
            background_tasks.add_task(
                procesar_operacion_telegram_background,
                solicitud_id,
                aprobar_operacion,
                str(user_id),
                chat_id,
                message_id,
            )

            # Telegram permite ejecutar un método Bot API directamente como
            # respuesta al webhook. Esto es más rápido que HTTP_CLIENT.post().
            return {
                "method": "answerCallbackQuery",
                "callback_query_id": callback_id,
                "text": "⏳ Procesando operación...",
                "show_alert": False,
            }

        # NAVEGACIÓN: DETALLE DE UNA CARPETA / LOTE
        # ----------------------------------------------------

        if action_data.startswith(
            "lot:"
        ):

            lote_id_raw = (
                action_data
                .split(
                    ":",
                    1
                )[1]
                .strip()
            )


            try:

                lote_id = str(
                    UUID(
                        lote_id_raw
                    )
                )


            except (
                ValueError,
                TypeError,
                AttributeError
            ):

                telegram_request(

                    "answerCallbackQuery",

                    {
                        "callback_query_id":
                            callback_id,

                        "text":
                            "❌ ID de lote inválido.",

                        "show_alert":
                            True,
                    },

                )


                return {
                    "status":
                        "callback_error"
                }


            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id
                },

            )


            mostrar_detalle_lote(
                chat_id,
                message_id,
                lote_id
            )


            return {
                "status":
                    "detalle_lote"
            }


        # ----------------------------------------------------
        # DECISIÓN: APROBAR / RECHAZAR CARPETA COMPLETA
        # ----------------------------------------------------

        if (
            action_data.startswith(
                "aplot:"
            )
            or
            action_data.startswith(
                "relot:"
            )
        ):

            accion_lote, lote_id_raw = (
                action_data.split(
                    ":",
                    1
                )
            )


            try:

                lote_id = str(
                    UUID(
                        lote_id_raw.strip()
                    )
                )


            except (
                ValueError,
                TypeError,
                AttributeError
            ):

                telegram_request(

                    "answerCallbackQuery",

                    {
                        "callback_query_id":
                            callback_id,

                        "text":
                            "❌ ID de lote inválido.",

                        "show_alert":
                            True,
                    },

                )


                return {
                    "status":
                        "callback_error"
                }


            aprobar_lote = (
                accion_lote
                == "aplot"
            )


            estado_lote = (

                "APROBADO"

                if aprobar_lote

                else

                "RECHAZADO"

            )


            drive_lote = None
            documentos_lote = []


            try:

                resultado_decision_lote = (
                    resolver_custodia_carpeta(
                        lote_id,
                        aprobar_lote,
                    )
                )


                if not resultado_decision_lote[
                    "procesado"
                ]:

                    telegram_request(

                        "answerCallbackQuery",

                        {
                            "callback_query_id":
                                callback_id,

                            "text":
                                (
                                    "⚠️ Esta carpeta ya fue procesada "
                                    "o no tiene pendientes."
                                ),

                            "show_alert":
                                True,
                        },

                    )


                    return {
                        "status":
                            "already_processed"
                    }


                documentos_lote = (
                    resultado_decision_lote[
                        "documentos"
                    ]
                )


                drive_lote = (
                    resultado_decision_lote[
                        "drive_lote"
                    ]
                )


            except Exception as error:

                print(
                    "[CALLBACK LOTE ERROR]",
                    error
                )


                telegram_request(

                    "answerCallbackQuery",

                    {
                        "callback_query_id":
                            callback_id,

                        "text":
                            (
                                "❌ No se pudo procesar "
                                "la carpeta: "
                                f"{error}"
                            )[:200],

                        "show_alert":
                            True,
                    },

                )


                return {

                    "status":
                        "callback_lote_error",

                    "error":
                        str(error),

                }


            primer = (
                documentos_lote[0]
            )


            solicitante_id_lote = str(
                primer.get(
                    "solicitante_id"
                )
                or ""
            ).strip()


            solicitante_nombre_lote = (
                primer.get(
                    "solicitante_nombre"
                )
                or
                "Usuario"
            )


            nombre_carpeta = (

                obtener_carpeta_desde_ruta(
                    primer.get(
                        "ruta_relativa"
                    )
                )

                or

                f"LOTE_{lote_id[:8]}"

            )


            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,

                    "text":
                        (
                            "Carpeta actualizada a: "
                            f"{estado_lote}"
                        ),
                },

            )


            icono = (

                "✅"

                if estado_lote
                == "APROBADO"

                else

                "❌"

            )


            nuevo_texto = (
                "🛡️ DataVault DLP - GM Ingenieros\n\n"
                f"📁 Carpeta: {nombre_carpeta}\n"
                f"👤 Solicitante: {solicitante_nombre_lote}\n"
                f"📄 Archivos procesados: "
                f"{len(documentos_lote)}\n\n"
                f"{icono} DECISIÓN: "
                f"Carpeta {estado_lote}\n"
            )


            if drive_lote:

                nuevo_texto += (
                    "☁️ Carpeta subida a Google Drive\n"
                    f"🆔 Drive folder ID: "
                    f"{drive_lote['folder_id']}\n"
                )


            nuevo_texto += (
                f"👤 Procesado por "
                f"Telegram ID: {user_id}"
            )


            botones_finales_lote = []


            if solicitante_id_lote:

                botones_finales_lote.append([
                    {
                        "text":
                            "🔙 Pendientes del usuario",

                        "callback_data":
                            f"usr:{solicitante_id_lote}",
                    }
                ])


            botones_finales_lote.append([
                {
                    "text":
                        "👥 Usuarios pendientes",

                    "callback_data":
                        "menu:usuarios",
                }
            ])


            telegram_request(

                "editMessageText",

                {
                    "chat_id":
                        chat_id,

                    "message_id":
                        message_id,

                    "text":
                        nuevo_texto,

                    "reply_markup": {
                        "inline_keyboard":
                            botones_finales_lote
                    },
                },

            )

            background_tasks.add_task(
                actualizar_alertas_subida_telegram,
                None,
                lote_id,
                estado_lote,
            )
            background_tasks.add_task(
                refrescar_bandejas_telegram,
                chat_id,
            )

            return {

                "status":
                    "ok",

                "estado_actualizado":
                    estado_lote,

                "lote_id":
                    lote_id,

                "drive_folder_id":
                    (
                        drive_lote[
                            "folder_id"
                        ]
                        if drive_lote
                        else None
                    ),

            }


        # ----------------------------------------------------
        # DECISIÓN: APROBAR / RECHAZAR ARCHIVO SUELTO
        # ----------------------------------------------------

        if ":" not in action_data:

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,

                    "text":
                        "❌ Formato de decisión inválido.",

                    "show_alert":
                        True,
                },

            )


            return {

                "status":
                    "callback_error",

                "error":
                    "Formato de callback inválido",

            }


        accion, auditoria_id_raw = (
            action_data.split(
                ":",
                1
            )
        )


        if accion not in (
            "aprobar",
            "rechazar"
        ):

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,

                    "text":
                        "❌ Acción inválida.",

                    "show_alert":
                        True,
                },

            )


            return {

                "status":
                    "callback_error",

                "error":
                    "Acción inválida",

            }


        try:

            auditoria_id = str(
                UUID(
                    auditoria_id_raw.strip()
                )
            )


        except (
            ValueError,
            AttributeError,
            TypeError
        ):

            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,

                    "text":
                        "❌ ID de auditoría inválido.",

                    "show_alert":
                        True,
                },

            )


            return {

                "status":
                    "callback_error",

                "error":
                    "ID de auditoría inválido",

            }


        id_auditoria_str = (
            auditoria_id
        )


        nuevo_estado = (

            "APROBADO"

            if accion
            == "aprobar"

            else

            "RECHAZADO"

        )


        drive_id = None
        resultado_update = None
        registro_actual = None


        try:

            resultado_decision = (
                resolver_custodia_archivo(

                    auditoria_id,

                    accion
                    == "aprobar",

                )
            )


            registro_actual = (
                resultado_decision[
                    "registro"
                ]
            )


            drive_id = (
                resultado_decision[
                    "drive_id"
                ]
            )


            resultado_update = (
                resultado_decision[
                    "resultado_update"
                ]
            )


            if not resultado_decision[
                "procesado"
            ]:

                telegram_request(

                    "answerCallbackQuery",

                    {
                        "callback_query_id":
                            callback_id,

                        "text":
                            (
                                "⚠️ Este documento ya fue procesado. "
                                f"Estado actual: "
                                f"{resultado_decision['estado']}."
                            ),

                        "show_alert":
                            True,
                    },

                )


                return {

                    "status":
                        "already_processed",

                    "estado_actual":
                        resultado_decision[
                            "estado"
                        ],

                }


        except Exception as error:

            print(
                "[CALLBACK ERROR]",
                error
            )


            telegram_request(

                "answerCallbackQuery",

                {
                    "callback_query_id":
                        callback_id,

                    "text":
                        (
                            "❌ No se pudo procesar: "
                            f"{error}"
                        )[:200],

                    "show_alert":
                        True,
                },

            )


            return {

                "status":
                    "callback_error",

                "error":
                    str(error),

            }


        print(

            "[SUPABASE UPDATE]",

            getattr(
                resultado_update,
                "data",
                None
            ),

        )


        telegram_request(

            "answerCallbackQuery",

            {
                "callback_query_id":
                    callback_id,

                "text":
                    (
                        "Estado actualizado a: "
                        f"{nuevo_estado}"
                    ),
            },

        )


        solicitante_id = str(

            (
                registro_actual
                or {}
            ).get(
                "solicitante_id"
            )

            or ""

        ).strip()


        nombre_archivo = (

            (
                registro_actual
                or {}
            ).get(
                "nombre_archivo"
            )

            or

            "Archivo"

        )


        solicitante_nombre = (

            (
                registro_actual
                or {}
            ).get(
                "solicitante_nombre"
            )

            or

            "Usuario"

        )


        icono = (

            "✅"

            if nuevo_estado
            == "APROBADO"

            else

            "❌"

        )


        nuevo_texto = (
            "🛡️ DataVault DLP - GM Ingenieros\n\n"
            f"📁 Archivo: {nombre_archivo}\n"
            f"👤 Solicitante: {solicitante_nombre}\n\n"
            f"{icono} DECISIÓN: "
            f"Documento {nuevo_estado}\n"
        )


        if drive_id:

            nuevo_texto += (
                "☁️ Subido a Google Drive "
                f"(ID: {drive_id})\n"
            )


        nuevo_texto += (
            f"👤 Procesado por "
            f"Telegram ID: {user_id}"
        )


        botones_finales = []


        if solicitante_id:

            botones_finales.append([
                {
                    "text":
                        "🔙 Archivos del usuario",

                    "callback_data":
                        f"usr:{solicitante_id}",
                }
            ])


        botones_finales.append([
            {
                "text":
                    "👥 Usuarios pendientes",

                "callback_data":
                    "menu:usuarios",
            }
        ])


        if chat_id and message_id:

            telegram_request(

                "editMessageText",

                {
                    "chat_id":
                        chat_id,

                    "message_id":
                        message_id,

                    "text":
                        nuevo_texto,

                    "reply_markup": {
                        "inline_keyboard":
                            botones_finales
                    },
                },

            )

        background_tasks.add_task(
            actualizar_alertas_subida_telegram,
            auditoria_id,
            None,
            nuevo_estado,
        )
        background_tasks.add_task(
            refrescar_bandejas_telegram,
            chat_id,
        )

        return {

            "status":
                "ok",

            "estado_actualizado":
                nuevo_estado,

            "autorizado_por":
                user_id,

            "drive_id":
                drive_id,

        }


    return {
        "status":
            "ignored"
    }
