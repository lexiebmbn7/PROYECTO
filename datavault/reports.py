from __future__ import annotations

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .states import ESTADO_APROBADO, ESTADO_PENDIENTE, ESTADO_RECHAZADO, normalizar_estado


def _safe_select_tabla(db, tabla: str, columnas: str = "*") -> list:
    try:
        respuesta = db.table(tabla).select(columnas).execute()
        return respuesta.data or []
    except Exception as error:
        print(f"[REPORT TABLE ERROR] {tabla}", error)
        return []


def _valor_excel(valor):
    if valor is None:
        return ""
    if isinstance(valor, (dict, list, tuple, set)):
        return json.dumps(valor, ensure_ascii=False, default=str)
    return valor


def _formatear_hoja_excel(ws):
    if ws.max_row < 1:
        return

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    borde = Border(bottom=Side(style="thin", color="B7B7B7"))

    for celda in ws[1]:
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for columna in ws.columns:
        letra = get_column_letter(columna[0].column)
        max_len = 0
        for celda in columna:
            texto = str(celda.value or "")
            max_len = max(max_len, min(len(texto), 60))
            if isinstance(celda.value, datetime):
                celda.number_format = "yyyy-mm-dd hh:mm:ss"
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[letra].width = max(12, min(max_len + 2, 55))


def _crear_hoja_datos(wb: Workbook, titulo: str, columnas: list, filas: list):
    ws = wb.create_sheet(title=titulo[:31])
    ws.append([col[0] for col in columnas])
    for fila in filas:
        ws.append([_valor_excel(fila.get(col[1])) for col in columnas])
    _formatear_hoja_excel(ws)
    return ws


def generar_reporte_excel_datavault(db, *, generado_en: str) -> bytes:
    auditoria = _safe_select_tabla(db, "auditoria_custodia")
    operaciones = _safe_select_tabla(db, "solicitudes_operacion")
    eliminados = _safe_select_tabla(db, "elementos_eliminados")
    eventos = _safe_select_tabla(db, "auditoria_eventos_v2")
    dlp = _safe_select_tabla(db, "dlp_auditoria")

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    total = len(auditoria)
    aprobados = sum(1 for x in auditoria if normalizar_estado(x.get("estado")) == ESTADO_APROBADO)
    pendientes = sum(1 for x in auditoria if normalizar_estado(x.get("estado")) == ESTADO_PENDIENTE)
    rechazados = sum(1 for x in auditoria if normalizar_estado(x.get("estado")) == ESTADO_RECHAZADO)

    ws.append(["INDICADOR", "VALOR"])
    ws.append(["Generado UTC", generado_en])
    ws.append(["Documentos auditados", total])
    ws.append(["Aprobados", aprobados])
    ws.append(["Pendientes", pendientes])
    ws.append(["Rechazados", rechazados])
    ws.append(["Operaciones registradas", len(operaciones)])
    ws.append(["Elementos eliminados", len(eliminados)])
    ws.append(["Eventos de auditoría", len(eventos)])
    ws.append(["Cambios DLP", len(dlp)])
    _formatear_hoja_excel(ws)

    _crear_hoja_datos(wb, "Custodia", [
        ("ID", "id"),
        ("Archivo", "nombre_archivo"),
        ("Usuario", "solicitante_nombre"),
        ("Correo", "solicitante_correo"),
        ("Estado", "estado"),
        ("Procesamiento", "estado_procesamiento"),
        ("Error procesamiento", "ultimo_error_procesamiento"),
        ("Estado archivo", "estado_archivo"),
        ("Tamaño bytes", "tamano_bytes"),
        ("SHA256", "hash_sha256"),
        ("Ruta relativa", "ruta_relativa"),
        ("Destino solicitado", "destino_solicitado_ruta"),
        ("Destino final", "destino_final_ruta"),
        ("Ubicación Drive", "ubicacion_drive"),
        ("Fecha solicitud", "fecha_solicitud"),
        ("Fecha transferencia", "fecha_transferencia"),
        ("Fecha eliminación", "fecha_eliminacion"),
    ], auditoria)

    _crear_hoja_datos(wb, "Operaciones", [
        ("ID", "id"),
        ("Tipo", "tipo_operacion"),
        ("Objeto", "objeto_tipo"),
        ("Nombre", "nombre_objeto"),
        ("Usuario", "solicitante_nombre"),
        ("Origen", "carpeta_origen"),
        ("Destino solicitado", "destino_solicitado_ruta"),
        ("Destino final", "destino_final_ruta"),
        ("Estado", "estado"),
        ("Procesamiento", "estado_procesamiento"),
        ("Error procesamiento", "ultimo_error_procesamiento"),
        ("Fecha solicitud", "fecha_solicitud"),
        ("Fecha resolución", "fecha_resolucion"),
        ("Resultado", "resultado"),
    ], operaciones)

    _crear_hoja_datos(wb, "Eliminaciones", [
        ("ID", "id"),
        ("Tipo", "objeto_tipo"),
        ("Nombre", "nombre_objeto"),
        ("Usuario", "solicitante_nombre"),
        ("Correo", "solicitante_correo"),
        ("Ruta anterior", "ubicacion_anterior"),
        ("Estado", "estado"),
        ("Fecha eliminación", "fecha_eliminacion"),
        ("Detalle", "detalle"),
    ], eliminados)

    _crear_hoja_datos(wb, "Auditoria", [
        ("Evento", "evento"),
        ("Categoría", "categoria"),
        ("Acción", "accion"),
        ("Actor", "actor_nombre"),
        ("Correo", "actor_correo"),
        ("Rol", "actor_rol"),
        ("Objeto", "objeto_tipo"),
        ("Objeto ID", "objeto_id"),
        ("Estado", "estado"),
        ("IP", "ip"),
        ("Fecha", "created_at"),
        ("Detalle", "detalle"),
    ], eventos)

    _crear_hoja_datos(wb, "Cambios DLP", [
        ("Política", "politica_codigo"),
        ("Acción", "accion"),
        ("Valor anterior", "valor_anterior"),
        ("Valor nuevo", "valor_nuevo"),
        ("Actor", "actor_nombre"),
        ("Correo", "actor_correo"),
        ("Fecha", "fecha"),
    ], dlp)

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida.getvalue()
